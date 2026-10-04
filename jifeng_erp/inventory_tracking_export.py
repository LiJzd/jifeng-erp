"""Export the complete checked tracking scope as literal Excel cells.

Balances use the same native cutoff projection as the two workbenches. Archive
prices, weights and units never recompute balances or unknown total weights.
No browser pages, ledger identifiers, source images or attachment URLs export.
"""

from io import BytesIO

import frappe
from openpyxl import Workbook
from openpyxl.comments import Comment

from jifeng_erp import inventory_tracking as T
from jifeng_erp.tool_archive_export import CONTENT_TYPE, _freeze, _fresh_snapshot


MAX_EXPORT_ROWS = 1_048_575
LABELS = {"warehouse": "仓库", "legacy_number": "编号", "stock_qty": "库存",
          "valuation_rate": "单价", "stock_value": "金额", "weight": "重量",
          "total_weight": "总重量", "major_category": "大类", "minor_category": "小类",
          "material": "材质", "stock_uom": "单位", "display_name": "名称",
          "specification": "规格", "location": "库位", "notes": "备注", "supplier": "供应商",
          "shape": "形状", "color": "颜色", "old_number": "旧编号",
          "owner_name": "负责", "region": "区域"}
SOURCES = (T.catalog.PROFILE, "Item", "Warehouse", "Stock Ledger Entry", "Company", "File",
           "Stock Entry", "Stock Reconciliation", "Stock Reconciliation Item", "Inventory Dimension")


def _deny():
    frappe.throw("没有完整读取及导出库存跟踪档案和库存流水的权限", frappe.PermissionError)


def _permission_state():
    user, state = frappe.session.user, []
    # Document permissions use the request's cached System Settings. A fresh
    # connection alone does not refresh that cache; compare it with the uncached
    # Check value before trusting either permission pass.
    strict = frappe.local.db.get_single_value("System Settings", "apply_strict_user_permissions", cache=False)
    effective_strict = frappe.get_system_settings("apply_strict_user_permissions")
    if bool(strict) != bool(effective_strict):
        _deny()
    state.append(("apply_strict_user_permissions", bool(strict), bool(effective_strict)))
    for doctype, filters, fields in (
        ("User", {"name": user}, ["name", "enabled", "user_type"]),
        ("Has Role", {"parenttype": "User", "parent": user}, ["*"]),
        ("User Permission", {"user": user}, ["*"]),
        ("DocPerm", {"parent": ["in", SOURCES]}, ["*"]),
        ("Custom DocPerm", {"parent": ["in", SOURCES]}, ["*"]),
        ("DocType", {"name": ["in", SOURCES]}, ["*"]),
        ("DocField", {"parent": ["in", SOURCES]}, ["*"]),
        ("Custom Field", {"dt": ["in", SOURCES]}, ["*"]),
        ("Property Setter", {"doc_type": ["in", SOURCES]}, ["*"]),
        ("DocShare", {"share_doctype": ["in", SOURCES]}, ["*"]),
    ):
        rows = frappe.get_all(doctype, filters=filters, fields=fields, limit_page_length=0)
        if doctype == "DocShare":
            rows = [row for row in rows if row.get("user") == user or row.get("everyone")]
        state.append((doctype, tuple(sorted((_freeze(row) for row in rows), key=repr))))
    return _freeze((user, state))


def _export_scope(kind, company):
    T.catalog._scope(kind, company)
    for doctype in (T.catalog.PROFILE, "Stock Ledger Entry"):
        if not frappe.permissions.can_export(doctype):
            _deny()


def _checked(kind, company, search, warehouse, as_of, conditions, order):
    _export_scope(kind, company)
    rows, entries, currency, cutoff = T._checked_tracking(kind, company, search, warehouse, as_of)
    # Export permission is independent of read. It covers every contributing
    # source before filters; an excluded pair must not hide a permission gap.
    for doctype, names in (
        (T.catalog.PROFILE, {row["profile_name"] for row in rows}),
        ("Stock Ledger Entry", {row["name"] for row in entries}),
    ):
        for name in sorted(names):
            frappe.get_doc(doctype, name).check_permission("export")
    matched = T._column_rows(rows, conditions, order)
    if len(matched) > MAX_EXPORT_ROWS:
        frappe.throw("库存跟踪超过Excel行数限制，请缩小筛选范围后完整导出")
    values = [[T._projected_value(row, field) for field in T.COLUMN_FIELDS[kind]] for row in matched]
    _export_scope(kind, company)
    return _freeze((rows, entries, currency, cutoff)), values, currency, cutoff


def _text(value):
    if value is None:
        return ""
    if not isinstance(value, str) or any(not (char in "\t\n\r" or " " <= char <= "\ud7ff"
            or "\ue000" <= char <= "\ufffd" or "\U00010000" <= char <= "\U0010ffff") for char in value):
        frappe.throw("库存跟踪包含无法完整导出的文本，请先核对来源")
    if len(value.encode("utf-16-le")) // 2 > 32767:
        frappe.throw("库存跟踪包含超过Excel单元格限制的文本，请缩小内容后导出")
    return value


def _cell(sheet, row, column, value):
    cell = sheet.cell(row, column)
    cell.value = _text(value)
    cell.data_type, cell.number_format = "s", "@"
    return cell


def _xlsx(kind, company, search, warehouse, conditions, order, checked):
    _, values, currency, cutoff = checked
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = kind + "跟踪"
        fields = T.COLUMN_FIELDS[kind]
        for column, field in enumerate(fields, 1):
            cell = _cell(sheet, 1, column, LABELS[field])
            basis = {"stock_qty": "stock_qty", "valuation_rate": "valuation_rate", "stock_value": "stock_value",
                     "stock_uom": "stock_uom", "weight": "profile_weight", "total_weight": "total_weight"}.get(field)
            if basis:
                cell.comment = Comment(T.BASIS[basis] + "；币种：" + currency + "。未知为空，合法零保留。", "集沣库存跟踪")
        for row_number, row in enumerate(values, 2):
            for column, value in enumerate(row, 1):
                _cell(sheet, row_number, column, value)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{sheet.cell(max(1, len(values) + 1), len(fields)).coordinate}"
        notes = workbook.create_sheet("查询条件")
        for index, pair in enumerate((
            ("公司", company), ("截止日期", cutoff), ("仓库", warehouse or "全部可读叶仓库"),
            ("关键词", search), ("列筛选", "见列筛选工作表；仅表头表示未应用列筛选"),
            ("排序", LABELS[order["field"]] + ("：升序" if order["direction"] == "asc" else "：降序") if order else "默认顺序"),
            ("币种", currency), ("数据行数", str(len(values))),
            ("范围", "当前筛选下全部匹配库存行，使用页面数据列和排序；不按当前页截断。"),
            ("数量金额", "库存单位、数量、估价单价和金额使用原生截止日库存账簿；精确十进制文本。"),
            ("重量", "重量为授权档案原值，总重量换算尚未核实，留空；不按档案单价估值。"),
        ), 1):
            for column, value in enumerate(pair, 1):
                _cell(notes, index, column, value)
        filters = workbook.create_sheet("列筛选")
        for column, value in enumerate(("列", "条件", "值类型", "值"), 1):
            _cell(filters, 1, column, value)
        operators = dict(zip(("eq", "ne", "starts", "not_starts", "ends", "not_ends", "contains", "not_contains", "in", "not_in"),
                             ("等于", "不等于", "开头是", "开头不是", "结尾是", "结尾不是", "包含", "不包含", "在里面", "不在里面")))
        index = 2
        # One bounded value per cell preserves legal 65K conditions without
        # truncating them or exceeding Excel's UTF-16 cell limit.
        for condition in conditions:
            selected = condition["value"] if condition["operator"] in {"in", "not_in"} else [condition["value"]]
            for value in selected or [None]:
                value_type = "空选择列表" if not selected else "空值" if value is None else "文本"
                for column, text in enumerate((LABELS[condition["field"]], operators[condition["operator"]], value_type, value), 1):
                    _cell(filters, index, column, text)
                index += 1
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


@frappe.whitelist(methods=["GET", "POST"])
def export_tracking(kind, company, search="", warehouse=None, as_of=None, column_filters=None, sort=None):
    conditions, order, _ = T._column_query(kind, column_filters, sort)
    # Bind a default date once so midnight cannot change the comparison scope.
    cutoff = str(T._date(as_of))
    _export_scope(kind, company)
    permissions = _permission_state()
    checked = _checked(kind, company, search, warehouse, cutoff, conditions, order)
    content = _xlsx(kind, company, search, warehouse, conditions, order, checked)
    with _fresh_snapshot():
        if permissions != _permission_state():
            _deny()
        if checked != _checked(kind, company, search, warehouse, cutoff, conditions, order):
            frappe.throw("库存跟踪来源或可见范围已改变，请重新查询后导出")
        if permissions != _permission_state():
            _deny()
    frappe.response.update({"type": "download", "content_type": CONTENT_TYPE, "filename": kind + "跟踪.xlsx",
                            "display_content_as": "attachment", "filecontent": content})
