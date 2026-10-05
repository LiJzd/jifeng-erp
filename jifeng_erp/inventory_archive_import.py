"""Bounded XLSX archive import using the existing native creation receipts."""

import base64
import binascii
from datetime import date, datetime, time
import hashlib
from io import BytesIO
import json
import math
import re
from uuid import UUID
from xml.etree import ElementTree as ET
import zipfile

import frappe
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, range_boundaries

from jifeng_erp import inventory_archive as A


MAX_FILE_BYTES = 1024 * 1024
MAX_ROWS = 100
_FIELDS = (
    "legacy_number", "legacy_code", "major_category", "minor_category", "old_number",
    "specification", "display_name", "location", "material", "shape", "color", "display_uom",
    "list_price", "processing_price", "weight", "supplier", "upper_limit", "lower_limit",
    "conversion_uom", "conversion_value", "owner_name", "region", "notes", "inactive",
)
_HEX = re.compile(r"^[0-9a-f]{64}$")
_INSTRUCTIONS = "导入说明"
_MARKER = "集沣库存档案导入说明"


def _hash(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                     separators=(",", ":")).encode("utf-8")).hexdigest()


def _row_id(batch_id, row_index):
    value = "jifeng-inventory-import:v1:" + batch_id + ":" + str(row_index)
    return str(UUID(bytes=hashlib.sha256(value.encode("utf-8")).digest()[:16], version=4))


def _metadata(kind, company, batch_id, row_digests):
    return {"version": 1, "kind": kind, "company": company, "batch_id": batch_id,
            "row_count": len(row_digests), "row_digests": row_digests}


def _scope(kind, company, create=True):
    A._login()
    A._kind(kind)
    A._company(company)
    profile = frappe.new_doc(A.PROFILE)
    profile.company, profile.archive_kind = company, kind
    A._permission(profile, "create" if create else "read", role_only=True)
    if create:
        A._permission(profile, "create")
        A._permission(frappe.new_doc("Item"), "create", role_only=True)
    return profile


def _columns(kind):
    excluded = {"shape", "color"} if kind == "物料" else {"material"}
    readable, writable = A._columns(A.PROFILE, "read"), A._columns(A.PROFILE, "write")
    result = []
    for name in _FIELDS:
        if name in excluded:
            continue
        field = frappe.get_meta(A.PROFILE).get_field(name)
        if not field or name not in A.PROFILE_EDIT:
            frappe.throw("库存档案导入字段尚未安装")
        result.append({"fieldname": name, "label": field.label, "fieldtype": field.fieldtype,
                       "required": bool(field.reqd), "options": field.options or "",
                       "read_only": bool(field.read_only or name not in readable & writable)})
    return result


@frappe.whitelist(methods=["GET"])
def init_import(kind, company):
    profile = _scope(kind, company)
    columns = _columns(kind)
    if any(column["read_only"] for column in columns if column["fieldname"] in ("legacy_number", "display_name")):
        frappe.throw("没有导入编号及名称字段的权限", frappe.PermissionError)
    item_fields = A._columns("Item", "read") & A._columns("Item", "write")
    item_meta = frappe.get_meta("Item")
    selectable = {name: bool(name in item_fields and item_meta.get_field(name)
                             and not item_meta.get_field(name).read_only
                             and frappe.has_permission(doctype, "select"))
                  for name, doctype in (("item_group", "Item Group"), ("stock_uom", "UOM"))}
    return {"kind": kind, "company": company, "columns": columns,
            "can_create_item": bool(profile.has_permission("create") and frappe.has_permission("Item", "create")),
            "can_select_item_group": selectable["item_group"], "can_select_stock_uom": selectable["stock_uom"],
            "max_file_bytes": MAX_FILE_BYTES, "max_rows": MAX_ROWS}


@frappe.whitelist(methods=["GET"])
def download_template(kind, company):
    config = init_import(kind, company)
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = kind + "档案"
        columns = [column for column in config["columns"] if not column["read_only"]]
        for index, column in enumerate(columns, 1):
            cell = sheet.cell(1, index, column["label"])
            cell.data_type, cell.number_format = "s", "@"
            # Empty text-formatted cells protect leading zeroes when filling the template.
            for row in range(2, MAX_ROWS + 2):
                sheet.cell(row, index).number_format = "@" if column["fieldtype"] not in ("Float", "Currency", "Int", "Check") else "General"
        notes = workbook.create_sheet(_INSTRUCTIONS)
        for text in (_MARKER, "编号、代码、旧号及名称请按文本填写，保留前导零。编号及名称必填。",
                     "每个文件最多100条数据，只能有一个可见数据页；超过时请分文件导入。",
                     "商品分组及真实库存单位在导入窗口明确选择；表内单位是档案原文。",
                     "供应商、换算单位填写有选择权限的系统ID，不按显示名称匹配。",
                     "不导入库存余额、财务金额凭证、图片、公式或链接；请勿添加其他字段。"):
            notes.append([text])
        output = BytesIO()
        workbook.save(output)
        frappe.response.update({"type": "download", "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                "filename": kind + "档案导入模板.xlsx", "display_content_as": "attachment",
                                "filecontent": output.getvalue()})
    finally:
        workbook.close()


def _xlsx_bytes(filename, content_base64):
    if not isinstance(filename, str) or not filename.lower().endswith(".xlsx"):
        frappe.throw("请选择 .xlsx 文件")
    if not isinstance(content_base64, str) or len(content_base64) > ((MAX_FILE_BYTES + 2) // 3) * 4:
        frappe.throw("单个.xlsx文件不能超过1MiB")
    try:
        raw = base64.b64decode(content_base64, validate=True)
    except (ValueError, binascii.Error):
        frappe.throw("文件内容不是有效的Base64")
    if not raw or len(raw) > MAX_FILE_BYTES:
        frappe.throw("单个.xlsx文件不能超过1MiB")
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            members = archive.infolist()
            names = [member.filename for member in members]
            if len(members) > 128 or len(names) != len(set(names)) or sum(member.file_size for member in members) > 8 * MAX_FILE_BYTES:
                frappe.throw("文件解包大小或数量超过导入限额")
            if not {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels"}.issubset(names):
                frappe.throw("无法读取.xlsx文件，请使用本系统模板")
            for member in members:
                name = member.filename
                if (name.startswith("/") or "\\" in name or ":" in name or ".." in name.split("/")
                        or member.flag_bits & 1 or member.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED)
                        or member.file_size > 2 * MAX_FILE_BYTES
                        or member.file_size > max(member.compress_size, 1) * 200):
                    frappe.throw("文件压缩结构不受支持或超过限额")
                lowered = name.lower()
                if any(token in lowered for token in ("vba", "macro", "activex", "externallinks/", "embeddings/", "media/", "drawings/")):
                    frappe.throw("导入文件不能含宏、外部链接或图片")
                if not lowered.endswith((".xml", ".rels")):
                    continue
                data = archive.read(member)
                if b"\x00" in data or b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
                    frappe.throw("文件XML结构不受支持")
                root = ET.fromstring(data)
                cells = 0
                for node in root.iter():
                    tag = node.tag.rsplit("}", 1)[-1]
                    if tag in ("f", "hyperlink", "hyperlinks"):
                        frappe.throw("导入文件不能含公式或链接")
                    if tag == "Relationship" and (node.get("TargetMode", "").lower() == "external"
                            or any(token in node.get("Type", "").lower() for token in ("hyperlink", "externallink", "vba", "activex", "oleobject"))
                            or ":" in node.get("Target", "")):
                        frappe.throw("导入文件不能含外部关系或链接")
                    if tag in ("Default", "Override") and any(token in node.get("ContentType", "").lower() for token in ("macro", "vba", "activex")):
                        frappe.throw("导入文件不能含宏")
                    if lowered.startswith("xl/worksheets/"):
                        if tag == "row" and (not re.fullmatch(r"[1-9][0-9]{0,3}", node.get("r", "")) or int(node.get("r")) > 1024):
                            frappe.throw("数据页行列范围超过导入限额")
                        if tag == "mergeCell":
                            frappe.throw("导入数据页不能含合并单元格")
                        if tag == "dimension":
                            bounds = range_boundaries(node.get("ref", "A1"))
                            if any(value is None for value in bounds) or bounds[2] > 64 or bounds[3] > 1024:
                                frappe.throw("数据页行列范围超过导入限额")
                        if tag == "c":
                            cells += 1
                            col, row = coordinate_from_string(node.get("r", "A1"))
                            if cells > 4096 or row > 1024 or column_index_from_string(col) > 64:
                                frappe.throw("数据页行列范围超过导入限额")
    except (zipfile.BadZipFile, ET.ParseError, ValueError, KeyError, OSError) as error:
        # frappe.ValidationError must keep its deliberate Chinese validation message.
        if isinstance(error, getattr(frappe, "ValidationError", type(None))):
            raise
        frappe.throw("无法读取.xlsx文件，请使用本系统模板")
    return raw


def _read_rows(kind, raw):
    workbook = None
    try:
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        sheets = []
        for sheet in workbook.worksheets:
            if sheet.title == _INSTRUCTIONS and sheet["A1"].value == _MARKER:
                sheet.reset_dimensions()
                if any(cell.value not in (None, "") and (index != 0 or not isinstance(cell.value, str))
                       for row in sheet.iter_rows() for index, cell in enumerate(row)):
                    frappe.throw("导入说明页只能包含单列说明文本，不能含数据")
                continue
            if sheet.sheet_state != "visible":
                frappe.throw("不能导入隐藏的数据页")
            sheets.append(sheet)
        if len(sheets) != 1:
            frappe.throw("导入文件必须有且仅有一个可见数据页")
        sheet = sheets[0]
        # Some spreadsheet writers emit a stale dimension. Read the actual rows.
        sheet.reset_dimensions()
        rows = sheet.iter_rows()
        header = next(rows, ())
        labels = {column["label"]: column["fieldname"] for column in _columns(kind)}
        headers = [cell.value for cell in header]
        populated = [value for value in headers if value not in (None, "")]
        if len(populated) != len(set(populated)) or any(not isinstance(value, str) or value not in labels for value in populated):
            frappe.throw("表头重复或含不支持的字段（图片、余额及未知列不能导入），请使用本系统模板")
        if not {"legacy_number", "display_name"}.issubset({labels[value] for value in populated}):
            frappe.throw("表头必须包含编号及名称")
        result = []
        for sheet_row, cells in enumerate(rows, 2):
            if not any(cell.value not in (None, "") for cell in cells):
                continue
            if any(cell.value not in (None, "") and (index >= len(headers) or headers[index] in (None, ""))
                   for index, cell in enumerate(cells)):
                frappe.throw("无表头或未知列含有数据，不能忽略；请使用本系统模板")
            result.append((sheet_row, {labels[headers[index]]: cell.value for index, cell in enumerate(cells)
                                       if index < len(headers) and headers[index] not in (None, "") and cell.value not in (None, "")}))
            if len(result) > MAX_ROWS:
                frappe.throw("每个文件最多100条数据，请分文件导入")
        if not result:
            frappe.throw("文件没有可导入的数据行")
        return result
    finally:
        if workbook:
            workbook.close()


def _prepare_row(kind, company, values, item_values):
    profile, item = frappe.new_doc(A.PROFILE), frappe.new_doc("Item")
    profile.company, profile.archive_kind = company, kind
    values = {"company": company, "conversion_value": 1, "inactive": 0, **values}
    normalized = A._normalized(A.PROFILE, values, A.PROFILE_EDIT | {"company"})
    A._updates(profile, normalized, A.PROFILE_EDIT | {"company"})
    A._permission(profile, "create")
    A._updates(item, item_values, A.ITEM_CREATE)
    item.item_name, item.is_stock_item, item.opening_stock, item.standard_rate = profile.display_name, 1, 0, 0
    A._permission(item, "create")
    for name in ("legacy_number", "display_name"):
        if not isinstance(profile.get(name), str) or not profile.get(name).strip():
            frappe.throw("请填写文本编号及名称；数值编号不能导入")
    for name in ("list_price", "processing_price", "weight", "upper_limit", "lower_limit", "conversion_value"):
        value = profile.get(name)
        if name == "conversion_value" and (value is None or value <= 0):
            frappe.throw("换算值必须大于0")
        if name != "conversion_value" and value is not None and value < 0:
            frappe.throw("价格、重量及库存上下限不能为负数")
    if profile.get("upper_limit") is not None and profile.get("lower_limit") is not None and profile.get("upper_limit") < profile.get("lower_limit"):
        frappe.throw("上限不能小于下限")
    return normalized


def _display_values(values):
    result = {}
    for name, value in values.items():
        if isinstance(value, (date, datetime, time)):
            value = value.isoformat()
        if isinstance(value, float) and not math.isfinite(value):
            value = str(value)
        if value is None or type(value) in (str, int, float, bool):
            result[name] = value
    return result


def _prepare(kind, company, filename, content_base64, item_group, stock_uom, batch_id):
    _scope(kind, company)
    A._uuid(batch_id)
    if not all(isinstance(value, str) and value.strip() for value in (item_group, stock_uom)):
        frappe.throw("请明确选择系统商品分组及真实库存单位")
    item_values = A._normalized("Item", {"item_group": item_group, "stock_uom": stock_uom}, A.ITEM_CREATE)
    source = _read_rows(kind, _xlsx_bytes(filename, content_base64))
    rows, errors, digests = [], [], []
    validation_errors = tuple(set([ValueError, frappe.PermissionError] + [getattr(frappe, name) for name in
        ("ValidationError", "DoesNotExistError", "LinkValidationError") if hasattr(frappe, name)]))
    for index, (sheet_row, values) in enumerate(source):
        row = {"row_index": index, "sheet_row": sheet_row, "request_id": _row_id(batch_id, index),
               "valid": False, "profile": _display_values(values)}
        try:
            normalized = _prepare_row(kind, company, values, item_values)
            row.update(valid=True, profile={key: value for key, value in normalized.items()
                                           if key in A.PROFILE_EDIT & A._columns(A.PROFILE, "read")})
            digests.append(A._digest(kind, normalized, item_values))
        except validation_errors as error:
            message = "没有创建档案或选择所填系统ID的权限" if isinstance(error, frappe.PermissionError) else (
                "所填系统ID不存在或不能选择" if isinstance(error, getattr(frappe, "DoesNotExistError", type(None))) else str(error))
            errors.append({"row_index": index, "sheet_row": sheet_row, "message": message})
        rows.append(row)
    metadata = _metadata(kind, company, batch_id, digests) if not errors else None
    return {"metadata": metadata, "preview_digest": _hash(metadata) if metadata else None,
            "row_count": len(source), "rows": rows, "errors": errors, "can_commit": not errors}, item_values


@frappe.whitelist(methods=["POST"])
def preview_import(kind, company, filename, content_base64, item_group, stock_uom, batch_id):
    return _prepare(kind, company, filename, content_base64, item_group, stock_uom, batch_id)[0]


@frappe.whitelist(methods=["POST"])
@A._atomic
def commit_import(kind, company, filename, content_base64, item_group, stock_uom, batch_id, preview_digest):
    preview, item_values = _prepare(kind, company, filename, content_base64, item_group, stock_uom, batch_id)
    if not preview["can_commit"]:
        frappe.throw("导入行存在错误，请修正后重新预览")
    if not isinstance(preview_digest, str) or not _HEX.fullmatch(preview_digest) or preview_digest != preview["preview_digest"]:
        frappe.throw("导入内容或默认设置已变化，请重新预览")
    for row in preview["rows"]:
        A.save_archive({"kind": kind, "request_id": row["request_id"],
                        "profile": {"company": company, **row["profile"]}, "item": item_values})
    return resolve_import(kind, company, batch_id, preview["row_count"], preview_digest, preview["metadata"]["row_digests"])


@frappe.whitelist(methods=["GET"])
def resolve_import(kind, company, batch_id, row_count, preview_digest, row_digests):
    _scope(kind, company, create=False)
    A._uuid(batch_id)
    if isinstance(row_count, str) and len(row_count) <= 3 and re.fullmatch(r"[1-9][0-9]*", row_count):
        row_count = int(row_count)
    if type(row_count) is not int or not 1 <= row_count <= MAX_ROWS:
        frappe.throw("批次行数必须为1到100")
    if isinstance(row_digests, str):
        if len(row_digests) > MAX_ROWS * 70:
            frappe.throw("批次核对信息超出限额")
        try:
            row_digests = json.loads(row_digests)
        except (TypeError, ValueError):
            frappe.throw("批次行校验必须是有效数组")
    if (not isinstance(row_digests, list) or len(row_digests) != row_count
            or any(not isinstance(value, str) or not _HEX.fullmatch(value) for value in row_digests)):
        frappe.throw("批次行校验信息不完整")
    metadata = _metadata(kind, company, batch_id, row_digests)
    if not isinstance(preview_digest, str) or not _HEX.fullmatch(preview_digest) or _hash(metadata) != preview_digest:
        frappe.throw("批次预览校验不匹配")
    receipts = []
    for index, digest in enumerate(row_digests):
        request_id = _row_id(batch_id, index)
        name = A._creation_name(request_id)
        receipt = {"row_index": index, "request_id": request_id, "status": "not_found", "name": None, "item_code": None}
        if name:
            profile, item = A._documents(name, kind)
            if profile.company != company or profile.request_id != request_id or profile.request_digest != digest:
                frappe.throw("批次行已用于不同内容，请核对原保存结果")
            visible = "item_code" in A._columns(A.PROFILE, "read") & A._columns("Item", "read")
            receipt.update(status="found", name=profile.name, item_code=item.name if visible else None)
        receipts.append(receipt)
    found = sum(row["status"] == "found" for row in receipts)
    return {"status": "found" if found == row_count else "partial" if found else "not_found",
            "metadata": metadata, "preview_digest": preview_digest, "row_count": row_count, "receipts": receipts}
