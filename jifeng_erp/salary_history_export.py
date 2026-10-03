"""Export the complete authorized saved payroll history as literal XLSX text.

No recomputation, payment inference, hidden-source omission, or page stitching.
The binary response is published only after an independent primary read-only
snapshot rechecks the complete source projections and native permissions.
"""

from decimal import Decimal, InvalidOperation
from io import BytesIO
import json

import frappe
from openpyxl import Workbook
from openpyxl.comments import Comment

from jifeng_erp import salary_history as H
from jifeng_erp.tool_archive_export import CONTENT_TYPE, _freeze, _fresh_snapshot


MAX_EXPORT_ROWS = 1_048_575
# Fixed, untranslated frontend labels are unique protocol IDs. Unknown old
# columns have no source mapping; notably default_amount is not old base pay.
COLUMNS = (
    ("说明", "description"), ("姓名", "employee_name"), ("部门", "department"),
    ("底薪", None), ("应出勤", None), ("实出勤", "actual_attendance"),
    ("出勤工资", "bucket:basic"), ("计时", "bucket:time"), ("计件", "bucket:piece"),
    ("合计", "core_total"), ("房补", "bucket:housing"), ("全勤", "bucket:full_attendance"),
    ("应发", "gross_pay"), ("水电", "bucket:utilities"), ("社保", "bucket:social_insurance"),
    ("罚款", "bucket:fine"), ("借款", "bucket:loan"), ("实发", "actual_paid"),
    ("开始", "start_date"), ("结束", "end_date"), ("类型", "type"),
    ("伙食", "bucket:meal"), ("交通", "bucket:transport"), ("夜宵", "bucket:night_snack"),
    ("加班", "bucket:overtime"), ("介绍费", "bucket:introduction"), ("考核", "bucket:assessment"),
    ("其他收入", "bucket:other_earning"), ("其他扣款", "bucket:other_deduction"),
    ("留存工资", "bucket:retained_pay"), ("住宿", "bucket:accommodation"),
)
COLUMN_FIELDS = dict(COLUMNS)
DATE_LABELS = {"period_start": "工资期开始", "period_end": "工资期结束",
    "period_overlap": "工资期间相交", "posting_date": "记账日"}
SIDE_LABELS = {"verified": "已完整分类并与原生金额核平", "unclassified": "分类或来源未完整核实",
    "unreconciled": "分类金额未与原生金额核平"}
UNKNOWN_LABELS = {"说明", "底薪", "应出勤", "实出勤", "合计", "实发", "类型"}


def _columns(value):
    try:
        if isinstance(value, str):
            value = json.loads(value)
        if (not isinstance(value, list) or not 1 <= len(value) <= len(COLUMNS)
                or any(not isinstance(label, str) or label not in COLUMN_FIELDS for label in value)
                or len(set(value)) != len(value)):
            raise ValueError
        return tuple(value)
    except (ValueError, TypeError, OverflowError):
        frappe.throw("工资历史导出列无效，请重新选择可见列")


def _deny():
    frappe.throw("没有完整读取及导出本次工资历史和所需来源的权限", frappe.PermissionError)


def _checked(context, columns):
    try:
        if not frappe.permissions.can_export(H.SLIP):
            raise H._Unavailable
        # This is the existing two complete 200-row scans, including child
        # table/field masks and rules. It never invokes the paginated endpoint.
        snapshot = H._stable_snapshot(context)
        if len(snapshot["rows"]) > MAX_EXPORT_ROWS:
            frappe.throw("工资历史超过Excel行数限制，请缩小筛选范围后完整导出")
        capture, checked_types = H._Capture(), {H.SLIP}
        selected_buckets = {COLUMN_FIELDS[label].split(":", 1)[1] for label in columns
            if (COLUMN_FIELDS[label] or "").startswith("bucket:")}

        def guard(doctype, name, fields):
            if doctype not in checked_types:
                if not frappe.permissions.can_export(doctype):
                    raise H._Unavailable
                checked_types.add(doctype)
            _, projected = capture.project(doctype, name, fields)
            # Native Data Exporter uses can_export(parent) then get_list's read
            # permissions, including System Manager's can_export semantics.
            # Linked/configuration doctypes retain read/select/field masks;
            # they are not exported contact or configuration tables.
            previous = snapshot["records"].get((doctype, name))
            if previous is None or previous != projected:
                raise H._Changed
            return projected

        for item in snapshot["rows"]:
            row = item["row"]
            for table, children in item["children"].items():
                if row["classification_state"][table] != "verified":
                    continue
                for child in children:
                    if child["do_not_include_in_total"] or not child["additional_salary"]:
                        continue
                    rule = H.rules.resolve_rule(snapshot["rules"], child["salary_component"], table, row["end_date"])
                    if not rule:
                        raise H._Changed
                    mode = rule.get("classification_mode") or "component"
                    contributes = {rule["bucket"]} if mode == "component" else {
                        "piece", "time", rule["extra_bucket"]}
                    if contributes & selected_buckets:
                        additional = guard("Additional Salary", child["additional_salary"], H.AS_FIELDS)
                        if additional["ref_doctype"] == "Jifeng Wage Entry":
                            guard("Jifeng Wage Entry", additional["ref_docname"], H.WAGE_FIELDS)
        return _freeze((snapshot, capture.records)), [item["row"] for item in snapshot["rows"]]
    except (H._Unavailable, frappe.PermissionError, frappe.DoesNotExistError,
            TypeError, InvalidOperation, OverflowError):
        _deny()
    except H._Changed:
        frappe.throw("工资历史来源或权限已改变，请重新查询后导出")


def _text(value):
    if value is None:
        return ""
    # The payroll projection already returns exact decimal strings. Accept no
    # float or implicit formatting here, including booleans disguised as zero.
    if not isinstance(value, str) or any(not (character in "\t\n\r" or " " <= character <= "\ud7ff"
            or "\ue000" <= character <= "\ufffd" or "\U00010000" <= character <= "\U0010ffff") for character in value):
        frappe.throw("工资历史包含无法完整导出的文本，请先核对原文")
    if len(value.encode("utf-16-le")) // 2 > 32767:
        frappe.throw("工资历史包含超过Excel单元格限制的文本，请先核对原文")
    return value


def _value(row, label):
    field = COLUMN_FIELDS[label]
    if label in UNKNOWN_LABELS:
        return None
    if field.startswith("bucket:"):
        bucket = field.split(":", 1)[1]
        side = "earnings" if bucket in H.rules.EARNING_BUCKETS else "deductions"
        if row["classification_state"][side] != "verified":
            return None
        value = row["buckets"].get(bucket)
    else:
        value = row[field]
    if field == "gross_pay" or field.startswith("bucket:"):
        try:
            finite = value is None or isinstance(value, str) and Decimal(value).is_finite()
        except (InvalidOperation, ValueError):
            finite = False
        if not finite:
            frappe.throw("工资历史包含无法完整导出的金额，请先核对原文")
    return value


def _cell(sheet, row, column, value):
    cell = sheet.cell(row, column)
    cell.value = _text(value)
    # Explicit string type makes = + - @, leading tabs and linebreaks literal.
    cell.data_type, cell.number_format = "s", "@"
    return cell


def _xlsx(context, columns, rows):
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = "工资历史"
        sheet.freeze_panes = "A2"
        for column, label in enumerate(columns, 1):
            cell = _cell(sheet, 1, column, label)
            if label in UNKNOWN_LABELS:
                note = "旧口径尚未核实，整列留空；未折算默认额不等于底薪，计薪天数不等于实际出勤，净应发不等于实际支付。"
            elif label == "应发":
                note = "原生工资单已保存的总收入（gross_pay），未扣员工扣款；不表示已付款。金额为十进制文本，币种见各金额单元格批注。"
            elif COLUMN_FIELDS[label].startswith("bucket:"):
                note = "本工资期间正式分类的已保存金额，必须全侧分类且与原生金额核平；未知留空，合法零保留。借款仅表示明确分类的工资扣款。"
            else:
                note = "工资单已保存的历史字段，不使用员工当前档案替代。"
            cell.comment = Comment(note, "集沣工资历史")
        for index, row in enumerate(rows, 2):
            for column, label in enumerate(columns, 1):
                cell = _cell(sheet, index, column, _value(row, label))
                field = COLUMN_FIELDS[label] or ""
                if field == "gross_pay" or field.startswith("bucket:"):
                    note = "币种：" + _text(row["currency"]) + "。金额保存为精确十进制文本，无跨币种合计。"
                    if field.startswith("bucket:"):
                        bucket = field.split(":", 1)[1]
                        side = "earnings" if bucket in H.rules.EARNING_BUCKETS else "deductions"
                        note += " 分类：" + SIDE_LABELS[row["classification_state"][side]] + "；空白不等于零。"
                    cell.comment = Comment(note, "集沣工资历史")
        sheet.auto_filter.ref = f"A1:{sheet.cell(max(1, len(rows) + 1), len(columns)).coordinate}"
        notes = workbook.create_sheet("口径说明")
        explanations = (
            ("公司", context["company"]), ("日期筛选", "启用" if context["use_date"] else "未启用"),
            ("日期口径", DATE_LABELS[context["date_basis"]]), ("开始日期", context["from_date"]),
            ("结束日期", context["to_date"]), ("关键词", context["search"]), ("工资单数量", str(len(rows))),
            ("数据范围", "当前筛选下已保存、已提交且原生可读的完整工资单；每行一单，不重算工资。"),
            ("应发", "原生总收入 gross_pay；净应发 net_pay、取整净应发 rounded_total 不作为实发，也未加入旧列。"),
            ("未知旧字段", "说明、底薪、应出勤、实出勤、合计、实发、类型留空；不自动填零或用默认额、计薪天数、净应发代替。"),
            ("分类金额", "只展示整个收入或扣款侧已完整分类并核平的金额；未知整侧留空，零是已核实的零。"),
            ("金额与币种", "金额写为精确十进制文本，币种见金额单元格批注；不进行跨币种合计。"),
            ("完整性", "按原生工资单导出权限和读取范围复核；分类未知留空，已输出分类金额所需金融来源的导出权限不足时整份拒绝；不删行、不截断。"),
        )
        for index, pair in enumerate(explanations, 1):
            for column, value in enumerate(pair, 1):
                _cell(notes, index, column, value)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


@frappe.whitelist(methods=["GET"])
def export_salary_history(context, columns):
    context, columns = H._context(context), _columns(columns)
    permissions = H._permission_state()
    checked, rows = _checked(context, columns)
    content = _xlsx(context, columns, rows)
    # Reuse the verified primary-connection, read-only transaction and cleanup
    # helper. A successful complete comparison precedes publishing any bytes.
    with _fresh_snapshot():
        if permissions != H._permission_state():
            frappe.throw("工资历史字段或权限已改变，请重新查询后导出", frappe.PermissionError)
        current, _ = _checked(context, columns)
        if current != checked:
            frappe.throw("工资历史来源或可见范围已改变，请重新查询后导出")
        if permissions != H._permission_state():
            frappe.throw("工资历史字段或权限已改变，请重新查询后导出", frappe.PermissionError)
    frappe.response.update({"type": "download", "content_type": CONTENT_TYPE, "filename": "工资历史.xlsx",
        "display_content_as": "attachment", "filecontent": content})
