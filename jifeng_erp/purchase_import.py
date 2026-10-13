"""Preview spreadsheet rows for native Purchase Orders without saving documents."""

import base64
import binascii
import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
import json
import zipfile

import frappe
from frappe.model import get_permitted_fields
from openpyxl import Workbook, load_workbook


ORDER = "Purchase Order"
CHILD = "Purchase Order Item"
PROFILE = "Jifeng Inventory Profile"
MAX_BYTES, MAX_ROWS = 1024 * 1024, 500
HEADERS = {
    "商品编码/条码": "code", "商品编码": "code", "物料编码": "code", "条码": "code",
    "编号": "code", "货号": "code", "item_code": "code", "barcode": "code",
    "数量": "qty", "采购数量": "qty", "qty": "qty",
    "单价": "rate", "采购单价": "rate", "rate": "rate",
    "单位": "uom", "uom": "uom", "交货日期": "schedule_date", "schedule_date": "schedule_date",
    "仓库": "warehouse", "warehouse": "warehouse",
}
TEMPLATE = ["商品编码/条码", "数量", "单价", "单位", "交货日期", "仓库"]
CONTEXT_FIELDS = {
    "company", "supplier", "transaction_date", "schedule_date", "currency", "conversion_rate",
    "buying_price_list", "price_list_currency", "plc_conversion_rate", "set_warehouse",
    "ignore_pricing_rule", "tax_category", "project", "cost_center",
}
ROW_FIELDS = {
    "item_code", "item_name", "description", "qty", "rate", "uom", "stock_uom",
    "conversion_factor", "stock_qty", "warehouse", "schedule_date", "amount",
    "price_list_rate", "discount_percentage", "discount_amount", "margin_type", "margin_rate_or_amount",
    "item_tax_template", "item_tax_rate", "expense_account", "cost_center", "project",
    "weight_per_unit", "weight_uom", "total_weight",
}


def _columns(doctype, permission="read", parenttype=None):
    fields = set(get_permitted_fields(doctype, parenttype=parenttype,
                                     permission_type=permission, ignore_virtual=True))
    if permission == "read" and frappe.session.user != "Administrator":
        meta = frappe.get_meta(doctype)
        levels = meta.get_permlevel_access(permission_type="mask", parenttype=parenttype, user=frappe.session.user)
        fields -= {df.fieldname for df in meta.fields if df.get("mask") and (df.permlevel or 0) not in levels}
    return fields


def _table_access(doctype, fieldname, childtype, write=False):
    # Native get_permitted_fields excludes Table fields from its scalar projection.
    meta = frappe.get_meta(doctype)
    table = meta.get_field(fieldname)
    permissions = ("read", "write") if write else ("read",)
    return bool(table and table.fieldtype == "Table" and table.options == childtype and (not write or not table.read_only)
                and all((table.permlevel or 0) in meta.get_permlevel_access(permission_type=permission)
                        for permission in permissions)
                and (frappe.session.user == "Administrator" or not table.get("mask")
                     or (table.permlevel or 0) in meta.get_permlevel_access(permission_type="mask")))


def _scope(context):
    if frappe.session.user == "Guest":
        frappe.throw("请登录后导入采购明细", frappe.PermissionError)
    if isinstance(context, str):
        context = json.loads(context)
    if not isinstance(context, dict) or set(context) - CONTEXT_FIELDS - {"name", "modified"}:
        frappe.throw("采购单上下文无效")
    for key in ("company", "supplier", "transaction_date"):
        if not isinstance(context.get(key), str) or not context[key].strip():
            frappe.throw("请先选择公司、供应商和采购日期")
    for doctype, name in (("Company", context["company"]), ("Supplier", context["supplier"])):
        frappe.get_doc(doctype, name).check_permission("select")
    if context.get("name"):
        document = frappe.get_doc(ORDER, context["name"])
        document.check_permission("read")
        document.check_permission("write")
        if document.docstatus != 0 or str(document.modified) != str(context.get("modified")):
            frappe.throw("采购单已变化，请刷新后重新导入")
        if document.company != context["company"]:
            frappe.throw("请先保存采购单的公司变更，再导入")
        if document.get("is_subcontracted") or document.get("is_old_subcontracting_flow") or document.get("has_unit_price_items"):
            frappe.throw("批量导入仅用于普通数量采购单")
    else:
        document = frappe.get_doc({"doctype": ORDER, **{key: context[key] for key in CONTEXT_FIELDS if key in context}})
        document.check_permission("create")
    required = {"item_code", "item_name", "qty", "rate", "uom", "stock_uom", "conversion_factor", "schedule_date", "warehouse"}
    readable = _columns(CHILD, parenttype=ORDER)
    writable = _columns(CHILD, "write", ORDER)
    table_editable = _table_access(ORDER, "items", CHILD, write=True)
    if (not required.issubset(readable) or not (required - {"item_name", "stock_uom", "conversion_factor"}).issubset(writable)
            or not table_editable or not {"company", "supplier"}.issubset(_columns(ORDER))):
        frappe.throw("没有导入采购明细字段的权限", frappe.PermissionError)
    return context, readable


def _raw_file(filename, content_base64):
    if not isinstance(content_base64, str) or len(content_base64) > ((MAX_BYTES + 2) // 3) * 4:
        frappe.throw("表格文件不能超过1MiB")
    try:
        raw = base64.b64decode(content_base64, validate=True)
    except (ValueError, binascii.Error):
        frappe.throw("文件内容无效")
    if not raw or len(raw) > MAX_BYTES:
        frappe.throw("请选择非空且不超过1MiB的表格文件")
    if not isinstance(filename, str) or not filename.lower().endswith((".csv", ".xlsx")):
        frappe.throw("请选择.xlsx或.csv文件")
    return raw


def _xlsx(raw):
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            members = archive.infolist()
            if (len(members) > 128 or sum(member.file_size for member in members) > 16 * MAX_BYTES
                    or len({member.filename for member in members}) != len(members)
                    or any(member.flag_bits & 1 or member.file_size > max(member.compress_size, 1) * 200 for member in members)):
                frappe.throw("表格解压大小超过限额")
            if any(token in member.filename.lower() for member in members for token in ("vba", "externallinks/", "embeddings/")):
                frappe.throw("表格不能包含宏或外部链接")
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) != 1:
                frappe.throw("请只保留一个采购明细工作表")
            sheet = book.worksheets[0]
            if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > 6:
                frappe.throw("每次最多500行、6列，请删除多余行列后重试")
            rows = []
            for cells in sheet.iter_rows():
                if any(cell.data_type in ("f", "e") for cell in cells):
                    frappe.throw("表格含公式或错误单元格，请粘贴为值后导入")
                rows.append([cell.value for cell in cells])
            return rows
        finally:
            book.close()
    except (zipfile.BadZipFile, ValueError, KeyError, OSError):
        frappe.throw("无法读取.xlsx文件，请重新另存为标准Excel文件")


def _text_rows(text):
    if not isinstance(text, str) or len(text.encode("utf-8")) > MAX_BYTES:
        frappe.throw("粘贴内容不能超过1MiB")
    text = text.lstrip("\ufeff")
    first = text.splitlines()[0] if text.splitlines() else ""
    delimiter = "\t" if "\t" in first else ";" if ";" in first and "," not in first else ","
    try:
        reader = csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True)
        rows = []
        for row in reader:
            rows.append(row)
            if len(rows) > MAX_ROWS + 1:
                frappe.throw("每次最多导入500行，请分批导入")
        return rows
    except csv.Error:
        frappe.throw("表格文本格式无效，请从Excel重新复制或另存为CSV")


def _data(filename=None, content_base64=None, text=None):
    if text is not None:
        if filename or content_base64:
            frappe.throw("请选择文件或粘贴内容中的一种")
        rows = _text_rows(text)
    else:
        raw = _raw_file(filename, content_base64)
        if filename.lower().endswith(".xlsx"):
            rows = _xlsx(raw)
        else:
            try:
                decoded = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                try:
                    decoded = raw.decode("gb18030")
                except UnicodeDecodeError:
                    frappe.throw("CSV编码无效，请使用UTF-8或中文CSV")
            rows = _text_rows(decoded)
    if len(rows) < 2:
        frappe.throw("表格需要列名和至少一行采购明细")
    headers = [HEADERS.get(str(value or "").strip().lower()) for value in rows[0]]
    if None in headers or len(headers) != len(set(headers)) or not {"code", "qty", "rate"}.issubset(headers):
        frappe.throw("列名须包含商品编码/条码、数量、单价；可选单位、交货日期、仓库")
    if len(rows) > MAX_ROWS + 1:
        frappe.throw("每次最多导入500行，请分批导入")
    return [(index, dict(zip(headers, row, strict=True))) if len(row) == len(headers) else
            (index, {"_error": "列数与表头不一致"}) for index, row in enumerate(rows[1:], 2)
            if any(value is not None and str(value).strip() for value in row)]


def _number(value, label, positive=False):
    try:
        number = Decimal(str(value).strip())
        if not number.is_finite() or number < 0 or (positive and number == 0) or number > Decimal("1000000000"):
            raise InvalidOperation
    except (InvalidOperation, ValueError):
        frappe.throw(label + ("须为大于0的有效数字" if positive else "须为不小于0的有效数字"))
    return number


def _date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    try:
        return date.fromisoformat(str(value).strip().replace("/", "-")).isoformat()
    except ValueError:
        frappe.throw("交货日期请填写YYYY-MM-DD")


def _resolve(code, company):
    if not isinstance(code, str) or not code.strip() or len(code.strip()) > 140:
        frappe.throw("商品编码/条码请按文本填写，保留前导零")
    code = code.strip()
    matches = set()
    if not {"name", "item_name", "stock_uom", "disabled", "is_purchase_item"}.issubset(_columns("Item")):
        frappe.throw("没有读取商品采购资料的权限", frappe.PermissionError)
    matches.update(row.name for row in frappe.get_list("Item", filters={"name": code}, fields=["name"], limit_page_length=2))
    if _table_access("Item", "barcodes", "Item Barcode") and "barcode" in _columns("Item Barcode", parenttype="Item"):
        matches.update(row.name for row in frappe.get_list("Item", filters=[["Item Barcode", "barcode", "=", code]],
                       fields=["name"], distinct=True, limit_page_length=2))
    if frappe.has_permission(PROFILE, "read") and {"company", "inactive", "item_code", "legacy_number", "legacy_code"}.issubset(_columns(PROFILE)):
        matches.update(row.item_code for row in frappe.get_list(PROFILE,
                       filters={"company": company, "inactive": 0}, or_filters={"legacy_number": code, "legacy_code": code},
                       fields=["item_code"], limit_page_length=2))
    if len(matches) != 1:
        frappe.throw("未找到唯一可用商品，请核对编码、条码或当前公司档案编号")
    item_code = matches.pop()
    item = frappe.get_doc("Item", item_code)
    item.check_permission("read")
    if item.disabled or not item.is_purchase_item:
        frappe.throw("该商品已停用或不允许采购")
    return item_code


@frappe.whitelist(methods=["POST"])
def preview_import(context, filename=None, content_base64=None, text=None):
    context, readable = _scope(context)
    source = _data(filename, content_base64, text)
    result, resolved = [], {}
    for index, values in source:
        entry = {"row_number": index, "code": str(values.get("code") or ""), "item": None, "error": ""}
        previous_mute = frappe.flags.mute_messages
        frappe.flags.mute_messages = True
        try:
            if values.get("_error"):
                frappe.throw(values["_error"])
            qty, rate = _number(values.get("qty"), "数量", True), _number(values.get("rate"), "单价")
            code = values.get("code")
            if code not in resolved:
                resolved[code] = _resolve(code, context["company"])
            schedule = _date(values.get("schedule_date") or context.get("schedule_date") or context["transaction_date"])
            if schedule < _date(context["transaction_date"]):
                frappe.throw("交货日期不能早于采购日期")
            item_values = {"item_code": resolved[code], "qty": float(qty), "rate": float(rate), "schedule_date": schedule}
            for field, doctype in (("uom", "UOM"), ("warehouse", "Warehouse")):
                value = values.get(field) or (context.get("set_warehouse") if field == "warehouse" else None)
                if value:
                    linked = frappe.get_doc(doctype, str(value).strip())
                    linked.check_permission("select")
                    if field == "warehouse" and (linked.company != context["company"] or linked.is_group or linked.disabled):
                        frappe.throw("仓库须为当前公司的在用末级仓库")
                    item_values[field] = linked.name
            document = frappe.get_doc({"doctype": ORDER, **{key: context[key] for key in CONTEXT_FIELDS if key in context},
                                       "items": [item_values]})
            document.set_missing_item_details(for_validate=True)
            item = document.items[0]
            if len(document.items) != 1 or Decimal(str(item.rate)) != rate or Decimal(str(item.qty)) != qty:
                frappe.throw("原生定价规则改变了数量或单价，请从采购单核对后录入")
            item.amount = float(qty * rate)
            for field in ROW_FIELDS - readable:
                if item.get(field) not in (None, "", 0, "{}"):
                    frappe.throw("没有读取该商品采购明细的完整字段权限", frappe.PermissionError)
            entry["item"] = {key: item.get(key) for key in sorted(ROW_FIELDS & readable) if item.get(key) is not None}
        except (frappe.ValidationError, frappe.PermissionError, frappe.DoesNotExistError) as error:
            entry["error"] = "没有读取或选择该商品及关联资料的权限" if isinstance(error, (frappe.PermissionError, frappe.DoesNotExistError)) else str(error)
        finally:
            frappe.flags.mute_messages = previous_mute
        result.append(entry)
    errors = sum(bool(row["error"]) for row in result)
    return {"rows": result, "row_count": len(result), "error_count": errors, "can_apply": bool(result) and errors == 0,
            "company": context["company"], "supplier": context["supplier"]}


@frappe.whitelist(methods=["GET"])
def download_template():
    if frappe.session.user == "Guest" or not frappe.has_permission(ORDER, "create"):
        frappe.throw("没有采购单新建权限", frappe.PermissionError)
    book = Workbook()
    try:
        sheet = book.active
        sheet.title = "采购明细"
        sheet.append(TEMPLATE)
        sheet.freeze_panes = "A2"
        for index in range(2, MAX_ROWS + 2):
            for column in (1, 4, 5, 6):
                sheet.cell(index, column).number_format = "@"
        for letter in "ABCDEF":
            sheet.column_dimensions[letter].width = 24
        output = BytesIO()
        book.save(output)
        frappe.response.update({"type": "download", "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                "filename": "采购明细导入模板.xlsx", "display_content_as": "attachment", "filecontent": output.getvalue()})
    finally:
        book.close()
