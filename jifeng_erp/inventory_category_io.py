"""XLSX category exchange. Parsing creates drafts, never saved dictionary rows."""

from io import BytesIO
from uuid import uuid4

import frappe
from openpyxl import Workbook, load_workbook

from jifeng_erp import inventory_category as C
from jifeng_erp.inventory_archive import _columns, _normalized, _permission


MAX_IMPORT_ROWS = 100
MAX_EXPORT_ROWS = 10000
MARKER = "集沣库存类别交换 v1"
CONTEXT_SHEET = "导入说明"
IDENTITY = (("reference", "引用"), ("name", "记录"), ("version", "版本"))
PARENT = (("parent_reference", "大类引用"), ("parent_version", "大类版本"))
LABELS = {"legacy_code": "代码", "owner_name": "负责", "region": "区域", "pinyin": "简拼", "inactive": "停用",
          "material": "材质", "display_uom": "单位", "conversion_uom_text": "换算单位", "conversion_value": "换算值", "notes": "备注"}


def _fields(record_type):
    fields = C.MAJOR_FIELDS if record_type == "major" else C.MINOR_FIELDS
    return [(field, "大类" if record_type == "major" else "小类") if field == "value" else (field, LABELS[field]) for field in fields]


def _columns_for(record_type):
    return [*IDENTITY, *(PARENT if record_type == "minor" else ()),
            *[(field, label) for field, label in _fields(record_type) if field in _columns(C.DICTIONARY, "read")]]


def _text(sheet, row, column, value):
    cell = sheet.cell(row, column, "" if value is None else str(value))
    cell.data_type, cell.number_format = "s", "@"


def _workbook(kind, company):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for record_type, title in (("major", "大类"), ("minor", "小类")):
        sheet = workbook.create_sheet(title)
        for index, (_, label) in enumerate(_columns_for(record_type), 1):
            _text(sheet, 1, index, label)
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["B"].hidden = sheet.column_dimensions["C"].hidden = True
        if record_type == "minor":
            sheet.column_dimensions["E"].hidden = True
    notes = workbook.create_sheet(CONTEXT_SHEET)
    for text in (MARKER, company, kind, "引用在文件内唯一；小类的大类引用必须指向大类页中的引用。",
                 "新增行的记录、版本及大类版本留空；已有行保留导出时的记录与版本。",
                 "编号、代码、名称和换算值按文本填写以保留前导零和小数。",
                 "导入只生成待保存修改，请核对两张表后点保存；不按名称猜测已有记录。",
                 "每次导入最多100条数据；导出超过100条时，请分段整理后导入。"):
        notes.append([text])
    return workbook


def _download(workbook, filename):
    try:
        output = BytesIO()
        workbook.save(output)
        frappe.response.update({"type": "download", "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                "filename": filename, "display_content_as": "attachment", "filecontent": output.getvalue()})
    finally:
        workbook.close()


@frappe.whitelist(methods=["GET"])
def download_template(kind, company):
    C._scope(kind, company)
    if not frappe.has_permission(C.DICTIONARY, "create"):
        frappe.throw("没有新增类别的权限", frappe.PermissionError)
    _download(_workbook(kind, company), kind + "类别导入模板.xlsx")


@frappe.whitelist(methods=["GET"])
def export_categories(kind, company):
    C._scope(kind, company)
    if not frappe.has_permission(C.DICTIONARY, "export"):
        frappe.throw("没有导出类别的权限", frappe.PermissionError)
    names = frappe.get_list(C.DICTIONARY, filters={"company": company, "archive_kind": kind,
                            "dictionary_type": ["in", ["大类", "小类"]]}, pluck="name",
                            order_by="creation asc, name asc", limit_page_length=MAX_EXPORT_ROWS + 1)
    if len(names) > MAX_EXPORT_ROWS:
        frappe.throw("单次导出不能超过10000条类别")
    documents = [C._record(name, kind, company) for name in names]
    parents = {document.name: document for document in documents if document.dictionary_type == "大类"}
    references = {name: "M-" + str(index).zfill(4) for index, name in enumerate(parents, 1)}
    workbook = _workbook(kind, company)
    try:
        positions, child_index = {"major": 2, "minor": 2}, 0
        for document in documents:
            _permission(document, "export")
            record_type = "major" if document.dictionary_type == "大类" else "minor"
            parent = None
            if record_type == "minor" and document.parent_dictionary:
                parent = parents.get(document.parent_dictionary)
                if not parent:
                    frappe.throw("类别关联或读取权限已变化，请重新查询后导出")
                _permission(parent, "read")
            child_index += record_type == "minor"
            values = {field: C._value(document.get(field)) for field, _ in _fields(record_type)}
            values.update(reference=references[document.name] if record_type == "major" else "C-" + str(child_index).zfill(4),
                          name=document.name, version=C._value(document.modified),
                          parent_reference=references[parent.name] if parent else "",
                          parent_version=C._value(parent.modified) if parent else "")
            sheet = workbook["大类" if record_type == "major" else "小类"]
            for index, (field, _) in enumerate(_columns_for(record_type), 1):
                _text(sheet, positions[record_type], index, values.get(field))
            positions[record_type] += 1
        _download(workbook, kind + "类别.xlsx")
    finally:
        workbook.close()


def _sheet_rows(sheet, record_type):
    columns = _columns_for(record_type)
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if headers != [label for _, label in columns]:
        frappe.throw("类别文件列名或读取权限不匹配，请使用当前公司模板")
    rows = []
    for line, cells in enumerate(sheet.iter_rows(min_row=2), 2):
        values = [cell.value for cell in cells]
        if all(value in (None, "") for value in values):
            continue
        if len(rows) >= MAX_IMPORT_ROWS:
            frappe.throw("每次类别导入最多100条数据")
        row = dict(zip((field for field, _ in columns), values, strict=True))
        for field in ("reference", "name", "version", "parent_reference", "parent_version"):
            value = row.get(field)
            if value not in (None, "") and (not isinstance(value, str) or value != value.strip() or len(value) > 140):
                frappe.throw("类别引用及身份须为准确文本，第" + str(line) + "行")
        if not row["reference"]:
            frappe.throw("请填写类别导入引用，第" + str(line) + "行")
        rows.append(row)
    return rows


def _values(row, record_type, document):
    values = {field: row.get(field) for field, _ in _fields(record_type) if field in _columns(C.DICTIONARY, "read")}
    if values.get("inactive") in (None, ""):
        values["inactive"] = 0
    if values.get("conversion_value") not in (None, "") and not isinstance(values["conversion_value"], str):
        values["conversion_value"] = str(values["conversion_value"])
    values = _normalized(C.DICTIONARY, values, set(C.MAJOR_FIELDS if record_type == "major" else C.MINOR_FIELDS))
    if document:
        return {field: value for field, value in values.items()
                if str("" if value is None else value) != str("" if document.get(field) is None else document.get(field))}
    return {field: value for field, value in values.items() if value is not None}


@frappe.whitelist(methods=["POST"])
def parse_import(kind, company, filename, content_base64):
    C._scope(kind, company)
    from jifeng_erp.inventory_archive_import import _xlsx_bytes

    raw = _xlsx_bytes(filename, content_base64)
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    try:
        if set(workbook.sheetnames) != {"大类", "小类", CONTEXT_SHEET}:
            frappe.throw("类别导入文件须包含大类、小类和导入说明三个页")
        notes = workbook[CONTEXT_SHEET]
        if [notes.cell(row, 1).value for row in (1, 2, 3)] != [MARKER, company, kind]:
            frappe.throw("类别文件与当前公司或档案种类不匹配")
        majors, minors = _sheet_rows(workbook["大类"], "major"), _sheet_rows(workbook["小类"], "minor")
        if not 1 <= len(majors) + len(minors) <= MAX_IMPORT_ROWS:
            frappe.throw("每次类别导入须有1至100条数据")
        references, identities, rows, sources = {}, set(), [], {}
        for record_type, input_rows in (("major", majors), ("minor", minors)):
            for incoming in input_rows:
                reference, name = incoming["reference"], incoming.get("name")
                if reference in references or name and name in identities:
                    frappe.throw("类别文件包含重复引用或记录")
                document = C._record(name, kind, company, "大类" if record_type == "major" else "小类") if name else None
                if document:
                    if not C.same_version(document.modified, incoming.get("version")):
                        frappe.throw("导入来源类别已有变化，请重新导出", frappe.TimestampMismatchError)
                    identities.add(name)
                    sources[name] = C._pack(document)
                elif incoming.get("version"):
                    frappe.throw("新增类别不能携带已有版本")
                row = {"row_id": str(uuid4()), "record_type": record_type, "operation": "upsert"}
                if document:
                    row.update(name=name, version=incoming["version"])
                row["values"] = _values(incoming, record_type, document)
                if record_type == "minor":
                    parent_entry = references.get(incoming.get("parent_reference"))
                    if not parent_entry or parent_entry[0]["record_type"] != "major":
                        frappe.throw("小类须明确引用文件大类页中的类别")
                    parent_row, parent_document = parent_entry
                    if parent_document:
                        if incoming.get("parent_version") not in (None, "", parent_row["version"]):
                            frappe.throw("小类的大类版本与文件不一致")
                        row.update(category=parent_document.name, category_version=parent_row["version"])
                    else:
                        row["parent_row_id"] = parent_row["row_id"]
                    if document and not document.parent_dictionary:
                        if row["values"]:
                            frappe.throw("未关联小类须先关联保存，再导入内容修改")
                        row.pop("values")
                        row["operation"] = "associate"
                references[reference] = (row, document)
                if not document or row.get("values") or row["operation"] == "associate":
                    rows.append(row)
        if rows:
            payload = {"kind": kind, "company": company, "request_id": str(uuid4()), "rows": rows}
            C._batch(payload)
            C._prepare(payload)
        return {"kind": kind, "company": company, "rows": rows, "sources": sources,
                "input_count": len(majors) + len(minors)}
    finally:
        workbook.close()
