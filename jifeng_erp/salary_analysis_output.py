"""Print/export the exact, authorized payroll settlement analysis in memory."""

from html import escape
from io import BytesIO
import re

import frappe
from openpyxl import Workbook

from jifeng_erp import salary_analysis as A
from jifeng_erp import salary_history as H
from jifeng_erp.archive_versions import same_version
from jifeng_erp.salary_history_export import DATE_LABELS, _text
from jifeng_erp.tool_archive_export import CONTENT_TYPE, _freeze, _fresh_snapshot


TOTALS = (("总收入", "gross_pay"), ("总扣款", "total_deduction"),
          ("净应发（未证明付款）", "net_pay"), ("舍入净应发", "rounded_total"))
CHARTS = (("部门收入占比", "department"), ("分组收入占比", "group"),
          ("工序计薪金额占比", "operation"))
SLIP_COLUMNS = (("工资条", "name"), ("员工编号", "employee"), ("姓名", "employee_name"),
                ("部门", "department"), ("币种", "currency"), ("开始", "start_date"), ("结束", "end_date"),
                *TOTALS, ("工资条版本", "version"))
SOURCE_COLUMNS = (("工资条", "salary_slip"), ("员工编号", "employee"), ("姓名", "employee_name"),
                  ("工序", "operation"), ("工序金额", "amount"), ("计薪记录", "wage_entry"),
                  ("附加薪资", "additional_salary"), ("工资条版本", "salary_slip_version"),
                  ("计薪记录版本", "wage_entry_version"), ("附加薪资版本", "additional_salary_version"))


def _deny():
    frappe.throw("没有完整打印或导出本次工资分析及其准确来源的权限", frappe.PermissionError)


def _checked(context, version, operation):
    result = A.query_salary_analysis(context)
    if result.get("requires_currency"):
        frappe.throw("请选择一个币种重新查询后打印或导出，不能跨币种合计")
    if result["result_version"] != version:
        frappe.throw("工资分析来源或金额已变化，请重新查询后打印或导出")
    references = {(H.SLIP, slip["name"]): slip["version"] for slip in result["slips"]}
    for source in result["sources"]:
        for doctype, name_field, version_field in (
            ("Jifeng Wage Entry", "wage_entry", "wage_entry_version"),
            ("Additional Salary", "additional_salary", "additional_salary_version"),
        ):
            key, saved_version = (doctype, source[name_field]), source[version_field]
            if key in references and references[key] != saved_version:
                frappe.throw("工资分析来源版本不一致，请重新查询")
            references[key] = saved_version
    checked_types, documents = {H.SLIP}, {}
    try:
        # Preserve native can_export semantics, including System Manager. A
        # readable report must not silently grant an export or print right.
        if operation == "export" and not frappe.permissions.can_export(H.SLIP):
            _deny()
        if operation == "print" and not frappe.has_permission(H.SLIP, "print"):
            _deny()
        for (doctype, name), saved_version in sorted(references.items()):
            if operation == "export" and doctype not in checked_types:
                if not frappe.permissions.can_export(doctype):
                    _deny()
                checked_types.add(doctype)
            document = frappe.get_doc(doctype, name)
            document.check_permission("read")
            if operation == "print":
                document.check_permission("print")
            if document.docstatus != 1 or not same_version(document.modified, saved_version):
                frappe.throw("工资分析来源已改变，请重新查询后打印或导出")
            documents[(doctype, name)] = _freeze(document.as_dict())
        return result, _freeze(documents)
    except (frappe.PermissionError, frappe.DoesNotExistError):
        _deny()


def _summary(result):
    context = result["context"]
    return [("公司", context["company"]), ("日期口径", DATE_LABELS[context["date_basis"]]),
            ("开始日期", context["from_date"]), ("结束日期", context["to_date"]),
            ("员工筛选", context["employee"]), ("关键词", context["search"]),
            ("币种", result["currency"]), ("已提交工资条数量", str(result["salary_slips"])),
            *[(label, result["settlements"][field]) for label, field in TOTALS]]


def _tables(result):
    for title, field in CHARTS:
        yield title, ("名称", "金额", "占比（%）"), [
            (row["name"], row["amount"], row["share_percent"]) for row in result[field]]
    yield "工资条来源", tuple(label for label, _ in SLIP_COLUMNS), [
        tuple(row[field] for _, field in SLIP_COLUMNS) for row in result["slips"]]
    yield "工序计薪来源", tuple(label for label, _ in SOURCE_COLUMNS), [
        tuple(row[field] for _, field in SOURCE_COLUMNS) for row in result["sources"]]


def _xlsx(result):
    workbook = Workbook()
    try:
        summary = workbook.active
        summary.title = "工资结算"
        tables = [(summary.title, ("项目", "值"), _summary(result)), *_tables(result),
                  ("口径说明", ("说明",), [(note,) for note in result["notes"]] + [
                      ("金额保存为精确十进制文本；净应发、舍入净应发均不代表实际付款。",),
                      ("来源仅含本次筛选已提交工资条，取消工资条不计入；未知分组不按当前人员档案补推。",)])]
        for index, (title, headings, rows) in enumerate(tables):
            if len(rows) > 1_048_575:
                frappe.throw("工资分析超过Excel行数限制，请缩小筛选范围后完整导出")
            sheet = summary if index == 0 else workbook.create_sheet(title)
            sheet.freeze_panes = "A2"
            for row_number, values in enumerate([headings, *rows], 1):
                for column, value in enumerate(values, 1):
                    cell = sheet.cell(row_number, column)
                    cell.value = _text(value)
                    cell.data_type, cell.number_format = "s", "@"
            sheet.auto_filter.ref = f"A1:{sheet.cell(len(rows) + 1, len(headings)).coordinate}"
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


def _table(headings, rows):
    head = "".join("<th>" + escape(_text(value)) + "</th>" for value in headings)
    body = "".join("<tr>" + "".join("<td>" + escape(_text(value)) + "</td>" for value in row) + "</tr>"
                   for row in rows)
    if not body:
        body = '<tr><td colspan="' + str(len(headings)) + '">当前筛选无来源记录</td></tr>'
    return "<table><thead><tr>" + head + "</tr></thead><tbody>" + body + "</tbody></table>"


def _html(result):
    sections = ["<h1>工资历史分析</h1>", _table(("项目", "值"), _summary(result))]
    for title, headings, rows in _tables(result):
        sections.extend(["<h2>" + escape(title) + "</h2>", _table(headings, rows)])
    sections.append("<h2>口径说明</h2>")
    sections.extend("<p>" + escape(_text(note)) + "</p>" for note in result["notes"])
    return ('<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">'
            '<title>工资历史分析</title><style>'
            '@page{size:A4 landscape;margin:10mm}body{font-family:sans-serif;font-size:10pt;color:#111}'
            'table{border-collapse:collapse;width:100%;table-layout:fixed;margin-bottom:1em}'
            'th,td{border:1px solid #aaa;padding:4px;overflow-wrap:anywhere;white-space:pre-wrap}'
            'thead{display:table-header-group}tr{break-inside:avoid}h2{break-after:avoid;font-size:13pt}'
            'h1{font-size:17pt}p{margin:.5em 0}</style></head><body>' + "".join(sections) + '</body></html>')


def _output(context, result_version, operation):
    if (frappe.session.user == "Guest" or not isinstance(result_version, str)
            or not re.fullmatch(r"[0-9a-f]{64}", result_version)):
        _deny()
    permissions = H._permission_state()
    result, documents = _checked(context, result_version, operation)
    content = _xlsx(result) if operation == "export" else _html(result)
    # Publish no bytes until a fresh primary read-only snapshot verifies the
    # exact output and source rights; reuse the installed connection helper.
    with _fresh_snapshot():
        if permissions != H._permission_state():
            _deny()
        current, current_documents = _checked(context, result_version, operation)
        if current != result or current_documents != documents:
            frappe.throw("工资分析来源已变化，请重新查询后打印或导出")
        if permissions != H._permission_state():
            _deny()
    frappe.response.update({"type": "download", "content_type": CONTENT_TYPE if operation == "export" else "text/html",
        "filename": "工资历史分析.xlsx" if operation == "export" else "工资历史分析.html",
        "display_content_as": "attachment", "filecontent": content})


@frappe.whitelist(methods=["GET"])
def export_salary_analysis(context, result_version):
    _output(context, result_version, "export")


@frappe.whitelist(methods=["GET"])
def print_salary_analysis(context, result_version):
    _output(context, result_version, "print")
