"""Complete factory statements over the existing native-source reconciliation."""

from datetime import date, datetime
from decimal import Decimal
from html import escape
from io import BytesIO
import json
import math
import re

import frappe
from openpyxl import Workbook

from jifeng_erp import subcontracting_reconciliation as R
from jifeng_erp.tool_archive_export import CONTENT_TYPE, _fresh_snapshot


CONTEXT_FIELDS = {"company", "supplier", "from_date", "to_date", "business_kind", "flow_order", "subcontracting_order"}
SECTIONS = (
    ("委外进度", "progress", (("subcontracting_order", "委外订单"), ("supplier_name", "供应商"), ("flow_order", "流程单"),
        ("item_code", "商品"), ("stock_uom", "单位"), ("planned_qty", "计划量"), ("accepted_qty", "净接收"),
        ("rejected_qty", "净拒收"), ("loss_qty", "确认损耗"), ("pending_qty", "待收"))),
    ("收回与库存价值", "receipts", (("receipt", "收回/退回单"), ("entry_type", "业务类型"),
        ("subcontracting_order", "委外订单"), ("flow_order", "流程单"), ("item_code", "商品"), ("stock_uom", "单位"),
        ("accepted_qty", "接收变动"), ("rejected_qty", "拒收变动"), ("loss_qty", "损耗变动"),
        ("inventory_qty_change", "库存数量变动"), ("inventory_value_change", "库存价值变动"), ("inventory_value_status", "估价状态"))),
    ("委外原料发退", "materials", (("stock_entry", "发退料凭证"), ("entry_type", "业务类型"),
        ("subcontracting_order", "委外订单"), ("flow_order", "流程单"), ("rm_item_code", "原料"), ("stock_uom", "库存单位"),
        ("issued_qty", "发出"), ("returned_qty", "退回"), ("net_qty", "净发出"))),
    ("应付与未核销", "payables", (("voucher_no", "采购发票/贷项凭证"), ("posting_date", "记账日期"),
        ("source_scope", "金额及余额依据"), ("payable_account", "应付科目"), ("currency", "金额币种"),
        ("invoiced", "应付来源金额"), ("paid", "原生账簿已核销"), ("credit_note", "贷项"), ("outstanding", "原生余额"))),
    ("正式付款", "payments", (("voucher_no", "付款凭证"), ("posting_date", "记账日期"), ("payment_scope", "来源关系"),
        ("paid_amount", "付款金额"), ("currency", "付款账户币种"))),
    ("已取消凭证", "history", (("doctype", "凭证类型"), ("voucher_no", "凭证编号"), ("posting_date", "日期"), ("status", "状态"))),
)
NOTES = ("数量按商品和库存单位分行，不汇总不同单位；金额按各行币种展示，不跨币种合计。",
         "库存价值使用公司本币，不是工价、应付或付款金额。",
         "空白或—表示未能准确归属的金额；不等于零。混合付款/发票保留来源与状态，不分摊或猜测金额。",
         "已取消凭证仅作历史追溯，不进入有效数量或金额。",
         "取消发票及付款沿所选订单来源追溯至结束日期，日期可早于开始日期；不作为本期有效金额。")


def _deny():
    frappe.throw("没有完整打印或导出本次工厂对账及准确业务凭证的权限", frappe.PermissionError)


def _context(value):
    try:
        if isinstance(value, str):
            value = json.loads(value)
        if not isinstance(value, dict) or set(value) != CONTEXT_FIELDS:
            raise ValueError
        return dict(value)
    except (ValueError, TypeError):
        frappe.throw("工厂对账输出筛选无效，请重新查询")


def _checked(context, version, operation):
    result, documents, state = R._capture_statement(**context)
    if result["result_version"] != version:
        frappe.throw("对账来源、金额或权限已变化，请重新查询后打印或导出")
    business = {key: doc for key, doc in documents.items() if key[0] in R.BUSINESS_DOCTYPES}
    types = {key[0] for key in business} or {"Subcontracting Order"}
    if operation == "export":
        if any(not bool(frappe.permissions.can_export(doctype)) for doctype in types):
            _deny()
    elif not business and not frappe.has_permission("Subcontracting Order", "print"):
        _deny()
    for document in business.values():
        document.check_permission("read")
        if operation == "print":
            document.check_permission("print")
    return result, state


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        value = value.isoformat()
    elif isinstance(value, Decimal):
        if not value.is_finite():
            frappe.throw("对账来源包含非有限金额，请核对原单据")
        value = str(value)
    elif type(value) in (int, float):
        if type(value) is float and not math.isfinite(value):
            frappe.throw("对账来源包含非有限数量或金额，请核对原单据")
        value = str(value)
    if not isinstance(value, str) or any(not (char in "\t\n\r" or " " <= char <= "\ud7ff"
            or "\ue000" <= char <= "\ufffd" or "\U00010000" <= char <= "\U0010ffff") for char in value):
        frappe.throw("对账来源包含无法完整输出的文本，请核对原单据")
    if len(value.encode("utf-16-le")) // 2 > 32767:
        frappe.throw("对账来源文本超过Excel单元格限制，请先核对原单据")
    return value


def _summary(result):
    return [("公司", result["company"]), ("供应商", result["supplier"]), ("工厂类型", result["supplier_kind"]),
            ("开始日期", result["from_date"]), ("结束/应付截止日期", result["to_date"]),
            ("流程单筛选", result["flow_order"]), ("委外订单筛选", result["subcontracting_order"]),
            ("库存价值本币", result["company_currency"])]


def _notes(result):
    return [result[field] for field in ("payable_basis", "payment_basis", "settlement_action")] + list(NOTES)


def _tables(result):
    yield "工厂对账范围", ("项目", "值"), _summary(result)
    for title, key, columns in SECTIONS:
        yield title, tuple(label for _, label in columns), [tuple(row.get(field) for field, _ in columns)
                                                         for row in result[key]]
    yield "口径说明", ("说明",), [(note,) for note in _notes(result)]


def _xlsx(result):
    workbook = Workbook()
    try:
        for index, (title, headings, rows) in enumerate(_tables(result)):
            if len(rows) > 1_048_575:
                frappe.throw("工厂对账超过Excel行数限制，请缩小期间后完整导出")
            sheet = workbook.active if index == 0 else workbook.create_sheet(title)
            sheet.title, sheet.freeze_panes = title, "A2"
            for row_number, row in enumerate([headings, *rows], 1):
                for column, value in enumerate(row, 1):
                    cell = sheet.cell(row_number, column)
                    cell.value, cell.data_type, cell.number_format = _text(value), "s", "@"
            sheet.auto_filter.ref = f"A1:{sheet.cell(len(rows) + 1, len(headings)).coordinate}"
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


def _html(result):
    title = "电镀对账" if result["business_kind"] == "plating" else "委外业务对账"
    parts = ["<h1>" + title + "</h1>"]
    for section, headings, rows in _tables(result):
        parts.append("<h2>" + escape(section) + "</h2><table><thead><tr>")
        parts.extend("<th>" + escape(label) + "</th>" for label in headings)
        parts.append("</tr></thead><tbody>")
        for row in rows:
            parts.append("<tr>" + "".join("<td>" + escape(_text(value) or "—") + "</td>" for value in row) + "</tr>")
        if not rows:
            parts.append('<tr><td colspan="' + str(len(headings)) + '">当前范围无记录</td></tr>')
        parts.append("</tbody></table>")
    return ('<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>' + title + '</title><style>'
            '@page{size:A4 landscape;margin:10mm}body{font-family:sans-serif;font-size:9pt;color:#111}'
            'h1{font-size:17pt}h2{font-size:12pt;break-after:avoid}table{border-collapse:collapse;width:100%;table-layout:fixed;margin-bottom:1em}'
            'th,td{border:1px solid #aaa;padding:4px;overflow-wrap:anywhere;white-space:pre-wrap}'
            'thead{display:table-header-group}tr{break-inside:avoid}</style></head><body>' + "".join(parts) + '</body></html>')


def _output(context, result_version, operation):
    if frappe.session.user == "Guest" or not isinstance(result_version, str) or not re.fullmatch(r"[0-9a-f]{64}", result_version):
        _deny()
    context = _context(context)
    result, state = _checked(context, result_version, operation)
    content = _xlsx(result) if operation == "export" else _html(result)
    with _fresh_snapshot():
        current, current_state = _checked(context, result_version, operation)
        if current != result or current_state != state:
            frappe.throw("对账来源、金额或权限已变化，请重新查询后打印或导出")
    title = "电镀对账" if result["business_kind"] == "plating" else "委外业务对账"
    frappe.response.update({"type": "download", "content_type": CONTENT_TYPE if operation == "export" else "text/html",
        "filename": title + (".xlsx" if operation == "export" else ".html"),
        "display_content_as": "attachment", "filecontent": content})


@frappe.whitelist(methods=["GET"])
def export_reconciliation(context, result_version):
    _output(context, result_version, "export")


@frappe.whitelist(methods=["GET"])
def print_reconciliation(context, result_version):
    _output(context, result_version, "print")
