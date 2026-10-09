"""Actual statement source binding and XLSX/HTML bytes in permission models.

Native financial calculations retain their existing separate checks. These
tests do not claim remote, browser, printer, or positive ledger acceptance.
"""

from copy import deepcopy
from datetime import date, timedelta
from decimal import Decimal
import hashlib
from io import BytesIO
import importlib.util
import json
from pathlib import Path
import sys
import types
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class TestFactoryStatementOutput(unittest.TestCase):
    def setUp(self):
        fixtures = load("factory_statement_connections", ROOT / "tests/test_salary_history_export.py")
        self.fixture = fixtures.TestSalaryHistoryExport()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.env, self.frappe = self.fixture.env, self.fixture.env.frappe
        environment = self.env

        def check_permission(document, right):
            if not environment.permission(document.doctype, right, document):
                raise PermissionError(right)

        self.env.Document.check_permission = check_permission
        self.env.Document.as_dict = lambda document: dict(document)
        utils = types.ModuleType("frappe.utils")
        utils.getdate = lambda value: date.fromisoformat(value)
        self.payable = types.ModuleType("jifeng_erp.supplier_reconciliation")
        self.payable._permission_state = lambda documents: self.env.policy
        self.financial_state = "initial financial closure"
        self.payable._capture = Mock(side_effect=lambda *args: (None, "CNY", {}, {
            ("Journal Entry", "JE", "AP"): self.env.get_doc("Journal Entry", "JE")}, self.financial_state))
        modules = patch.dict(sys.modules, {"frappe.utils": utils, "jifeng_erp.supplier_reconciliation": self.payable})
        modules.start()
        self.addCleanup(modules.stop)
        self.R = load("factory_statement_sources", ROOT / "jifeng_erp/subcontracting_reconciliation.py")
        canonical = patch.dict(sys.modules, {"jifeng_erp.subcontracting_reconciliation": self.R})
        canonical.start()
        self.addCleanup(canonical.stop)
        self.E = load("factory_statement_output", ROOT / "jifeng_erp/subcontracting_reconciliation_output.py")
        for doctype, name in (("Subcontracting Order", "SCO"), ("Subcontracting Receipt", "SCR"),
            ("Purchase Order", "PO"), ("Purchase Invoice", "PI"), ("Payment Entry", "PE"),
            ("Stock Entry", "STE"), ("Jifeng Flow Order", "FLOW"), ("Journal Entry", "JE")):
            self.env.add(doctype, name, docstatus=1, company="COMPANY", supplier="SUP", private_note="PRIVATE")
        self.env.add("Subcontracting Receipt", "CANCELLED", docstatus=2, company="COMPANY", supplier="SUP")
        self.data = dict(company="COMPANY", supplier="SUP", supplier_kind="电镀厂", business_kind="plating",
            from_date="2026-10-01", to_date="2026-10-08", flow_order="FLOW", subcontracting_order="SCO", company_currency="CNY",
            progress=[dict(subcontracting_order="SCO", supplier_name="电镀厂一", flow_order="FLOW", item_code="FG", stock_uom="Nos",
                planned_qty="5", accepted_qty="2", rejected_qty="1", loss_qty="1", pending_qty="1")],
            receipts=[dict(receipt="SCR", entry_type="退回", subcontracting_order="SCO", flow_order="FLOW", item_code="FG", stock_uom="Nos",
                accepted_qty="-1", rejected_qty="0", loss_qty="0", inventory_qty_change="-1", inventory_value_change="-12.50", inventory_value_status="已过账")],
            materials=[dict(stock_entry="STE", entry_type="发料", subcontracting_order="SCO", flow_order="FLOW", rm_item_code="RM", stock_uom="Kg",
                issued_qty="2.5", returned_qty="1.25", net_qty="1.25")],
            payables=[dict(voucher_no="PI", posting_date=date(2026, 10, 2), source_scope="发票原生余额（当前；AP截至日未列示）",
                payable_account="AP", currency="CNY", invoiced="4", paid=None, credit_note=None, outstanding="0")],
            payments=[dict(voucher_no="PE", posting_date="2026-10-03", payment_scope="混合来源，无法归入本委外范围", paid_amount=None, currency=None)],
            history=[dict(doctype="Subcontracting Receipt", voucher_no="CANCELLED", posting_date="2026-10-03", status="已取消")],
            payable_basis="AP按截至日；发票当前余额逐行注明", payment_basis="仅精确PE付款来源，不以应付余额推实付",
            settlement_action="请打开原生凭证办理")
        self.context = {key: self.data[key] for key in self.E.CONTEXT_FIELDS}

        def query(*args, **kwargs):
            # The existing source-query functions are tested independently;
            # model them here while exercising the real collector and output.
            self.R._remember(self.env.get_doc("Purchase Order", "PO"))
            return deepcopy(self.data)

        self.R._query_reconciliation = Mock(side_effect=query)

    def query(self):
        return self.R.query_reconciliation(**self.context)

    def output(self, operation="export", version=None):
        getattr(self.E, operation + "_reconciliation")(json.dumps(self.context), version or self.query()["result_version"])
        return self.frappe.response["filecontent"]

    def workbook(self, content):
        workbook = load_workbook(BytesIO(content))
        self.addCleanup(workbook.close)
        return workbook

    def test_statement_keeps_six_business_sections_units_cutoff_basis_and_unknown_amounts(self):
        workbook = self.workbook(self.output())
        self.assertEqual(workbook.sheetnames, ["工厂对账范围", "委外进度", "收回与库存价值", "委外原料发退",
            "应付与未核销", "正式付款", "已取消凭证", "口径说明"])
        self.assertEqual(list(workbook["收回与库存价值"].values)[1][6:11], ("-1", "0", "0", "-1", "-12.50"))
        self.assertEqual(list(workbook["委外原料发退"].values)[1][5:], ("Kg", "2.5", "1.25", "1.25"))
        payable = list(workbook["应付与未核销"].values)[1]
        self.assertEqual(payable[-4:], ("4", None, None, "0"))
        self.assertIn("当前；AP截至日未列示", payable[2])
        self.assertEqual(list(workbook["正式付款"].values)[1][-2:], (None, None))
        self.assertEqual(list(workbook["已取消凭证"].values)[1][1], "CANCELLED")
        self.assertNotIn("PRIVATE", str([list(sheet.values) for sheet in workbook]))
        self.assertEqual(self.fixture.events[-2:], [("rollback", None), ("close", None)])

    def test_complete_rows_and_exact_decimal_and_foreign_currency_are_not_aggregated(self):
        self.data["payments"] = [dict(voucher_no="PE", posting_date="2026-10-03", payment_scope="发票已核销付款",
                                     paid_amount="9007199254740993.017", currency="USD") for _ in range(21)]
        workbook = self.workbook(self.output())
        self.assertEqual(workbook["正式付款"].max_row, 22)
        self.assertEqual(workbook["正式付款"].cell(2, 4).value, "9007199254740993.017")
        self.assertEqual(workbook["正式付款"].cell(2, 5).value, "USD")
        self.assertEqual(workbook["正式付款"].cell(2, 4).data_type, "s")

    def test_print_and_excel_source_text_cannot_execute_html_or_formulas(self):
        self.data["progress"][0]["supplier_name"] = '=HYPERLINK("https://example.invalid")'
        self.data["payables"][0]["source_scope"] = '<script>alert("bad")</script>'
        html = self.output("print")
        self.assertIn("&lt;script&gt;", html)
        self.assertNotIn("<script>", html)
        self.assertIn("thead{display:table-header-group}", html)
        workbook = self.workbook(self.output())
        cell = workbook["委外进度"].cell(2, 2)
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(cell.value.startswith("="))

    def test_each_nonempty_business_source_requires_its_own_print_and_export_right(self):
        sources = {doctype for doctype, _ in self.env.records if doctype in self.R.BUSINESS_DOCTYPES}
        for doctype in sorted(sources):
            with self.subTest(doctype=doctype):
                self.env.denials.add((doctype, "print"))
                with self.assertRaises(PermissionError):
                    self.output("print")
                self.env.denials.clear()
                self.fixture.roles.clear()
                self.fixture.export_grants = sources - {doctype}
                with self.assertRaises(PermissionError):
                    self.output()
                self.fixture.roles.add("System Manager")
                self.assertNotIn("filecontent", self.frappe.response)

    def test_hidden_source_and_changed_source_version_reject_entire_output(self):
        version = self.query()["result_version"]
        self.env.records["Subcontracting Receipt", "SCR"].modified = self.fixture.F.LATER
        with self.assertRaisesRegex(ValueError, "重新查询"):
            self.output(version=version)
        self.env.denials.add(("Subcontracting Receipt", "SCR", "read"))
        with self.assertRaises(PermissionError):
            self.output()
        self.assertNotIn("filecontent", self.frappe.response)

    def test_financial_ledger_or_permission_change_before_publish_leaks_no_bytes(self):
        self.fixture.on_begin = lambda: setattr(self, "financial_state", "ledger changed")
        with self.assertRaisesRegex(ValueError, "重新查询"):
            self.output()
        self.fixture.on_begin = lambda: self.env.denials.add(("Purchase Invoice", "print"))
        with self.assertRaises(PermissionError):
            self.output("print")
        self.assertNotIn("filecontent", self.frappe.response)

    def test_guest_malformed_scope_and_unknown_query_version_never_output(self):
        self.frappe.session.user = "Guest"
        with self.assertRaises(PermissionError):
            self.output(version="a" * 64)
        self.frappe.session.user = "qa@example.test"
        with self.assertRaises(ValueError):
            self.E.export_reconciliation(json.dumps({**self.context, "unexpected": 1}), "a" * 64)
        with self.assertRaises(PermissionError):
            self.output(version="bad")
        self.assertNotIn("filecontent", self.frappe.response)

    def test_collector_resets_after_exception_and_no_plaintext_source_state_is_returned(self):
        data = self.query()
        self.assertIsNone(self.R._STATEMENT.get())
        self.assertNotIn("documents", data)
        self.assertNotIn("states", data)
        self.assertNotIn("PRIVATE", str(data))
        self.R._query_reconciliation.side_effect = ValueError("fixture failure")
        with self.assertRaises(ValueError):
            self.query()
        self.assertIsNone(self.R._STATEMENT.get())

    def test_all_endpoints_get_only_and_text_limits_do_not_silently_truncate(self):
        self.assertIn(("query_reconciliation", {"methods": ["GET"]}), self.env.whitelisted)
        self.assertIn(("export_reconciliation", {"methods": ["GET"]}), self.env.whitelisted)
        self.assertIn(("print_reconciliation", {"methods": ["GET"]}), self.env.whitelisted)
        with self.assertRaises(ValueError):
            self.E._text("a" * 32768)
        with self.assertRaises(ValueError):
            self.E._text(float("nan"))

    def test_native_stock_entry_time_and_nested_child_times_keep_exact_query_and_fresh_output(self):
        document = self.env.records["Stock Entry", "STE"]
        document.posting_time = timedelta(hours=8, minutes=23, seconds=17, microseconds=123456)
        document["items"] = [{"doctype": "Stock Entry Detail", "name": "STE-ROW", "time": timedelta(microseconds=-1)}]
        first = self.query()
        self.assertEqual(first["result_version"], self.query()["result_version"])
        self.assertTrue(self.output("print", first["result_version"]).startswith("<!doctype html>"))
        self.assertEqual(self.fixture.events[-2:], [("rollback", None), ("close", None)])
        document["items"][0]["time"] = timedelta(0)
        self.assertNotEqual(first["result_version"], self.query()["result_version"])

    def test_one_microsecond_native_time_change_rejects_stale_and_fresh_outputs(self):
        document = self.env.records["Stock Entry", "STE"]
        document.posting_time = timedelta(hours=12, microseconds=1)
        version = self.query()["result_version"]
        document.posting_time += timedelta(microseconds=1)
        with self.assertRaisesRegex(ValueError, "重新查询"):
            self.output(version=version)
        self.assertNotIn("filecontent", self.frappe.response)
        self.fixture.on_begin = lambda: document.update(posting_time=document.posting_time + timedelta(microseconds=1))
        with self.assertRaisesRegex(ValueError, "重新查询"):
            self.output("print")
        self.assertNotIn("filecontent", self.frappe.response)

    def test_native_time_is_not_a_lookalike_json_value_and_unknown_types_are_rejected(self):
        actual = self.R._statement_version({"posting_time": timedelta(days=-1, seconds=86399, microseconds=999999)}, ())
        literal = self.R._statement_version({"posting_time": {"native_time_delta": [-1, 86399, 999999]}}, ())
        self.assertNotEqual(actual, literal)
        self.assertEqual(self.R._fingerprint_native(timedelta(microseconds=-1)),
                         {"native_time_delta": [-1, 86399, 999999]})
        for invalid in (object(), float("nan"), float("inf")):
            with self.subTest(invalid=type(invalid).__name__), self.assertRaisesRegex(ValueError, "未支持"):
                self.R._statement_version({"native_value": invalid}, ())

    def test_streamed_fingerprint_keeps_exact_prior_bytes_without_full_json_copy(self):
        freeze = self.fixture.tools._freeze
        result = {"scope": "电镀厂", "time": timedelta(microseconds=-1), "amount": Decimal("10.00001"),
                  "date": date(2026, 10, 9), "rows": [None, True, -2, 1.25, {"x": "\\\"\n"}]}
        state = freeze(({("Stock Entry", "STE"): {"time": timedelta(days=2, seconds=3)}},
                        ("permission", [("field", "company", False), ("field", "qty", 0)])))
        previous = hashlib.sha256(json.dumps(freeze((result, state)), sort_keys=True, ensure_ascii=False,
            allow_nan=False, default=self.R._fingerprint_native).encode("utf-8")).hexdigest()
        with patch.object(self.R.json, "dumps", side_effect=AssertionError("no whole JSON document")):
            self.assertEqual(previous, self.R._statement_version(result, state))
        self.assertNotEqual(previous, self.R._statement_version(result, state + ("permission changed",)))

    def test_already_frozen_financial_and_permission_closure_is_retained_without_recopy(self):
        self.financial_state = ("ledger", (("source", "JE"), ("amount", "4")))
        self.payable._permission_state = lambda documents: permission_state
        permission_state = ("raw permissions", (("DocField", "company"), ("mask", False)))
        _, _, state = self.R._capture_statement(**self.context)
        self.assertIs(state[2], self.financial_state)
        self.assertIs(state[3], permission_state)
        self.assertEqual(state, self.fixture.tools._freeze(state))
        saved = deepcopy(state)
        self.env.records["Stock Entry", "STE"].modified = self.fixture.F.LATER
        self.assertEqual(state, saved)


if __name__ == "__main__":
    unittest.main()
