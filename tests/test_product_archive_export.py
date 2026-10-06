"""Actual product export source and XLSX bytes under isolated native models.

These tests do not access a site, a database, a browser or HTTP. A modeled
export grant is explicit; it does not certify an ordinary deployed role.
"""

from contextlib import contextmanager
from copy import deepcopy
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import importlib.util
from pathlib import Path
import sys
import types
import unittest
from unittest.mock import Mock, patch
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook


ROOT = Path(__file__).resolve().parents[1]


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class Database:
    def __init__(self, model, **settings):
        self.__dict__.update(settings)
        self.begin = Mock(side_effect=lambda **kw: model.event("begin", kw))
        self.rollback = Mock(side_effect=lambda: model.event("rollback"))
        self.close = Mock(side_effect=lambda: model.event("close"))


class TestProductArchiveExport(unittest.TestCase):
    def setUp(self):
        self.F = load("product_export_fixtures", ROOT / "tests/test_product_archive.py")
        self.env = self.F.TestProductArchive("test_init_has_explicit_company_and_business_fields_without_native_defaults")
        self.env.setUp()
        self.addCleanup(self.env.doCleanups)
        self.P = self.env.module
        self.env.frappe.session.user = "export-enabled-model@example.invalid"
        for cls, name, action in ((self.F.NativeDocument, "as_dict", lambda doc: deepcopy(dict(doc))),
                (self.F.Meta, "get_permlevel_access", lambda meta, permission, parenttype=None: set())):
            patcher = patch.object(cls, name, action, create=True)
            patcher.start()
            self.addCleanup(patcher.stop)
        fields = ("file_url", "is_private", "attached_to_doctype", "attached_to_name", "attached_to_field")
        self.env.metas["File"] = self.F.Meta([{"fieldname": field, "fieldtype": "Data"} for field in fields])
        self.env.profile.update(list_price=Decimal("17.50"), blank_price=0, weight="0.123456789012345678901234",
            silver_weight=0, stone_weight=0, profile_image="/private/files/product.png",
            patent_image="/private/files/patent.png")
        self.env.attachment()
        self.env.attachment("PATENT-FILE", "patent_image", "/private/files/patent.png")
        self.events, self.failures, self.old, self.omissions = [], set(), None, {}
        self.settings = {"socket": "/synthetic/product.sock", "host": "synthetic-primary", "port": 13316,
            "user": "synthetic-export", "password": "synthetic-only", "cur_db_name": "synthetic-db"}
        self.original = Database(self, **self.settings)
        self.env.frappe.local = self.F.Record(db=self.original, site="synthetic.localhost", conf=self.F.Record(
            db_type="mariadb", db_socket=self.settings["socket"], db_host=self.settings["host"], db_port=self.settings["port"],
            db_user=self.settings["user"], db_password=self.settings["password"], db_name=self.settings["cur_db_name"]))
        self.factory = Mock(side_effect=lambda **kw: Database(self, **kw))
        dbmodule = types.ModuleType("frappe.database")
        dbmodule.get_db = self.factory
        module_scope = patch.dict(sys.modules, {"frappe.database": dbmodule,
            "jifeng_erp.tool_archive": types.SimpleNamespace(DOCTYPES={}),
            "jifeng_erp.plate_modeling": types.SimpleNamespace(DOCTYPE="Jifeng Plate Modeling")})
        module_scope.start()
        self.addCleanup(module_scope.stop)
        self.C = load("product_export_field_masks", ROOT / "jifeng_erp/inventory_catalog.py")
        self.tools = load("product_export_fresh_snapshot", ROOT / "jifeng_erp/tool_archive_export.py")
        imports = patch.dict(sys.modules, {"jifeng_erp.inventory_catalog": self.C,
            "jifeng_erp.tool_archive_export": self.tools})
        imports.start()
        self.addCleanup(imports.stop)
        self.auth = {"User": [{"name": self.env.frappe.session.user, "enabled": 1, "user_type": "System User"}],
            "DocPerm": [{"name": "explicit-model-export", "parent": self.P.PROFILE, "read": 1, "export": 1}],
            "Property Setter": [{"name": "field-mask", "doc_type": self.P.PROFILE, "property": "mask", "value": "0"}]}
        self.env.frappe.get_all = Mock(side_effect=self.get_all)
        self.env.frappe.get_doc = Mock(side_effect=self.get_doc)
        self.env.frappe.get_list = Mock(side_effect=self.get_list)
        model = self

        class Response(self.F.Record):
            def update(response, *args, **kw):
                model.assertIs(model.env.frappe.local.db, model.original)
                model.assertEqual(model.events[-2:], [("rollback", None), ("close", None)])
                super().update(*args, **kw)
        self.env.frappe.response = Response()
        self.E = load("product_archive_export_under_test", ROOT / "jifeng_erp/product_archive_export.py")

    def event(self, name, options=None):
        self.events.append((name, options))
        if name in self.failures:
            raise RuntimeError("synthetic " + name + " failure")

    @contextmanager
    def view(self):
        current = self.env.documents
        if self.old and self.env.frappe.local.db is self.original:
            self.env.documents = self.old["documents"]
        try:
            yield
        finally:
            self.env.documents = current

    def get_doc(self, *args, **kw):
        with self.view():
            return self.env.get_doc(*args, **kw)

    def get_list(self, doctype, filters, fields, **kw):
        with self.view():
            result = self.env.get_list(doctype, filters, fields, **kw)
            # SQL projection includes an explicit None for an unset column.
            result = [self.F.Record({field: row.get(field) for field in fields
                if field not in self.omissions.get(doctype, set())}) for row in result]
            length = kw.get("limit_page_length")
            return result[:length] if length else result

    def get_all(self, doctype, filters=None, fields=None, **kw):
        with self.view():
            if doctype == "File":
                result = self.env.get_all(doctype, filters, **kw)
                length = kw.get("limit_page_length")
                return result[:length] if length else result
            auth = self.old["auth"] if self.old and self.env.frappe.local.db is self.original else self.auth
            def matches(row):
                return all(row.get(field) in value[1] if isinstance(value, (list, tuple)) and value[0] == "in"
                    else row.get(field) == value for field, value in (filters or {}).items())
            return [self.F.Record(deepcopy(row) if fields == ["*"] else {field: row.get(field) for field in fields})
                for row in auth.get(doctype, []) if matches(row)]

    def export(self, company="JF", search="", include_inactive=0):
        self.env.frappe.response.clear()
        self.E.export_archives(company, search, include_inactive)

    def sheet(self):
        workbook = load_workbook(BytesIO(self.env.frappe.response["filecontent"]), data_only=False)
        self.addCleanup(workbook.close)
        return workbook.active

    def no_binary(self):
        self.assertNotIn("filecontent", self.env.frappe.response)
        self.assertIs(self.env.frappe.local.db, self.original)

    def change(self, action, stale=True):
        if stale:
            self.old = {"documents": deepcopy(self.env.documents), "auth": deepcopy(self.auth)}
        original = Workbook.save
        def changed(workbook, stream):
            result = original(workbook, stream)
            action()
            return result
        return patch.object(Workbook, "save", changed)

    def add_profile(self, name, **changes):
        self.env.add("Item", "ITEM-" + name, item_code="ITEM-" + name)
        values = {key: deepcopy(value) for key, value in self.env.profile.items()
            if key not in ("name", "doctype", "modified", "item_code", "profile_image", "patent_image")}
        self.env.add(self.P.PROFILE, name, item_code="ITEM-" + name, profile_image=None, patent_image=None, **(values | changes))

    def test_complete_28_fields_literal_cells_and_readonly_get(self):
        before = deepcopy((self.env.documents, self.env.versions, self.env.saves))
        self.export()
        sheet = self.sheet()
        self.assertEqual(sheet.max_column, 28)
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(tuple(cell.value for cell in sheet[1]), tuple(label for _, label in self.E.COLUMNS))
        values = [cell.value for row in sheet for cell in row]
        for forbidden in ("BUSINESS-ITEM", "PRIVATE-FILE", "PATENT-FILE", "JF", "/private/files/product.png",
                          "/private/files/patent.png?fid=PATENT-FILE", "SUP", "CATEGORY-1"):
            self.assertNotIn(forbidden, values)
        self.assertEqual(sheet.cell(2, 19).value, "17.50")
        self.assertEqual(sheet.cell(2, 21).value, "0.123456789012345678901234")
        self.assertEqual([sheet.cell(2, column).value for column in (27, 28)], ["已上传", "已上传"])
        self.assertEqual(self.env.frappe.response["filename"], "产品档案.xlsx")
        self.assertEqual(self.E.export_archives.allowed_methods, ["GET"])
        self.assertEqual(before, (self.env.documents, self.env.versions, self.env.saves))
        self.env.frappe.new_doc.assert_not_called()
        self.assertEqual(self.events, [("begin", {"read_only": True}), ("rollback", None), ("close", None)])
        for value in ("00123", "=1+1", "+001", "#N/A", "<b>&原文</b>", "  文本\t\n "):
            self.env.profile["notes"] = value
            self.export()
            self.assertEqual(self.sheet().cell(2, 26).value, value)
            self.assertTrue(all(cell.data_type in ("s", "inlineStr") and cell.number_format == "@" and cell.hyperlink is None
                for row in self.sheet() for cell in row))
            with ZipFile(BytesIO(self.env.frappe.response["filecontent"])) as archive:
                xml = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
                ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                self.assertFalse(xml.findall(".//s:f", ns))
                self.assertFalse(xml.findall(".//s:hyperlinks", ns))
                self.assertFalse(any(path.startswith("xl/externalLinks/") for path in archive.namelist()))

    def test_all_saved_rows_beyond_1000_and_exact_filter_scope(self):
        for index in range(1000):
            self.add_profile("SYNTHETIC-%04d" % index)
        self.add_profile("OTHER-COMPANY", company="OTHER")
        self.add_profile("INACTIVE", inactive=1, display_name="INACTIVE-ONLY")
        self.export()
        self.assertEqual(self.sheet().max_row, 1002)
        self.export(include_inactive=1)
        self.assertEqual(self.sheet().max_row, 1003)
        self.export(search="INACTIVE-ONLY", include_inactive=1)
        self.assertEqual(self.sheet().max_row, 2)
        self.export(search="INACTIVE-ONLY")
        self.assertEqual(self.sheet().max_row, 1)
        calls = [call for call in self.env.frappe.get_list.call_args_list if call.args[0] == self.P.PROFILE]
        self.assertTrue(all(call.kwargs["limit_page_length"] == 1048576 for call in calls))

    def test_initial_or_mid_read_hidden_masked_missing_fields_refuse_entire_file(self):
        for doctype, field in ((self.P.PROFILE, "notes"), (self.P.PROFILE, "profile_image"), ("Item", "item_code"), ("File", "attached_to_field")):
            self.env.hidden[doctype, "read"] = {field}
            with self.subTest(doctype=doctype, field=field), self.assertRaises(PermissionError):
                self.export()
            self.no_binary()
            self.env.hidden.clear()
        masked = self.env.metas[self.P.PROFILE].get_field("notes")
        masked.mask, masked.permlevel = 1, 0
        with self.assertRaises(PermissionError):
            self.export()
        self.no_binary()
        masked.mask = 0
        for doctype, field in ((self.P.PROFILE, "material"), ("Item", "modified")):
            self.omissions[doctype] = {field}
            with self.assertRaises(PermissionError):
                self.export()
            self.no_binary()
            self.omissions.clear()
        self.add_profile("ZZZ-LATER")
        original = self.get_doc
        def withdraw(doctype, name, **kw):
            result = original(doctype, name, **kw)
            if doctype == self.P.PROFILE and name == "ZZZ-LATER":
                self.env.hidden[self.P.PROFILE, "read"] = {"notes"}
            return result
        with patch.object(self.env.frappe, "get_doc", side_effect=withdraw), self.assertRaises(PermissionError):
            self.export()
        self.no_binary()

    def test_role_document_guest_company_source_and_empty_export_denials(self):
        actions = (lambda: self.env.denials.add((self.P.PROFILE, "export")),
            lambda: self.env.denials.add((self.P.PROFILE, "read")),
            lambda: self.env.denials.add(("Item", "read")),
            lambda: self.env.document_denials.add((self.P.PROFILE, "BUSINESS-ITEM", "export")),
            lambda: self.env.document_denials.add(("Item", "BUSINESS-ITEM", "read")),
            lambda: self.env.document_denials.add(("Company", "JF", "select")),
            lambda: setattr(self.env.frappe.session, "user", "Guest"))
        for action in actions:
            action()
            with self.assertRaises(PermissionError):
                self.export()
            self.no_binary()
            self.env.denials.clear()
            self.env.document_denials.clear()
            self.env.frappe.session.user = "export-enabled-model@example.invalid"
        self.env.denials.add((self.P.PROFILE, "export"))
        with self.assertRaises(PermissionError):
            self.export(search="no-result")
        self.no_binary()

    def test_two_exact_image_states_never_export_alias_or_unreadable_url(self):
        actions = ((lambda: self.env.profile.update(profile_image=None), "未上传"),
            (lambda: self.env.documents.pop(("File", "PRIVATE-FILE")), "附件不可用"),
            (lambda: self.env.document_denials.add(("File", "PRIVATE-FILE", "read")), "附件不可用"),
            (lambda: self.env.denied_downloads.add("PRIVATE-FILE"), "附件不可用"),
            (lambda: self.env.documents["File", "PRIVATE-FILE"].update(attached_to_field="patent_image"), "附件不可用"),
            (lambda: self.env.profile.update(profile_image="/files/public.png"), "附件不可用"))
        for action, state in actions:
            previous = deepcopy(self.env.documents)
            action()
            self.export()
            self.assertEqual(self.sheet().cell(2, 27).value, state)
            self.assertEqual(self.sheet().cell(2, 28).value, "已上传")
            self.env.documents = previous
            self.env.profile = self.env.documents[self.P.PROFILE, "BUSINESS-ITEM"]
            self.env.document_denials.clear()
            self.env.denied_downloads.clear()

    def test_post_serialization_sources_range_and_permissions_rechecked_fresh(self):
        actions = (lambda: self.env.profile.update(notes="committed without modified"),
            lambda: self.env.item.update(item_name="changed source"),
            lambda: self.env.documents["File", "PRIVATE-FILE"].update(attached_to_field="other"),
            lambda: self.add_profile("NEW-MATCH"),
            lambda: self.auth["Property Setter"][0].update(value="1"),
            lambda: self.env.document_denials.add((self.P.PROFILE, "BUSINESS-ITEM", "export")),
            lambda: self.env.document_denials.add(("Item", "BUSINESS-ITEM", "read")))
        for action in actions:
            documents, auth = deepcopy(self.env.documents), deepcopy(self.auth)
            with self.change(action), self.assertRaises((ValueError, PermissionError)):
                self.export()
            self.no_binary()
            self.env.documents, self.auth, self.old = documents, auth, None
            self.env.profile = self.env.documents[self.P.PROFILE, "BUSINESS-ITEM"]
            self.env.item = self.env.documents["Item", "BUSINESS-ITEM"]
            self.env.document_denials.clear()

    def test_duplicate_row_projection_scope_and_versions_cannot_silently_skip(self):
        projection = self.get_list(self.P.PROFILE, {"company": "JF"}, sorted(self.E.PROFILE_FIELDS))
        bad_rows = (projection * 2, [self.F.Record(projection[0] | {"notes": "not the native projection"})],
            [self.F.Record(projection[0] | {"company": "OTHER"})],
            [self.F.Record(projection[0] | {"modified": datetime(2026, 10, 1, 9, 0, 0, 1)})])
        for rows in bad_rows:
            self.env.forced_lists[self.P.PROFILE] = rows
            with self.assertRaises((ValueError, PermissionError)):
                self.export()
            self.no_binary()
        self.env.forced_lists.clear()
        self.env.item["item_code"] = "OTHER-IDENTITY"
        with self.assertRaises(PermissionError):
            self.export()
        self.no_binary()

    def test_invalid_values_excel_bounds_and_attachment_overflow_refuse(self):
        for field, value in (("notes", "x" * 32768), ("notes", "😀" * 16384), ("notes", "x\x00y"),
                ("notes", "x\ud800y"), ("notes", "x\ufffey"), ("notes", 123), ("inactive", 2),
                ("weight", float("nan")), ("weight", "Infinity"), ("profile_image", 12)):
            previous = self.env.profile.get(field)
            self.env.profile[field] = value
            with self.assertRaises((ValueError, PermissionError)):
                self.export(include_inactive=1)
            self.no_binary()
            self.env.profile[field] = previous
        with patch.object(self.E, "MAX_EXPORT_ROWS", 0), self.assertRaisesRegex(ValueError, "Excel行数限制"):
            self.export()
        self.no_binary()
        with patch.object(self.E, "MAX_IMAGE_ATTACHMENTS", 0), self.assertRaises(ValueError):
            self.export()
        self.no_binary()

    def test_invalid_scope_and_readonly_connection_cleanup_failures_never_publish(self):
        for args in ((" JF", "", 0), ("JF", [], 0), ("JF", "x" * 141, 0), ("JF", "", 2)):
            with self.assertRaises((ValueError, PermissionError)):
                self.export(*args)
            self.no_binary()
        for failure in ("begin", "rollback", "close"):
            self.failures.add(failure)
            with self.assertRaises(RuntimeError):
                self.export()
            self.no_binary()
            self.failures.clear()
        self.env.frappe.local.conf.db_user = "root"
        with self.assertRaises(ValueError):
            self.export()
        self.no_binary()


if __name__ == "__main__":
    unittest.main()
