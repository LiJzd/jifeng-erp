"""Real XLSX serialization with native archive/source permission models; no DB."""

from copy import deepcopy
from io import BytesIO
import importlib.util
from pathlib import Path
import sys
import types
import unittest
from unittest.mock import Mock, patch
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook


ROOT = Path(__file__).resolve().parents[1]
TEXT_FIELDS = ("legacy_number", "display_name", "owner_name", "mobile", "reception", "phone",
               "address", "notes", "bank_name", "account_name", "bank_account")
HEADERS = ("编号", "名称", "负责", "手机", "接货", "电话", "地址", "备注", "开户", "户名", "账号")
PLATING_HEADERS = (*HEADERS[:8], "户行", *HEADERS[9:], "停用", "操作", "日期", "修改人", "修改日期")
XML_NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class TestPartnerArchiveExport(unittest.TestCase):
    def setUp(self):
        search_path = patch.object(sys, "path", [str(ROOT), *sys.path])
        search_path.start()
        self.addCleanup(search_path.stop)
        fixtures = load_module("archive_export_fixtures", ROOT / "tests/test_partner_archive.py")
        self.fixtures = fixtures
        self.env = fixtures.TestPartnerArchive("test_joint_save_preserves_raw_and_shared_supplier")
        self.addCleanup(self.env.doCleanups)
        self.env.setUp()
        self.P = self.env.module
        self.env.metas[self.P.PROFILE].fields.append(fixtures.Record(fieldname="modified", fieldtype="Datetime"))
        self.env.frappe.response = fixtures.Record()
        self.env.frappe._ = lambda text: text
        self.methods = []
        self.env.frappe.whitelist = lambda **kwargs: self.decorator(kwargs)
        self.auth = {
            "User": [{"name": self.env.frappe.session.user, "enabled": 1, "user_type": "System User"}],
            "Has Role": [{"name": "EXPORT-ROLE", "parenttype": "User", "parent": self.env.frappe.session.user,
                          "role": "Purchase Manager"}],
            "DocPerm": [{"name": "EXPORT-PERM", "parent": self.P.PROFILE, "role": "Purchase Manager", "read": 1, "export": 1}],
            "DocField": [{"name": "BANK-FIELD", "parent": self.P.PROFILE, "fieldname": "bank_account", "mask": 0}],
            "Property Setter": [{"name": "BANK-MASK", "doc_type": self.P.PROFILE, "field_name": "bank_account",
                                 "property": "mask", "value": "0"}],
            "DocShare": [{"name": "EXPORT-SHARE", "share_doctype": self.P.PROFILE, "share_name": "SYNTHETIC-000",
                          "user": self.env.frappe.session.user, "everyone": 0, "read": 1}],
        }
        self.env.frappe.get_all = Mock(side_effect=self.get_all)
        self.env.records["Supplier", "SUP-A"].update(bank_account="SUPPLIER-FALLBACK-FORBIDDEN", mobile="SUPPLIER-RAW-FORBIDDEN")
        package = types.ModuleType("jifeng_erp")
        package.__path__ = [str(ROOT / "jifeng_erp")]
        package.partner_archive = self.P
        self.dependencies = patch.dict(sys.modules, {"jifeng_erp": package, "jifeng_erp.partner_archive": self.P})
        self.dependencies.start()
        self.addCleanup(self.dependencies.stop)
        self.module = load_module("partner_archive_export_model", ROOT / "jifeng_erp/partner_archive_export.py")
        self.add_profile()
        self.original_records = deepcopy(self.env.records)

    def decorator(self, options):
        self.methods.append(options)
        return lambda function: function

    def get_all(self, doctype, filters=None, fields=None, **kwargs):
        def matches(row):
            for field, value in (filters or {}).items():
                if isinstance(value, list) and value[0] == "in":
                    if row.get(field) not in value[1]:
                        return False
                elif row.get(field) != value:
                    return False
            return True
        rows = [deepcopy(row) for row in self.auth.get(doctype, []) if matches(row)]
        if fields and fields != ["*"]:
            rows = [{field: row.get(field) for field in fields} for row in rows]
        return [self.fixtures.Record(row) for row in rows]

    def add_profile(self, index=0, kind="外发工厂", **values):
        name = f"SYNTHETIC-{index:03d}"
        row = {"doctype": self.P.PROFILE, "name": name, "company": "Company A", "archive_kind": kind,
               "supplier": "SUP-A", "modified": "2026-10-02T00:00:00.123456", "creation": "2026-10-01T00:00:00.123456",
               "owner": "synthetic-creator", "modified_by": "synthetic-modifier", "request_id": "INTERNAL-REQUEST-FORBIDDEN",
               "request_digest": "INTERNAL-DIGEST-FORBIDDEN", **self.env.payload()["profile"], **values}
        self.env.records[self.P.PROFILE, name] = row
        return row

    def export(self, kind="外发工厂", company="Company A", search="", status="在用"):
        return self.module.export_archives(kind, company, search, status)

    def worksheet(self):
        content = self.env.frappe.response.get("filecontent")
        self.assertIsInstance(content, bytes)
        workbook = load_workbook(BytesIO(content), data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(len(workbook.sheetnames), 1)
        return workbook.active

    def assert_no_binary(self):
        self.assertTrue("filecontent" not in self.env.frappe.response, "Refused export must not publish binary content")
        self.assertNotEqual(self.env.frappe.response.get("type"), "binary")

    def during_serialization(self, change, events=None):
        original = Workbook.save
        def save(workbook, target):
            if events is not None:
                events.append(("xlsx", "save"))
            result = original(workbook, target)
            change()
            return result
        return patch.object(Workbook, "save", new=save)

    def test_get_response_whitelist_exact_rows_without_writes_or_internal_identity(self):
        for kind, headers in (("外发工厂", HEADERS), ("电镀厂", PLATING_HEADERS)):
            with self.subTest(kind=kind):
                self.env.records[self.P.PROFILE, "SYNTHETIC-000"]["archive_kind"] = kind
                self.env.frappe.response.clear()
                before = deepcopy(self.env.records)
                self.export(kind=kind)
                self.assertEqual(self.env.records, before)
                sheet = self.worksheet()
                self.assertEqual(tuple(cell.value for cell in sheet[1]), headers)
                self.assertEqual(sheet.max_column, 11 if kind == "外发工厂" else 16)
                self.assertEqual(sheet.max_row, 2)
                row = self.env.records[self.P.PROFILE, "SYNTHETIC-000"]
                self.assertEqual(tuple(cell.value for cell in sheet[2][:11]), tuple(row[field] for field in TEXT_FIELDS))
                if kind == "电镀厂":
                    self.assertEqual(sheet.cell(2, 13).value, row["owner"])
                    self.assertEqual(sheet.cell(2, 14).value, row["creation"])
                    self.assertEqual(sheet.cell(2, 15).value, row["modified_by"])
                    self.assertEqual(sheet.cell(2, 16).value, row["modified"])
                self.assertEqual(self.env.frappe.response["filename"], kind + "档案.xlsx")
                self.assertEqual(self.env.frappe.response["type"], "binary")
                self.assertEqual(self.env.frappe.response["content_type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.assertEqual(self.env.frappe.response["display_content_as"], "attachment")
                cells = [str(cell.value) for rows in sheet for cell in rows]
                for forbidden in ("SYNTHETIC-000", "SUP-A", "Company A", "INTERNAL-REQUEST-FORBIDDEN", "INTERNAL-DIGEST-FORBIDDEN",
                                  "selection", "修改", "删除", "supplier", "request_id"):
                    self.assertNotIn(forbidden, cells)
        self.assertEqual(self.methods[-1], {"methods": ["GET"]})
        self.assertEqual(self.env.saves, [])
        self.assertEqual(self.env.snapshots, {})

    def test_real_xlsx_preserves_zero_html_formula_error_url_and_whitespace_as_text(self):
        values = ("00123", "<b>&amp;HTML</b>", "=1+1", "{=1+1}", "#N/A", "https://example.invalid/literal",
                  "+123", "@SUM(A1)", "\t=1+1", "  leading and trailing  ", "-123")
        row = self.env.records[self.P.PROFILE, "SYNTHETIC-000"]
        row.update(zip(TEXT_FIELDS, values))
        self.export()
        sheet = self.worksheet()
        self.assertEqual(tuple(cell.value for cell in sheet[2]), values)
        self.assertTrue(all(cell.data_type == "s" and cell.hyperlink is None for cell in sheet[2]))
        with ZipFile(BytesIO(self.env.frappe.response["filecontent"])) as archive:
            self.assertFalse(any(name.startswith("xl/externalLinks/") for name in archive.namelist()))
            for name in archive.namelist():
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    tree = ElementTree.fromstring(archive.read(name))
                    self.assertEqual(tree.findall(".//s:f", XML_NS), [])
                    self.assertEqual(tree.findall(".//s:hyperlinks", XML_NS), [])

    def test_oversize_or_illegal_xml_rejects_whole_response_without_partial_binary(self):
        row = self.env.records[self.P.PROFILE, "SYNTHETIC-000"]
        for value in ("x" * 32768, "before\x00after", "before\x0bafter", "before\ud800after", "before\ufffeafter"):
            with self.subTest(value_type="oversize" if len(value) > 32767 else "illegal XML"):
                row["notes"] = value
                with self.assertRaises((ValueError, PermissionError)):
                    self.export()
                self.assert_no_binary()
        row["notes"] = "x" * 32767
        self.export()
        self.assertEqual(self.worksheet().cell(2, 8).value, row["notes"])

    def test_masked_optional_fields_are_blank_without_supplier_fallback(self):
        self.env.records[self.P.PROFILE, "SYNTHETIC-000"]["archive_kind"] = "电镀厂"
        self.env.hidden_fields[self.P.PROFILE, "read"] = {"mobile", "bank_account", "owner", "creation", "modified_by", "modified"}
        self.export(kind="电镀厂")
        sheet = self.worksheet()
        for column in (4, 11, 13, 14, 15, 16):
            self.assertIsNone(sheet.cell(2, column).value)
        self.assertNotIn("SUPPLIER-FALLBACK-FORBIDDEN", str(list(sheet.values)))
        self.env.frappe.response.clear()
        self.env.hidden_fields.clear()
        self.env.metas[self.P.PROFILE].get_field("bank_account").mask = 1
        self.export(kind="电镀厂")
        self.assertIsNone(self.worksheet().cell(2, 11).value)

    def test_irregular_native_values_do_not_get_silently_stringified(self):
        row = self.env.records[self.P.PROFILE, "SYNTHETIC-000"]
        for value in (123, True, float("nan"), {"nested": "not text"}, ["not", "text"]):
            with self.subTest(type=type(value).__name__):
                row["bank_account"] = value
                with self.assertRaises((ValueError, PermissionError)):
                    self.export()
                self.assert_no_binary()

    def test_native_inactive_flags_are_exact_text_and_irregular_flags_reject(self):
        row = self.env.records[self.P.PROFILE, "SYNTHETIC-000"]
        row["archive_kind"] = "电镀厂"
        for value, expected in ((False, "0"), (True, "1"), (0, "0"), (1, "1")):
            with self.subTest(value=value):
                row["inactive"] = value
                self.export(kind="电镀厂", status="全部")
                cell = self.worksheet().cell(2, 12)
                self.assertEqual(cell.value, expected)
                self.assertEqual(cell.data_type, "s")
                self.env.frappe.response.clear()
        for value in ("0", "1", "unexpected flag", 2, -1, {}, float("nan")):
            with self.subTest(type=type(value).__name__):
                self.env.frappe.response.clear()
                row["inactive"] = value
                with self.assertRaises((ValueError, PermissionError)):
                    self.export(kind="电镀厂", status="全部")
                self.assert_no_binary()

    def test_export_permission_checked_before_serialization_and_for_every_profile_after(self):
        self.add_profile(1)
        events = []
        original = self.fixtures.Document.check_permission
        def check(document, permission):
            if document.doctype == self.P.PROFILE and permission == "export":
                events.append((document.name, permission))
            return original(document, permission)
        with patch.object(self.fixtures.Document, "check_permission", new=check), self.during_serialization(lambda: None, events):
            self.export()
        split = events.index(("xlsx", "save"))
        for name in ("SYNTHETIC-000", "SYNTHETIC-001"):
            self.assertIn((name, "export"), events[:split])
            self.assertIn((name, "export"), events[split + 1:])
        self.env.frappe.response.clear()
        def deny_one(document, permission):
            if document.name == "SYNTHETIC-001" and permission == "export":
                raise PermissionError("Synthetic profile lacks native export permission")
            return original(document, permission)
        with patch.object(self.fixtures.Document, "check_permission", new=deny_one), \
             patch.object(Workbook, "save") as save, self.assertRaises(PermissionError):
            self.export()
        save.assert_not_called()
        self.assert_no_binary()
        self.env.denials.add((self.P.PROFILE, "export"))
        with patch.object(Workbook, "save") as save, self.assertRaises(PermissionError):
            self.export()
        save.assert_not_called()
        self.assert_no_binary()

    def test_export_revoked_during_serialization_or_after_final_source_recheck_rejects(self):
        with self.during_serialization(lambda: self.env.denials.add((self.P.PROFILE, "export"))):
            with self.assertRaises(PermissionError):
                self.export()
        self.assert_no_binary()
        self.env.denials.clear()
        original = self.P._recheck_archives
        calls = 0
        def recheck(*args):
            nonlocal calls
            result = original(*args)
            calls += 1
            if calls == 2:
                self.env.denials.add((self.P.PROFILE, "export"))
            return result
        with patch.object(self.P, "_recheck_archives", side_effect=recheck), self.assertRaises(PermissionError):
            self.export()
        self.assertEqual(calls, 2)
        self.assert_no_binary()

    def test_late_source_collection_mask_and_raw_authorization_changes_reject_binary(self):
        self.env.records[self.P.PROFILE, "SYNTHETIC-000"]["archive_kind"] = "电镀厂"
        self.assertIn("modified", self.P._columns(self.P.PROFILE, "read"))
        self.assertTrue(self.env.records[self.P.PROFILE, "SYNTHETIC-000"]["modified"])
        cases = (
            lambda: self.env.records[self.P.PROFILE, "SYNTHETIC-000"].update(notes="changed without modified"),
            lambda: self.env.records["Supplier", "SUP-A"].update(disabled=1),
            lambda: self.add_profile(1, kind="电镀厂"),
            lambda: self.env.hidden_fields.update({(self.P.PROFILE, "read"): {"bank_account"}}),
            lambda: self.env.hidden_fields.update({(self.P.PROFILE, "read"): {"modified"}}),
            lambda: self.auth["DocPerm"][0].update(export=0),
            lambda: self.auth["Has Role"][0].update(role="No Export"),
            lambda: self.auth["Property Setter"][0].update(value="1"),
        )
        for change in cases:
            with self.subTest(change=cases.index(change)):
                records, auth, hidden = deepcopy(self.env.records), deepcopy(self.auth), deepcopy(self.env.hidden_fields)
                with self.during_serialization(change), self.assertRaises((ValueError, PermissionError)):
                    self.export(kind="电镀厂")
                self.assert_no_binary()
                self.env.records, self.auth, self.env.hidden_fields = records, auth, hidden

    def test_supplier_if_owner_permission_gain_rejects_changed_visible_collection(self):
        self.env.records["Supplier", "SUP-HIDDEN"] = dict(self.env.records["Supplier", "SUP-A"],
            name="SUP-HIDDEN", owner="different-synthetic-owner")
        self.add_profile(1, supplier="SUP-HIDDEN")
        permissions_before = deepcopy(self.auth)
        original = self.fixtures.Document.has_permission
        def permission(document, action):
            if document.doctype == "Supplier" and document.name == "SUP-HIDDEN" and action == "read":
                return document.owner == self.env.frappe.session.user
            return original(document, action)
        def gain_owner():
            self.env.records["Supplier", "SUP-HIDDEN"]["owner"] = self.env.frappe.session.user
        with patch.object(self.fixtures.Document, "has_permission", new=permission), \
             patch.object(self.P, "_checked_archives", wraps=self.P._checked_archives) as checked, \
             self.during_serialization(gain_owner), self.assertRaises((ValueError, PermissionError)):
            self.export()
        checked.assert_called_once_with("外发工厂", "Company A", "", "在用")
        self.assertEqual(self.env.records["Supplier", "SUP-HIDDEN"]["owner"], self.env.frappe.session.user)
        self.assertEqual(self.auth, permissions_before)
        self.assert_no_binary()

    def test_checked_collection_once_exports_all_101_rows_without_paging(self):
        for index in range(1, 101):
            self.add_profile(index, legacy_number=f"ZERO-{index:04d}")
        with patch.object(self.P, "_checked_archives", wraps=self.P._checked_archives) as checked, \
             patch.object(self.P, "query_archives", side_effect=AssertionError("must not export by pages")):
            self.export()
        checked.assert_called_once_with("外发工厂", "Company A", "", "在用")
        self.assertEqual(self.worksheet().max_row, 102)
        self.assertEqual(len(list(self.worksheet().iter_rows(min_row=2))), 101)

    def test_native_query_1000_cap_rejects_export_before_serialization(self):
        for index in range(1, 1001):
            self.add_profile(index)
        with patch.object(Workbook, "save") as save, self.assertRaises(ValueError):
            self.export()
        save.assert_not_called()
        self.assert_no_binary()

    def test_invalid_scope_and_guest_rejected_without_binary(self):
        for args in (("外发", "Company A", "", "在用"), ("外发工厂", " Company A", "", "在用"),
                     ("外发工厂", "Company A", "x" * 141, "在用"), ("外发工厂", "Company A", "", "invalid")):
            with self.subTest(args=args), self.assertRaises((ValueError, PermissionError)):
                self.export(*args)
            self.assert_no_binary()
        self.env.frappe.session.user = "Guest"
        with self.assertRaises(PermissionError):
            self.export()
        self.assert_no_binary()


if __name__ == "__main__":
    unittest.main(verbosity=2)
