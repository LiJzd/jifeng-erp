"""Real payroll projections and XLSX bytes under isolated native models only."""

from copy import deepcopy
from io import BytesIO
import importlib.util
import json
from pathlib import Path
import sys
import types
import unittest
from unittest.mock import Mock, patch
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
DEFAULT_COLUMNS = ["说明", "姓名", "部门", "底薪", "应出勤", "实出勤", "出勤工资", "计时", "计件", "合计",
    "房补", "全勤", "应发", "水电", "社保", "罚款", "借款", "实发", "开始", "结束", "类型"]
OPTIONAL_COLUMNS = ["伙食", "交通", "夜宵", "加班", "介绍费", "考核", "其他收入", "其他扣款", "留存工资", "住宿"]


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class Database:
    def __init__(self, model, **settings):
        self.__dict__.update(settings)
        self.begin = Mock(side_effect=lambda **kw: model.event("begin", kw))
        self.rollback = Mock(side_effect=lambda: model.event("rollback"))
        self.close = Mock(side_effect=lambda: model.event("close"))


class TestSalaryHistoryExport(unittest.TestCase):
    def setUp(self):
        self.F = load("salary_export_fixtures", ROOT / "tests/test_salary_history.py")
        self.env = self.F.TestSalaryHistory("test_paid_leave_and_net_pay_never_become_attendance_or_payment")
        self.env.setUp()
        self.addCleanup(self.env.doCleanups)
        self.H = self.env.module
        self.env.rule_module.EARNING_BUCKETS |= {"meal", "transport", "night_snack", "overtime", "introduction", "assessment"}
        self.env.rule_module.DEDUCTION_BUCKETS |= {"retained_pay", "accommodation"}
        package = types.ModuleType("jifeng_erp")
        package.__path__ = [str(ROOT / "jifeng_erp")]
        package.salary_history = self.H
        modules = patch.dict(sys.modules, {"jifeng_erp": package, "jifeng_erp.salary_history": self.H,
            "jifeng_erp.tool_archive": types.ModuleType("jifeng_erp.tool_archive"),
            "jifeng_erp.plate_modeling": types.ModuleType("jifeng_erp.plate_modeling")})
        modules.start()
        self.addCleanup(modules.stop)
        sys.modules["jifeng_erp.tool_archive"].DOCTYPES = {}
        sys.modules["jifeng_erp.plate_modeling"].DOCTYPE = "Jifeng Plate Modeling"
        self.tools = load("salary_export_snapshot", ROOT / "jifeng_erp/tool_archive_export.py")
        helper = patch.dict(sys.modules, {"jifeng_erp.tool_archive_export": self.tools})
        helper.start()
        self.addCleanup(helper.stop)
        self.roles, self.export_grants = {"System Manager"}, set()
        self.can_export = Mock(side_effect=lambda doctype: "System Manager" in self.roles or doctype in self.export_grants)
        self.env.frappe.permissions = self.F.Record(can_export=self.can_export)
        self.settings = {"socket": "/synthetic/payroll.sock", "host": "synthetic-primary", "port": 13316,
            "user": "synthetic-export", "password": "synthetic-only", "cur_db_name": "synthetic-db"}
        self.events, self.failures, self.on_begin = [], set(), None
        self.original = Database(self, **self.settings)
        self.env.frappe.local = self.F.Record(db=self.original, site="synthetic.localhost", conf=self.F.Record(
            db_type="mariadb", db_socket=self.settings["socket"], db_host=self.settings["host"],
            db_port=self.settings["port"], db_user=self.settings["user"], db_password=self.settings["password"],
            db_name=self.settings["cur_db_name"]))
        self.factory = Mock(side_effect=lambda **kw: Database(self, **kw))
        dbmodule = types.ModuleType("frappe.database")
        dbmodule.get_db = self.factory
        patchdb = patch.dict(sys.modules, {"frappe.database": dbmodule})
        patchdb.start()
        self.addCleanup(patchdb.stop)
        model = self

        class Response(self.F.Record):
            def update(response, *args, **kwargs):
                model.assertIs(model.env.frappe.local.db, model.original)
                model.assertEqual(model.events[-2:], [("rollback", None), ("close", None)])
                super().update(*args, **kwargs)

        self.env.frappe.response = Response()
        self.E = load("salary_history_export_under_test", ROOT / "jifeng_erp/salary_history_export.py")

    def event(self, name, options=None):
        self.events.append((name, options))
        if name == "begin" and self.on_begin:
            self.on_begin()
        if name in self.failures:
            raise RuntimeError("synthetic " + name + " failure")

    def export(self, columns=None, context=None):
        self.E.export_salary_history(json.dumps(context if context is not None else self.env.context),
            json.dumps(columns if columns is not None else DEFAULT_COLUMNS))
        return self.env.frappe.response["filecontent"]

    def workbook(self, content):
        book = load_workbook(BytesIO(content), data_only=False)
        self.addCleanup(book.close)
        return book

    def row(self, sheet, index=2):
        return {sheet.cell(1, column).value: sheet.cell(index, column).value for column in range(1, sheet.max_column + 1)}

    def assert_unpublished(self):
        self.assertNotIn("filecontent", self.env.frappe.response)
        self.assertIs(self.env.frappe.local.db, self.original)
        self.env.frappe.db.commit.assert_not_called()
        self.env.frappe.db.sql.assert_not_called()
        self.env.frappe.db.set_value.assert_not_called()

    def test_fixed_31_label_protocol_get_only_and_complete_snapshot_not_public_pages(self):
        self.assertEqual([label for label, _ in self.E.COLUMNS], DEFAULT_COLUMNS + OPTIONAL_COLUMNS)
        self.assertIn(("export_salary_history", {"methods": ["GET"]}), self.env.whitelisted)
        with patch.object(self.H, "query_salary_history", side_effect=AssertionError("no paged API")), \
                patch.object(self.H, "get_salary_history_detail", side_effect=AssertionError("no public detail stitching")):
            self.export()
        self.assertEqual(self.env.frappe.response["filename"], "工资历史.xlsx")
        self.assertEqual(self.env.frappe.response["content_type"], self.tools.CONTENT_TYPE)
        self.factory.assert_called_once_with(**self.settings)
        self.assertEqual(self.events, [("begin", {"read_only": True}), ("rollback", None), ("close", None)])

    def test_native_system_manager_can_export_without_extra_document_or_configuration_export(self):
        for doctype in ("Salary Slip", "Company", "Employee", "Currency", "Salary Component", "Jifeng Salary History Rule"):
            self.env.denials.add((doctype, "export"))
        content = self.export()
        self.assertGreater(len(content), 0)
        self.assertEqual({call.args[0] for call in self.can_export.call_args_list},
            {"Salary Slip", "Additional Salary", "Jifeng Wage Entry"})
        self.roles = {"HR User"}
        self.env.frappe.response.clear()
        with self.assertRaises(PermissionError):
            self.export()
        self.assert_unpublished()

    def test_custom_native_export_grant_does_not_require_contact_or_config_exports(self):
        self.roles, self.export_grants = {"Synthetic Payroll Exporter"}, {"Salary Slip", "Additional Salary", "Jifeng Wage Entry"}
        self.export()
        self.assertFalse(any(call.args[0] in {"Employee", "Company", "Currency", "Salary Component", "Jifeng Salary History Rule"}
            for call in self.can_export.call_args_list))

    def test_actual_bucket_financial_sources_require_native_export_without_omitting_salary_row(self):
        for blocked in ("Additional Salary", "Jifeng Wage Entry"):
            with self.subTest(blocked=blocked):
                self.roles, self.export_grants = set(), {"Salary Slip", "Additional Salary", "Jifeng Wage Entry"} - {blocked}
                with self.assertRaises(PermissionError):
                    self.export(["姓名", "计件", "应发"])
                self.assert_unpublished()
        self.factory.assert_not_called()

    def test_unselected_financial_columns_do_not_require_financial_source_export(self):
        self.roles, self.export_grants = set(), {"Salary Slip"}
        sheet = self.workbook(self.export(["应发", "姓名"]))["工资历史"]
        self.assertEqual(self.row(sheet), {"应发": "1120", "姓名": "Period employee name"})
        self.assertEqual({call.args[0] for call in self.can_export.call_args_list}, {"Salary Slip"})

    def test_unknown_unreadable_sources_keep_whole_salary_and_blank_side_like_query(self):
        for blocked in ("Additional Salary", "Jifeng Wage Entry"):
            with self.subTest(blocked=blocked):
                self.env.denials.add((blocked, "read"))
                self.roles, self.export_grants = set(), {"Salary Slip"}
                self.env.frappe.response.clear()
                sheet = self.workbook(self.export())["工资历史"]
                exported = self.row(sheet)
                self.assertEqual(sheet.max_row, 2)
                self.assertEqual(exported["应发"], "1120")
                for label in ("出勤工资", "计时", "计件", "房补", "全勤"):
                    self.assertIsNone(exported[label])
                self.assertEqual(exported["水电"], "50")
                self.assertIn("未完整核实", sheet.cell(2, 9).comment.text)
                self.env.denials.remove((blocked, "read"))

    def test_missing_rules_or_unreadable_rule_config_do_not_invent_wages(self):
        for rules in ([], None):
            with self.subTest(rules=rules):
                if rules is None:
                    self.env.rule_module.load_rules.side_effect = PermissionError("synthetic hidden rule")
                else:
                    self.env.rule_rows.clear()
                self.roles, self.export_grants = set(), {"Salary Slip"}
                self.env.frappe.response.clear()
                sheet = self.workbook(self.export())["工资历史"]
                row = self.row(sheet)
                self.assertEqual(row["应发"], "1120")
                self.assertIsNone(row["计件"])
                self.assertIsNone(row["出勤工资"])
                self.assertEqual(sheet.max_row, 2)

    def test_mixed_exact_money_unknown_legacy_columns_and_zero_are_literal_public_text(self):
        content = self.export(DEFAULT_COLUMNS + OPTIONAL_COLUMNS)
        book = self.workbook(content)
        sheet = book["工资历史"]
        row = self.row(sheet)
        self.assertEqual(row["出勤工资"], "1000")
        self.assertEqual(row["计时"], "30")
        self.assertEqual(row["计件"], "80")
        self.assertEqual(row["其他收入"], "10")
        self.assertEqual(row["房补"], "0")
        self.assertEqual(row["住宿"], "0")
        for label in ("说明", "底薪", "应出勤", "实出勤", "合计", "实发", "类型"):
            self.assertIsNone(row[label])
        self.assertEqual(row["应发"], "1120")
        self.assertIn("gross_pay", sheet.cell(1, 13).comment.text)
        self.assertIn("CNY", sheet.cell(2, 13).comment.text)
        self.assertIn("net_pay", book["口径说明"].cell(9, 2).value)
        with ZipFile(BytesIO(content)) as archive:
            texts = {node.text for name in archive.namelist() if name.endswith(".xml")
                for node in ElementTree.fromstring(archive.read(name)).iter() if node.text}
        for private in ("PRIVATE-BANK", "PRIVATE-NAME", "Current changed name", "1200", "1070", "RULE-WAGES", "WAGE", "SLIP", "AS"):
            self.assertNotIn(private, texts)

    def test_formula_like_unicode_and_leading_zero_values_are_never_excel_formulas(self):
        values = ["=HYPERLINK(\"https://synthetic.invalid\",\"x\")", "+SUM(1,2)", "-1+2", "@SUM(1,2)",
            "\t=1+2", "\n=1+2", "000姓名", "字\U0001f600"]
        for index, value in enumerate(values):
            slip = self.env.slip("TEXT-" + str(index))
            slip.employee_name = value
        content = self.export(["姓名", "应发"])
        sheet = self.workbook(content)["工资历史"]
        self.assertTrue(set(values).issubset({sheet.cell(index, 1).value for index in range(2, sheet.max_row + 1)}))
        with ZipFile(BytesIO(content)) as archive:
            for name in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
                tree = ElementTree.fromstring(archive.read(name))
                self.assertFalse(tree.findall(".//s:f", NS))
                self.assertTrue(all(cell.attrib["t"] in {"s", "inlineStr"} for cell in tree.findall(".//s:c", NS)))

    def test_decimal_strings_are_not_rounded_through_binary_float(self):
        slip = self.env.records["Salary Slip", "SLIP"]
        slip.gross_pay = "9007199254740993.000000001"
        sheet = self.workbook(self.export(["应发", "出勤工资"]))["工资历史"]
        self.assertEqual(self.row(sheet)["应发"], slip.gross_pay)
        self.assertIsNone(self.row(sheet)["出勤工资"])

    def test_over_200_saved_salary_rows_are_all_exported_in_complete_batches(self):
        for index in range(201):
            self.env.slip("BATCH-" + str(index).zfill(3))
        sheet = self.workbook(self.export(["姓名", "应发"]))["工资历史"]
        self.assertEqual(sheet.max_row, 203)
        batch_calls = [call for call in self.env.calls if call[0] == "list" and call[1] == "Salary Slip"
            and call[2] == {"company": "COMPANY", "docstatus": 1}]
        self.assertEqual([call[4] for call in batch_calls], [0, 200] * 4)
        self.assertTrue(all(call[3] == 200 for call in batch_calls))

    def test_current_company_keyword_and_date_context_use_saved_history_not_current_employee(self):
        extra = self.env.slip("OUTSIDE")
        extra.update(employee_name="Other saved employee", start_date="2026-01-01", end_date="2026-01-31", posting_date="2026-02-01")
        foreign = self.env.slip("FOREIGN")
        foreign.company = "OTHER"
        context = {**self.env.context, "use_date": True, "date_basis": "posting_date", "from_date": "2026-11-01",
            "to_date": "2026-11-01", "search": "Period employee"}
        sheet = self.workbook(self.export(["姓名", "开始", "结束"], context))["工资历史"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(self.row(sheet)["姓名"], "Period employee name")

    def test_source_or_policy_change_at_fresh_snapshot_refuses_all_bytes_and_restores_connection(self):
        mutations = [lambda: self.env.records["Salary Slip", "SLIP"].update(modified=self.F.LATER),
            lambda: self.env.records["Additional Salary", "AS"].update(amount=121),
            lambda: self.env.records["Jifeng Wage Entry", "WAGE"].update(piece_amount=81),
            lambda: self.env.rule_rows[0].update(modified=self.F.LATER),
            lambda: setattr(self.env, "policy", "changed-export-policy")]
        for mutation in mutations:
            with self.subTest(mutation=mutation):
                records, rules, policy = deepcopy(self.env.records), deepcopy(self.env.rule_rows), self.env.policy
                self.on_begin = mutation
                with self.assertRaises((ValueError, PermissionError)):
                    self.export()
                self.assert_unpublished()
                self.assertEqual(self.events[-2:], [("rollback", None), ("close", None)])
                self.env.records, self.env.rule_rows, self.env.policy = records, rules, policy

    def test_financial_export_revocation_while_generating_workbook_refuses_whole_file(self):
        self.roles, self.export_grants = set(), {"Salary Slip", "Additional Salary", "Jifeng Wage Entry"}
        self.on_begin = lambda: self.export_grants.remove("Jifeng Wage Entry")
        with self.assertRaises(PermissionError):
            self.export()
        self.assert_unpublished()

    def test_native_parent_table_child_fields_and_masks_remain_required(self):
        mutations = [lambda: self.env.denials.add(("Salary Slip", "read")),
            lambda: self.env.denials.add(("Employee", "select")),
            lambda: self.env.hidden_fields.update({("Salary Detail", "Salary Slip"): {"amount"}}),
            lambda: self.env.metas["Salary Slip"].get_field("earnings").update(mask=1),
            lambda: self.env.metas["Salary Slip"].get_field("deductions").update(permlevel=1)]
        for mutation in mutations:
            with self.subTest(mutation=mutation):
                denials, fields, metas = self.env.denials.copy(), deepcopy(self.env.hidden_fields), deepcopy(self.env.metas)
                mutation()
                with self.assertRaises(PermissionError):
                    self.export()
                self.assert_unpublished()
                self.env.denials, self.env.hidden_fields, self.env.metas = denials, fields, metas

    def test_masking_financial_source_amount_keeps_unknown_side_empty_not_partly_exported(self):
        self.env.hidden_fields["Additional Salary", None] = {"amount"}
        self.roles, self.export_grants = set(), {"Salary Slip"}
        row = self.row(self.workbook(self.export())["工资历史"])
        self.assertEqual(row["应发"], "1120")
        self.assertIsNone(row["计件"])
        self.assertIsNone(row["出勤工资"])

    def test_unknown_side_source_changes_cannot_replace_saved_gross_or_zero_unknowns(self):
        self.env.denials.add(("Jifeng Wage Entry", "read"))
        self.roles, self.export_grants = set(), {"Salary Slip"}
        self.on_begin = lambda: self.env.records["Jifeng Wage Entry", "WAGE"].update(piece_amount=99999)
        row = self.row(self.workbook(self.export())["工资历史"])
        self.assertEqual(row["应发"], "1120")
        self.assertIsNone(row["计件"])
        self.assertIsNone(row["实发"])

    def test_bank_values_present_absent_and_concurrently_changed_never_enter_file(self):
        initial = self.row(self.workbook(self.export())["工资历史"])
        for doctype, name in (("Salary Slip", "SLIP"), ("Employee", "EMP")):
            self.env.records[doctype, name].pop("bank_account_no", None)
            self.env.records[doctype, name].pop("account_name", None)
        self.env.frappe.response.clear()
        self.on_begin = lambda: self.env.records["Salary Slip", "SLIP"].update(bank_account_no="PRIVATE-CHANGED")
        content = self.export()
        self.assertEqual(self.row(self.workbook(content)["工资历史"]), initial)
        with ZipFile(BytesIO(content)) as archive:
            self.assertFalse(any(b"PRIVATE-CHANGED" in archive.read(name) for name in archive.namelist()))

    def test_invisible_employee_or_company_link_rejects_all_rows_not_a_readable_prefix(self):
        self.env.slip("SECOND")
        for hidden in (("Employee", "EMP"), ("Company", "COMPANY")):
            with self.subTest(hidden=hidden):
                self.env.hidden.add(hidden)
                with self.assertRaises(PermissionError):
                    self.export()
                self.assert_unpublished()
                self.env.hidden.remove(hidden)
        self.factory.assert_not_called()

    def test_excluded_statistical_source_not_in_native_sum_never_requires_money_export(self):
        self.roles, self.export_grants = set(), {"Salary Slip"}
        child = self.env.records["Salary Detail", "SLIP-W"]
        child.update(do_not_include_in_total=1, statistical_component=1)
        slip = self.env.records["Salary Slip", "SLIP"]
        slip.earnings[1].update(do_not_include_in_total=1, statistical_component=1)
        slip.gross_pay, slip.net_pay, slip.rounded_total = 1000, 950, 950
        row = self.row(self.workbook(self.export())["工资历史"])
        self.assertEqual(row["应发"], "1000")
        self.assertEqual(row["计件"], "0")
        self.assertEqual(row["计时"], "0")
        self.assertEqual({call.args[0] for call in self.can_export.call_args_list}, {"Salary Slip"})

    def test_bad_columns_and_context_protocol_are_rejected_before_native_reads(self):
        for value in (None, {}, [], ["银行账号"], ["姓名", "姓名"], [True], ["姓名", None], '["姓名",]'):
            with self.subTest(value=value):
                with self.assertRaises(ValueError):
                    self.E.export_salary_history(self.env.context, value)
                self.assert_unpublished()
        for change in ({"use_date": 1}, {"date_basis": "legacy_guess"}, {"page": 1}, {"search": "=1\x00"}):
            with self.subTest(change=change):
                with self.assertRaises(ValueError):
                    self.export(context={**self.env.context, **change})
        self.factory.assert_not_called()
        self.assertFalse(self.env.calls)

    def test_excel_limits_reject_whole_file_no_prefix_or_utf16_truncation(self):
        self.env.slip("SECOND")
        with patch.object(self.E, "MAX_EXPORT_ROWS", 1):
            with self.assertRaisesRegex(ValueError, "行数限制"):
                self.export()
        self.assert_unpublished()
        for text in ("x" * 32768, "\U0001f600" * 16384, "unexportable\x01", "\ud800"):
            with self.subTest(length=len(text)):
                self.env.records["Salary Slip", "SLIP"].employee_name = text
                with self.assertRaises(ValueError):
                    self.export(["姓名"])
                self.assert_unpublished()
        self.factory.assert_not_called()

    def test_primary_connection_uncertainty_and_cleanup_failure_never_publish(self):
        self.env.frappe.local.conf.db_host = "different-primary"
        with self.assertRaises(ValueError):
            self.export()
        self.assert_unpublished()
        self.factory.assert_not_called()
        self.env.frappe.local.conf.db_host = self.settings["host"]
        for failure in ("begin", "rollback", "close"):
            with self.subTest(failure=failure):
                self.failures = {failure}
                with self.assertRaises(RuntimeError):
                    self.export()
                self.assert_unpublished()


if __name__ == "__main__":
    unittest.main()
