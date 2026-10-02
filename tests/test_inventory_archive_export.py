"""Real inventory DTOs and XLSX serialization under isolated native models."""

from contextlib import contextmanager
from copy import deepcopy
from io import BytesIO
import importlib.util
from pathlib import Path
import sys
import types
import unittest
from unittest.mock import Mock, patch
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook


ROOT = Path(__file__).resolve().parents[1]


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class Database:
    def __init__(self, model, **settings):
        self.__dict__.update(settings)
        self.begin = Mock(side_effect=lambda **kw: model.event("begin", kw))
        self.rollback = Mock(side_effect=lambda: model.event("rollback"))
        self.close = Mock(side_effect=lambda: model.event("close"))


class TestInventoryArchiveExport(unittest.TestCase):
    def setUp(self):
        self.F = load("inventory_export_fixtures", ROOT / "tests/test_inventory_catalog.py")
        self.env = self.F.TestInventoryCatalog("test_pagination_counts_only_exact_readable_sources_and_preserves_saved_values")
        self.env.setUp()
        self.addCleanup(self.env.doCleanups)
        self.C = self.env.module
        docdict = patch.object(self.F.NativeDocument, "as_dict", lambda doc: deepcopy(dict(doc)), create=True)
        docdict.start()
        self.addCleanup(docdict.stop)
        self.env.metas["Supplier"] = {field: self.F.Record(fieldname=field, label=field, fieldtype="Data")
                                      for field in ("name", "modified", "supplier_name")}
        self.env.permitted["Supplier"] = set(self.env.metas["Supplier"])
        self.env.add("Supplier", "SUP", supplier_name=" 000供应商 <b>原文</b> ")
        self.env.file()
        package = types.ModuleType("jifeng_erp")
        package.__path__ = [str(ROOT / "jifeng_erp")]
        package.inventory_catalog = self.C
        modules = patch.dict(sys.modules, {"jifeng_erp": package, "jifeng_erp.inventory_catalog": self.C,
            "jifeng_erp.tool_archive": types.ModuleType("jifeng_erp.tool_archive"),
            "jifeng_erp.plate_modeling": types.ModuleType("jifeng_erp.plate_modeling")})
        modules.start()
        self.addCleanup(modules.stop)
        sys.modules["jifeng_erp.tool_archive"].DOCTYPES = {}
        sys.modules["jifeng_erp.plate_modeling"].DOCTYPE = "Jifeng Plate Modeling"
        self.tools = load("inventory_export_snapshot", ROOT / "jifeng_erp/tool_archive_export.py")
        snapshot = patch.dict(sys.modules, {"jifeng_erp.tool_archive_export": self.tools})
        snapshot.start()
        self.addCleanup(snapshot.stop)
        self.settings = {"socket": "/synthetic/inventory.sock", "host": "synthetic-primary", "port": 13316,
            "user": "synthetic-export", "password": "synthetic-only", "cur_db_name": "synthetic-db"}
        self.events, self.failures, self.old = [], set(), None
        self.original = Database(self, **self.settings)
        self.env.frappe.local = self.F.Record(db=self.original, site="synthetic.localhost", conf=self.F.Record(
            db_type="mariadb", db_socket=self.settings["socket"], db_host=self.settings["host"], db_port=self.settings["port"],
            db_user=self.settings["user"], db_password=self.settings["password"], db_name=self.settings["cur_db_name"]))
        self.factory = Mock(side_effect=lambda **kw: Database(self, **kw))
        dbmodule = types.ModuleType("frappe.database")
        dbmodule.get_db = self.factory
        patchdb = patch.dict(sys.modules, {"frappe.database": dbmodule})
        patchdb.start()
        self.addCleanup(patchdb.stop)
        self.auth = {"User": [{"name": self.env.frappe.session.user, "enabled": 1, "user_type": "System User"}],
            "DocPerm": [{"name": "native-export", "parent": self.C.PROFILE, "read": 1, "export": 1}],
            "Property Setter": [{"name": "field-mask", "doc_type": self.C.PROFILE, "property": "mask", "value": "0"}]}
        self.env.frappe.get_all = Mock(side_effect=self.get_all)
        self.env.frappe.get_doc = Mock(side_effect=self.get_doc)
        self.env.frappe.get_list = Mock(side_effect=self.get_list)
        self.methods = []
        self.env.frappe.whitelist = lambda **kw: self.decorator(kw)
        model = self

        class Response(self.F.Record):
            def update(response, *args, **kw):
                model.assertIs(model.env.frappe.local.db, model.original)
                model.assertEqual(model.events[-2:], [("rollback", None), ("close", None)])
                super().update(*args, **kw)
        self.env.frappe.response = Response()
        self.E = load("inventory_archive_export_under_test", ROOT / "jifeng_erp/inventory_archive_export.py")

    def decorator(self, options):
        self.methods.append(options)
        return lambda fn: fn

    def event(self, name, options=None):
        self.events.append((name, options))
        if name in self.failures:
            raise RuntimeError("synthetic " + name + " failure")

    @contextmanager
    def view(self):
        current = self.env.documents
        if self.old and self.env.frappe.local.db is self.original:
            self.env.documents = self.old["documents"]
        try:
            yield
        finally:
            self.env.documents = current

    def get_doc(self, *args, **kw):
        with self.view():
            return self.env.get_doc(*args, **kw)

    def get_list(self, *args, **kw):
        with self.view():
            return self.env.get_list(*args, **kw)

    def get_all(self, doctype, filters=None, fields=None, **kw):
        with self.view():
            if doctype == "File":
                return self.env.get_all(doctype, filters, **kw)
            auth = self.old["auth"] if self.old and self.env.frappe.local.db is self.original else self.auth
            def matches(row):
                return all(row.get(field) in value[1] if isinstance(value, (list, tuple)) and value[0] == "in"
                           else row.get(field) == value for field, value in (filters or {}).items())
            return [self.F.Record(deepcopy(row) if fields == ["*"] else {field: row.get(field) for field in fields})
                    for row in auth.get(doctype, []) if matches(row)]

    def export(self, kind="物料", company="JF", search="", status="全部"):
        self.env.frappe.response.clear()
        self.E.export_archives(kind, company, search, status)

    def sheet(self):
        workbook = load_workbook(BytesIO(self.env.frappe.response["filecontent"]), data_only=False)
        self.addCleanup(workbook.close)
        return workbook.active

    def no_binary(self):
        self.assertNotIn("filecontent", self.env.frappe.response)
        self.assertIs(self.env.frappe.local.db, self.original)

    def change(self, action, stale=True):
        if stale:
            self.old = {"documents": deepcopy(self.env.documents), "auth": deepcopy(self.auth)}
        original = Workbook.save
        def changed(workbook, stream):
            result = original(workbook, stream)
            action()
            return result
        return patch.object(Workbook, "save", changed)

    def test_verified_20_19_columns_no_internal_ids_urls_or_native_writes(self):
        self.env.add(self.C.PROFILE, "MATERIAL-A", **{key: value for key, value in self.env.profile.items()
            if key not in {"doctype", "name", "modified", "archive_kind"}}, archive_kind="材料")
        before = deepcopy(self.env.documents)
        for kind, count in (("物料", 20), ("材料", 19)):
            self.export(kind)
            sheet = self.sheet()
            self.assertEqual(sheet.max_column, count)
            self.assertEqual(tuple(cell.value for cell in sheet[1]), tuple(label for _, label in self.E.COLUMNS[kind]))
            self.assertEqual(sheet.max_row, 2)
            values = [cell.value for row in sheet for cell in row]
            for value in ("SUP", "ITEM-A", "PROFILE-A", "JF", "/private/files/item.png?fid=EXACT"):
                self.assertNotIn(value, values)
            self.assertIn(" 000供应商 <b>原文</b> ", values)
            self.assertIn("已上传", values)
            self.assertEqual(self.env.frappe.response["filename"], kind + "档案.xlsx")
        self.assertEqual(before, self.env.documents)
        self.assertEqual(self.methods, [{"methods": ["GET"]}])
        self.assertEqual(self.events.count(("begin", {"read_only": True})), 2)
        self.env.frappe.db.commit.assert_not_called()
        self.env.frappe.new_doc.assert_not_called()
        self.env.frappe.delete_doc.assert_not_called()

    def test_literal_text_zeros_formula_html_errors_and_precision(self):
        for value in ("00123", "=1+1", "+001", "#N/A", "<b>&原文</b>", "  文本\t\n "):
            self.env.profile["notes"] = value
            self.export()
            sheet = self.sheet()
            self.assertEqual(sheet.cell(2, 18).value, value)
            self.assertEqual(sheet.cell(2, 10).value, "17.50")
            self.assertTrue(all(cell.data_type in ("s", "inlineStr") and cell.hyperlink is None for row in sheet for cell in row))
            with ZipFile(BytesIO(self.env.frappe.response["filecontent"])) as archive:
                xml = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
                ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                self.assertFalse(xml.findall(".//s:f", ns))
                self.assertFalse(xml.findall(".//s:hyperlinks", ns))
                self.assertFalse(any(path.startswith("xl/externalLinks/") for path in archive.namelist()))

    def test_complete_saved_set_beyond_1000_and_explicit_kind_company_status_search(self):
        self.env.add_archives(1000)
        self.env.add(self.C.PROFILE, "OTHER", company="OTHER", archive_kind="物料", item_code="ITEM-A", inactive=0)
        self.env.add(self.C.PROFILE, "MATERIAL", company="JF", archive_kind="材料", item_code="ITEM-A", inactive=0)
        self.export()
        self.assertEqual(self.sheet().max_row, 1002)
        self.env.profile["inactive"] = 1
        self.export(search="旧名称", status="停用")
        self.assertEqual(self.sheet().max_row, 2)
        self.export(search="旧名称", status="在用")
        self.assertEqual(self.sheet().max_row, 1)
        profile_calls = [call for call in self.env.frappe.get_list.call_args_list if call.args[0] == self.C.PROFILE]
        self.assertTrue(all(call.kwargs["limit_page_length"] == 1048576 for call in profile_calls))

    def test_hidden_or_omitted_projection_never_filled_from_full_document(self):
        self.env.profile["notes"] = "PRIVATE-HIDDEN"
        self.env.hidden[self.C.PROFILE] = {"notes"}
        self.env.projection_omissions[self.C.PROFILE] = {"material"}
        self.env.hidden["Supplier"] = {"supplier_name"}
        self.export()
        sheet = self.sheet()
        self.assertIsNone(sheet.cell(2, 18).value)
        self.assertIsNone(sheet.cell(2, 7).value)
        self.assertIsNone(sheet.cell(2, 17).value)

    def test_read_export_guest_company_supplier_and_empty_export_refuse(self):
        actions = [lambda: self.env.role_denials.add((self.C.PROFILE, "export")),
            lambda: self.env.doc_denials.add((self.C.PROFILE, "PROFILE-A", "export")),
            lambda: self.env.doc_denials.add(("Company", "JF", "select")),
            lambda: self.env.doc_denials.add(("Supplier", "SUP", "read")),
            lambda: self.env.hidden.update({self.C.PROFILE: {"company"}}),
            lambda: setattr(self.env.frappe.session, "user", "Guest")]
        for action in actions:
            action()
            with self.assertRaises(PermissionError):
                self.export()
            self.no_binary()
            self.env.role_denials.clear()
            self.env.doc_denials.clear()
            self.env.hidden.clear()
            self.env.frappe.session.user = "stock-user@example.invalid"
        self.env.role_denials.add((self.C.PROFILE, "export"))
        with self.assertRaises(PermissionError):
            self.export(search="no-result")
        self.no_binary()

    def test_fresh_source_supplier_file_scope_or_permission_changes_reject_entire_file(self):
        actions = [lambda: self.env.profile.update(notes="committed without modified"),
            lambda: self.env.documents["Supplier", "SUP"].update(supplier_name="changed"),
            lambda: self.env.documents["File", "EXACT"].update(attached_to_field="other"),
            lambda: self.env.add_archives(1),
            lambda: self.auth["Property Setter"][0].update(value="1"),
            lambda: self.env.doc_denials.add((self.C.PROFILE, "PROFILE-A", "export")),
            lambda: self.env.doc_denials.add(("Item", "ITEM-A", "read"))]
        for action in actions:
            saved_documents, saved_auth = deepcopy(self.env.documents), deepcopy(self.auth)
            with self.change(action), self.assertRaises((ValueError, PermissionError)):
                self.export()
            self.no_binary()
            self.env.documents, self.auth, self.old = saved_documents, saved_auth, None
            self.env.profile = self.env.documents[self.C.PROFILE, "PROFILE-A"]
            self.env.doc_denials.clear()

    def test_invalid_values_and_excel_limits_are_not_silently_truncated(self):
        for field, value in (("notes", "x" * 32768), ("notes", "😀" * 16384), ("notes", "x\x00y"),
                ("notes", "x\ud800y"), ("notes", "x\ufffey"), ("notes", 123), ("inactive", 2),
                ("weight", float("nan")), ("weight", "Infinity")):
            old = self.env.profile.get(field)
            self.env.profile[field] = value
            with self.assertRaises((ValueError, PermissionError)):
                self.export()
            self.no_binary()
            self.env.profile[field] = old
        with patch.object(self.E, "MAX_EXPORT_ROWS", 0), self.assertRaisesRegex(ValueError, "Excel行数限制"):
            self.export()
        self.no_binary()

    def test_invalid_scope_and_fresh_cleanup_failures_never_publish(self):
        for args in (("产品", "JF", "", "全部"), ("物料", " JF", "", "全部"), ("物料", "JF", [], "全部"),
                     ("物料", "JF", "x" * 141, "全部"), ("物料", "JF", "", "任意")):
            with self.assertRaises((ValueError, PermissionError)):
                self.export(*args)
            self.no_binary()
        for failure in ("begin", "rollback", "close"):
            self.failures.add(failure)
            with self.assertRaises(RuntimeError):
                self.export()
            self.no_binary()
            self.failures.clear()


if __name__ == "__main__":
    unittest.main()
