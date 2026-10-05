"""Bounded files and batch native-save/receipt checks without a live database."""

import base64
from copy import deepcopy
import importlib.util
from io import BytesIO
import json
from pathlib import Path
import re
import sys
import unittest
from unittest.mock import Mock, patch
from uuid import UUID, uuid4
import zipfile

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).parents[1]
spec = importlib.util.spec_from_file_location("archive_import_native_model", ROOT / "tests/test_inventory_archive.py")
MODEL = importlib.util.module_from_spec(spec)
spec.loader.exec_module(MODEL)


def xlsx(headers, rows, mutate=None):
    book = Workbook()
    book.active.append(headers)
    for row in rows:
        book.active.append(row)
    if mutate:
        mutate(book)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class TestInventoryArchiveImport(unittest.TestCase):
    def setUp(self):
        self.native = MODEL.TestInventoryArchive(methodName="runTest")
        self.native.setUp()
        self.frappe = self.native.frappe
        self.frappe.response = {}
        self.frappe.ValidationError = ValueError
        self.frappe.DoesNotExistError = type("DoesNotExistError", (ValueError,), {})
        original_get = self.native.get_doc

        def get_doc(*args, **kwargs):
            try:
                return original_get(*args, **kwargs)
            except KeyError:
                raise self.frappe.DoesNotExistError("private native missing ID")

        self.frappe.get_doc.side_effect = get_doc
        with patch.dict(sys.modules, {"frappe": self.frappe, "jifeng_erp.inventory_archive": self.native.module}):
            spec = importlib.util.spec_from_file_location("archive_import_under_test", ROOT / "jifeng_erp/inventory_archive_import.py")
            self.module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(self.module)
        self.batch_id = str(uuid4())

    def payload(self, raw=None, **changes):
        raw = raw or xlsx(["编号", "名称", "单位", "单价", "备注"],
                           [[" 001 ", " 中文一 ", "原文斤", 0, "=原文备注"], ["002", "中文二", "原文包", 1.25, "测试"]],
                           lambda book: setattr(book.active["E2"], "data_type", "s"))
        return {"kind": "物料", "company": "JF", "filename": "档案.xlsx", "content_base64": base64.b64encode(raw).decode(),
                "item_group": "Raw Material", "stock_uom": "Nos", "batch_id": self.batch_id, **changes}

    def preview(self, **changes):
        return self.module.preview_import(**self.payload(**changes))

    def resolve(self, preview, **changes):
        metadata = preview["metadata"]
        return self.module.resolve_import(**{key: metadata[key] for key in ("kind", "company", "batch_id", "row_count", "row_digests")},
                                          preview_digest=preview["preview_digest"], **changes)

    def test_template_and_preview_are_readonly_and_preserve_text_and_separate_units(self):
        before = deepcopy(self.native.documents)
        init = self.module.init_import("物料", "JF")
        fields = {column["fieldname"] for column in init["columns"]}
        self.assertNotIn("shape", fields)
        self.assertIn("material", fields)
        self.assertNotIn("item_code", fields)
        self.assertTrue(init["can_create_item"] and init["can_select_item_group"] and init["can_select_stock_uom"])
        self.module.download_template("物料", "JF")
        book = load_workbook(BytesIO(self.frappe.response["filecontent"]))
        self.assertEqual(book.sheetnames, ["物料档案", "导入说明"])
        self.assertEqual(book.active["A2"].number_format, "@")
        self.assertNotIn("图片", [cell.value for cell in book.active[1]])
        book.close()
        preview = self.preview()
        self.assertTrue(preview["can_commit"])
        self.assertEqual(preview["rows"][0]["profile"]["legacy_number"], " 001 ")
        self.assertEqual(preview["rows"][0]["profile"]["display_name"], " 中文一 ")
        self.assertEqual(preview["rows"][0]["profile"]["notes"], "=原文备注")
        self.assertEqual(preview["metadata"]["row_count"], 2)
        self.assertEqual(set(preview["metadata"]), {"version", "kind", "company", "batch_id", "row_count", "row_digests"})
        self.assertEqual(self.native.documents, before)
        self.assertEqual(self.native.saves, [])
        self.assertEqual(self.resolve(preview)["status"], "not_found")

    def test_batch_creation_reuse_and_metadata_only_resolution(self):
        payload = self.payload()
        preview = self.module.preview_import(**payload)
        result = self.module.commit_import(**payload, preview_digest=preview["preview_digest"])
        self.assertEqual(result["status"], "found")
        self.assertEqual(result["metadata"], preview["metadata"])
        for index, receipt in enumerate(result["receipts"]):
            self.assertEqual(receipt["request_id"], preview["rows"][index]["request_id"])
            self.assertEqual(UUID(receipt["request_id"]).version, 4)
            item = self.native.documents[("Item", receipt["item_code"])]
            profile = self.native.documents[(self.native.module.PROFILE, receipt["name"])]
            self.assertEqual(item["stock_uom"], "Nos")
            self.assertEqual(profile["display_uom"], ("原文斤", "原文包")[index])
            self.assertEqual(profile["request_digest"], preview["metadata"]["row_digests"][index])
            self.assertNotIn("profile", receipt)
        before = deepcopy(self.native.documents)
        self.assertEqual(self.module.commit_import(**payload, preview_digest=preview["preview_digest"]), result)
        self.assertEqual(self.native.documents, before)
        self.assertEqual(self.resolve(preview), result)
        self.native.hidden[(self.native.module.PROFILE, "read")] = {"item_code"}
        self.assertTrue(all(row["item_code"] is None for row in self.resolve(preview)["receipts"]))

    def test_commit_rechecks_digest_defaults_contents_and_current_permissions_before_writes(self):
        payload = self.payload()
        preview = self.module.preview_import(**payload)
        before = deepcopy(self.native.documents)
        changed = xlsx(["编号", "名称"], [["001", "不同内容"], ["002", "中文二"]])
        for change in ({"stock_uom": "Kg"}, {"content_base64": base64.b64encode(changed).decode()}, {"batch_id": str(uuid4())}):
            with self.assertRaisesRegex(ValueError, "重新预览"):
                self.module.commit_import(**{**payload, **change}, preview_digest=preview["preview_digest"])
            self.assertEqual(self.native.documents, before)
        for denial in (("Item", "create"), ("Company", "select"), ("Item Group", "select")):
            self.native.denials.add(denial)
            with self.assertRaises((ValueError, PermissionError)):
                self.module.commit_import(**payload, preview_digest=preview["preview_digest"])
            self.native.denials.remove(denial)
            self.assertEqual(self.native.documents, before)

    def test_middle_native_failure_rolls_back_entire_batch_and_callbacks(self):
        payload = self.payload()
        preview = self.module.preview_import(**payload)
        before = deepcopy(self.native.documents)
        previous = Mock()
        self.frappe.db.after_commit.add(previous)
        self.frappe.local._realtime_log = ["previous"]
        self.frappe.flags.currently_saving = ["previous"]
        save = self.native.module.save_archive
        calls = []

        def fail_second(values):
            calls.append(values)
            self.native.fail_profile = len(calls) == 2
            return save(values)

        with patch.object(self.native.module, "save_archive", side_effect=fail_second):
            with self.assertRaisesRegex(ValueError, "profile native failure"):
                self.module.commit_import(**payload, preview_digest=preview["preview_digest"])
        self.assertEqual(len(calls), 2)
        self.assertEqual(self.native.documents, before)
        self.assertEqual(self.frappe.db.after_commit._functions, [previous])
        self.assertEqual(self.frappe.local._realtime_log, ["previous"])
        self.assertEqual(self.frappe.flags.currently_saving, ["previous"])
        self.assertEqual(self.resolve(preview)["status"], "not_found")

    def test_all_row_errors_are_visible_and_prevent_whole_batch_commit(self):
        raw = xlsx(["编号", "名称", "单价", "上限", "下限", "换算值", "供应商"],
                   [[123, "名称", 0, None, None, 1, None], ["002", "名称", -1, None, None, 1, None],
                    ["003", "名称", 0, 1, 2, 1, None], ["004", "名称", 0, None, None, 0, None],
                    ["005", "名称", 0, None, None, 1, "显示名不存在"]])
        preview = self.preview(raw=raw)
        self.assertFalse(preview["can_commit"])
        self.assertIsNone(preview["preview_digest"])
        self.assertIsNone(preview["metadata"])
        self.assertEqual(preview["rows"][0]["profile"]["legacy_number"], 123)
        self.assertEqual(preview["rows"][1]["profile"]["list_price"], -1)
        self.assertEqual(preview["rows"][4]["profile"]["supplier"], "显示名不存在")
        self.assertEqual([error["sheet_row"] for error in preview["errors"]], [2, 3, 4, 5, 6])
        self.assertNotIn("private native", str(preview))
        before = deepcopy(self.native.documents)
        with self.assertRaisesRegex(ValueError, "修正"):
            self.module.commit_import(**self.payload(raw=raw), preview_digest="a" * 64)
        self.assertEqual(self.native.documents, before)

    def test_100_rows_limit_and_blank_rows_do_not_silently_truncate(self):
        data = [[str(index).zfill(3), "中文"] for index in range(101)]
        preview = self.preview(raw=xlsx(["编号", "名称"], data[:100]))
        self.assertEqual(preview["row_count"], 100)
        with self.assertRaisesRegex(ValueError, "分文件"):
            self.preview(raw=xlsx(["编号", "名称"], data))
        preview = self.preview(raw=xlsx(["编号", "名称"], [[None, None], ["001", "中文"]]))
        self.assertEqual(preview["rows"][0]["sheet_row"], 3)
        self.assertEqual(preview["rows"][0]["request_id"], self.module._row_id(self.batch_id, 0))
        # A stale dimension must not hide rows that are physically in the sheet.
        raw = xlsx(["编号", "名称"], data[:2])
        output = BytesIO()
        with zipfile.ZipFile(BytesIO(raw)) as archive, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as corrected:
            for info in archive.infolist():
                content = archive.read(info)
                if info.filename == "xl/worksheets/sheet1.xml":
                    content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:B2"', content)
                corrected.writestr(info.filename, content)
        self.assertEqual(self.preview(raw=output.getvalue())["row_count"], 2)

    def test_rejects_unknown_picture_columns_formulas_links_and_extra_or_hidden_sheets(self):
        examples = [xlsx(["编号", "名称", "图片"], [["001", "中文", "file.jpg"]]),
                    xlsx(["编号", "名称", "余额"], [["001", "中文", 123]]),
                    xlsx(["编号", "名称"], [["001", "=1+1"]]),
                    xlsx(["编号", "名称"], [["001", "中文"]], lambda book: setattr(book.active["B2"], "hyperlink", "https://example.invalid")),
                    xlsx(["编号", "名称"], [["001", "中文"]], lambda book: book.create_sheet("其他数据")),
                    xlsx(["编号", "名称"], [["001", "中文"]], lambda book: setattr(book.create_sheet("隐藏数据"), "sheet_state", "hidden")),
                    xlsx(["编号", "名称"], [["001", "中文"]], lambda book: book.active.merge_cells("B2:C2"))]
        for raw in examples:
            with self.subTest(raw_size=len(raw)), self.assertRaises(ValueError):
                self.preview(raw=raw)
        self.assertEqual(self.native.saves, [])

    def test_zip_bomb_duplicate_members_and_nonxlsx_are_rejected(self):
        buffers = []
        original = xlsx(["编号", "名称"], [["001", "中文"]])
        for members in ((('xl/vbaProject.bin', b'x'),), (('xl/oversize.xml', b'x' * (2 * 1024 * 1024 + 1)),),
                        (('same.xml', b'<x/>'), ('same.xml', b'<x/>'))):
            output = BytesIO()
            with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
                with zipfile.ZipFile(BytesIO(original)) as source:
                    for info in source.infolist():
                        archive.writestr(info.filename, source.read(info))
                for name, content in members:
                    archive.writestr(name, content)
            buffers.append(output.getvalue())
        for raw in buffers + [b'not a zip']:
            with self.assertRaises(ValueError):
                self.preview(raw=raw)
        with self.assertRaises(ValueError):
            self.preview(filename="档案.xlsm")
        with self.assertRaises(ValueError):
            self.preview(content_base64="A" * (2 * 1024 * 1024))

    def test_receipt_resolution_is_readonly_validates_all_metadata_and_current_access(self):
        payload = self.payload()
        preview = self.module.preview_import(**payload)
        row = preview["rows"][0]
        self.native.module.save_archive({"kind": "物料", "request_id": row["request_id"], "profile": {"company": "JF", **row["profile"]},
                                         "item": {"item_group": "Raw Material", "stock_uom": "Nos"}})
        before = deepcopy(self.native.documents)
        self.assertEqual(self.resolve(preview)["status"], "partial")
        metadata = preview["metadata"]
        args = {key: metadata[key] for key in ("kind", "company", "batch_id", "row_count", "row_digests")}
        args["preview_digest"] = preview["preview_digest"]
        self.assertEqual(self.module.resolve_import(**{**args, "row_count": "2", "row_digests": json.dumps(args["row_digests"])} )["status"], "partial")
        for change in ({"row_digests": ["b" * 64, "c" * 64]}, {"company": "OTHER"}, {"batch_id": str(uuid4())}, {"row_count": 1}):
            with self.assertRaises(ValueError):
                self.module.resolve_import(**{**args, **change})
        for denial in ((self.native.module.PROFILE, "read"), ("Item", "read"), ("Company", "select")):
            self.native.denials.add(denial)
            with self.assertRaises(PermissionError):
                self.resolve(preview)
            self.native.denials.remove(denial)
        self.assertEqual(self.native.documents, before)

    def test_guest_and_read_only_role_cannot_preview_commit_or_download(self):
        preview = self.preview()
        for guest in (False, True):
            if guest:
                self.frappe.session.user = "Guest"
            else:
                self.native.denials.add((self.native.module.PROFILE, "create"))
            for action in (lambda: self.preview(), lambda: self.module.init_import("物料", "JF"),
                           lambda: self.module.download_template("物料", "JF"),
                           lambda: self.module.commit_import(**self.payload(), preview_digest=preview["preview_digest"])):
                with self.assertRaises(PermissionError):
                    action()
        self.assertEqual(self.native.saves, [])


if __name__ == "__main__":
    unittest.main()
