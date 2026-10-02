"""Actual tool DTOs and real XLSX packages under isolated permission models."""

from contextlib import contextmanager
from copy import deepcopy
from io import BytesIO
import importlib.util
import json
from pathlib import Path
import sys
import types
import unittest
from unittest.mock import Mock, patch
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook


ROOT = Path(__file__).resolve().parents[1]
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class Database:
    def __init__(self, model, **settings):
        self.__dict__.update(settings)
        self.begin = Mock(side_effect=lambda **kw: model.event("begin", kw))
        self.rollback = Mock(side_effect=lambda: model.event("rollback"))
        self.close = Mock(side_effect=lambda: model.event("close"))


class TestToolArchiveExport(unittest.TestCase):
    def setUp(self):
        fixtures = load("tool_export_fixtures", ROOT / "tests/test_tool_archive.py")
        self.fixtures = fixtures
        self.env = fixtures.TestToolArchive("test_metadata_preserves_existing_roles_business_identity_and_original_order")
        self.env.setUp()
        self.addCleanup(self.env.doCleanups)
        self.T = self.env.module
        schema = json.loads((ROOT / "jifeng_erp/jifeng_erp/doctype/jifeng_plate_modeling/jifeng_plate_modeling.json").read_text(encoding="utf-8"))
        dt = "Jifeng Plate Modeling"
        self.env.fields[dt] = {row["fieldname"]: fixtures.Record(row) for row in schema["fields"]}
        self.env.metadata[dt] = types.SimpleNamespace(get_field=lambda field: self.env.fields[dt].get(field))
        self.env.fields["Jifeng Expense Voucher"] = {"name": fixtures.Record(fieldtype="Data")}
        self.env.metadata["Jifeng Expense Voucher"] = types.SimpleNamespace(get_field=lambda field: self.env.fields["Jifeng Expense Voucher"].get(field))
        self.M = load("modeling_export_dto", ROOT / "jifeng_erp/plate_modeling.py")
        docdict = patch.object(fixtures.Document, "as_dict", lambda doc: deepcopy(dict(doc)), create=True)
        docdict.start()
        self.addCleanup(docdict.stop)
        package = types.ModuleType("jifeng_erp")
        package.__path__ = [str(ROOT / "jifeng_erp")]
        package.tool_archive, package.plate_modeling = self.T, self.M
        modules = patch.dict(sys.modules, {"jifeng_erp": package, "jifeng_erp.tool_archive": self.T, "jifeng_erp.plate_modeling": self.M})
        modules.start()
        self.addCleanup(modules.stop)
        self.settings = {"socket": "/synthetic/export.sock", "host": "synthetic-primary", "port": 13316,
                         "user": "synthetic-export", "password": "synthetic-only", "cur_db_name": "synthetic-db"}
        self.events, self.failures, self.old = [], set(), None
        self.original = Database(self, **self.settings)
        self.env.frappe.local.db, self.env.frappe.local.site = self.original, "synthetic.localhost"
        self.env.frappe.local.conf = fixtures.Record(db_type="mariadb", db_socket=self.settings["socket"],
            db_host=self.settings["host"], db_port=self.settings["port"], db_user=self.settings["user"],
            db_password=self.settings["password"], db_name=self.settings["cur_db_name"])
        self.databases = []
        self.factory = Mock(side_effect=self.database)
        dbmodule = types.ModuleType("frappe.database")
        dbmodule.get_db = self.factory
        patchdb = patch.dict(sys.modules, {"frappe.database": dbmodule})
        patchdb.start()
        self.addCleanup(patchdb.stop)
        self.auth = {"User": [{"name": self.env.frappe.session.user, "enabled": 1, "user_type": "System User"}],
            "DocPerm": [{"name": "native-export", "parent": "Jifeng Mold", "read": 1, "export": 1}],
            "Property Setter": [{"name": "field-mask", "doc_type": "Jifeng Mold", "property": "mask", "value": "0"}]}
        self.env.frappe.get_all = Mock(side_effect=self.get_all)
        self.env.frappe.get_doc = Mock(side_effect=self.get_doc)
        self.env.frappe.get_list = Mock(side_effect=self.get_list)
        self.methods = []
        self.env.frappe.whitelist = lambda **kw: self.decorator(kw)
        model = self
        class Response(fixtures.Record):
            def update(response, *args, **kw):
                model.assertIs(model.env.frappe.local.db, model.original)
                if "filecontent" in dict(*args, **kw):
                    model.assertEqual(model.events[-2:], [("rollback", None), ("close", None)])
                super().update(*args, **kw)
        self.env.frappe.response = Response()
        self.E = load("tool_archive_export_under_test", ROOT / "jifeng_erp/tool_archive_export.py")
        for kind in self.E.DOCTYPES:
            self.add(kind)

    def decorator(self, options):
        self.methods.append(options)
        return lambda fn: fn

    def event(self, name, options=None):
        self.events.append((name, options))
        if name in self.failures:
            raise RuntimeError("synthetic " + name + " failure")

    def database(self, **kw):
        self.assertEqual(kw, self.settings)
        database = Database(self, **kw)
        self.databases.append(database)
        return database

    @contextmanager
    def view(self):
        current = self.env.documents
        if self.old is not None and self.env.frappe.local.db is self.original:
            self.env.documents = self.old["documents"]
        try:
            yield
        finally:
            self.env.documents = current

    def get_doc(self, *args, **kw):
        with self.view():
            return self.env.get_doc(*args, **kw)

    def get_list(self, *args, **kw):
        with self.view():
            return self.env.get_list(*args, **kw)

    def get_all(self, doctype, filters=None, fields=None, **kw):
        def matches(row):
            return all(row.get(field) in value[1] if isinstance(value, list) and value[0] == "in" else row.get(field) == value
                       for field, value in (filters or {}).items())
        auth = self.old["auth"] if self.old and self.env.frappe.local.db is self.original else self.auth
        with self.view():
            if doctype in self.E.DOCTYPES.values():
                rows = [row for row in self.env.documents.values() if row["doctype"] == doctype and matches(row)]
                if kw.get("or_filters"):
                    rows = [row for row in rows if any(value.strip("%") in str(row.get(field) or "") for field, _, value in kw["or_filters"])]
                rows = sorted(rows, key=lambda row: row["name"])
            elif doctype == "File":
                rows = [row for row in self.env.documents.values() if row["doctype"] == doctype and matches(row)]
            else:
                rows = [row for row in auth.get(doctype, []) if matches(row)]
            if kw.get("pluck"):
                return [row[kw["pluck"]] for row in rows]
            return [self.fixtures.Record(deepcopy(row) if fields == ["*"] else {field: row.get(field) for field in fields}) for row in rows]

    def add(self, kind, index=0, **values):
        name = "SYNTHETIC-" + kind.upper() + f"-{index:03d}"
        row = {"doctype": self.E.DOCTYPES[kind], "name": name, "company": "C-1", "__islocal": 0,
               "modified": "2026-10-01T12:00:00.123456", "legacy_number": "00123", "product_code": "原产品",
               "request_id": "INTERNAL-REQUEST-FORBIDDEN", "request_digest": "INTERNAL-DIGEST-FORBIDDEN",
               "maker": None, "modeler": None, "legacy_copy": 0, "image": None, "copy_from": None, "item_code": None,
               "mold_no": name + "-BUSINESS", "plate_no": name + "-BUSINESS", "quantity": 2.5 if kind == "mold" else "0001.0000000000000007",
               "amount": 12.25, "mold_date": "2026-10-01", "scrapped": 0, "retired": 0, "assessment_score": "-0",
               "model_number": name + "-BUSINESS", "model_date": "2026-10-01", "reference_unit_price": "0002.000000000000001",
               **values}
        self.env.documents[name] = self.fixtures.Record(row)
        return self.env.documents[name]

    def export(self, kind="mold", company="C-1", search="", include_inactive=0):
        self.env.frappe.response.clear()
        return self.E.export_archives(kind, company, search, include_inactive)

    def sheet(self):
        workbook = load_workbook(BytesIO(self.env.frappe.response["filecontent"]), data_only=False)
        self.addCleanup(workbook.close)
        return workbook.active

    def no_binary(self):
        self.assertNotIn("filecontent", self.env.frappe.response)
        self.assertNotIn(self.env.frappe.response.get("type"), ("download", "binary"))
        self.assertIs(self.env.frappe.local.db, self.original)

    def serialization_change(self, action, stale=False):
        if stale:
            self.old = {"documents": deepcopy(self.env.documents), "auth": deepcopy(self.auth)}
        original = Workbook.save
        def changed(workbook, stream):
            result = original(workbook, stream)
            action()
            return result
        return patch.object(Workbook, "save", changed)

    def test_exact_current_verified_columns_all_kinds_no_internal_source_no_writes(self):
        before = deepcopy(self.env.documents)
        for kind in self.E.DOCTYPES:
            self.export(kind)
            sheet = self.sheet()
            self.assertEqual(tuple(cell.value for cell in sheet[1]), tuple(label for _, label in self.E.COLUMNS[kind]))
            self.assertEqual(sheet.max_row, 2)
            self.assertEqual(self.env.frappe.response["filename"], self.E.TITLES[kind] + "档案.xlsx")
            self.assertEqual(self.env.frappe.response["content_type"], self.E.CONTENT_TYPE)
            self.assertEqual(self.env.frappe.response["display_content_as"], "attachment")
            self.assertEqual(self.env.frappe.response["type"], "download")
            for value in ("C-1", "INTERNAL-REQUEST-FORBIDDEN", "INTERNAL-DIGEST-FORBIDDEN", "ITEM"):
                self.assertNotIn(value, [cell.value for row in sheet for cell in row])
        self.assertEqual(before, self.env.documents)
        self.assertEqual(self.methods, [{"methods": ["GET"]}])
        self.assertEqual(self.events.count(("begin", {"read_only": True})), 3)

    def test_modeling_export_declaration_only_existing_three_managers(self):
        schema = json.loads((ROOT / "jifeng_erp/jifeng_erp/doctype/jifeng_plate_modeling/jifeng_plate_modeling.json").read_text(encoding="utf-8"))
        self.assertEqual({row["role"] for row in schema["permissions"]},
                         {"System Manager", "Stock Manager", "Manufacturing Manager"})
        self.assertTrue(all(row["export"] == 1 for row in schema["permissions"]))
        self.assertTrue(all(not row.get("delete") for row in schema["permissions"]))

    def test_literal_zero_html_formula_error_precision_and_whitespace(self):
        for kind, text_field in (("mold", "remark"), ("plate", "notes"), ("modeling", "remarks")):
            for value in ("00123", "<b>&amp;字面HTML</b>", "=1+1", "#N/A", "https://example.invalid/raw", "  +001\t\n  "):
                self.env.documents["SYNTHETIC-" + kind.upper() + "-000"][text_field] = value
                self.export(kind)
                sheet = self.sheet()
                column = [field for field, _ in self.E.COLUMNS[kind]].index(text_field) + 1
                self.assertEqual(sheet.cell(2, column).value, value)
                self.assertTrue(all(cell.data_type in ("s", "inlineStr") and cell.hyperlink is None for row in sheet for cell in row))
                with ZipFile(BytesIO(self.env.frappe.response["filecontent"])) as archive:
                    xml = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
                    self.assertEqual(xml.findall(".//s:f", NS), [])
                    self.assertEqual(xml.findall(".//s:hyperlinks", NS), [])
            if kind == "plate":
                self.assertEqual(sheet.cell(2, 6).value, "0001.0000000000000007")
            if kind == "modeling":
                self.assertEqual(sheet.cell(2, 4).value, "0002.000000000000001")

    def test_complete_filter_set_beyond_frontend_100_company_kind_search_and_inactive(self):
        for index in range(1, 106):
            self.add("plate", index, notes="=literal", product_code="match")
        self.add("plate", 201, company="C-2", product_code="match")
        self.add("plate", 202, product_code="match", retired=1)
        self.add("mold", 203, product_code="match")
        self.export("plate", search="match")
        self.assertEqual(self.sheet().max_row, 106)
        self.export("plate", search="match", include_inactive=1)
        self.assertEqual(self.sheet().max_row, 107)
        candidate_calls = [call for call in self.env.frappe.get_all.call_args_list if call.args[0] == "Jifeng Plate"]
        self.assertTrue(all(call.kwargs["limit_page_length"] == 0 for call in candidate_calls))

    def test_unreadable_optional_fields_export_blank_without_raw_fallback(self):
        self.env.documents["SYNTHETIC-MOLD-000"]["remark"] = "PRIVATE-FORBIDDEN"
        self.env.hidden["Jifeng Mold", "read"] = {"remark"}
        self.export()
        self.assertIsNone(self.sheet().cell(2, 8).value)

    def test_native_read_export_company_and_guest_rejections_publish_nothing(self):
        actions = [lambda: self.env.denials.add(("Jifeng Mold", "read")),
                   lambda: self.env.denials.add(("SYNTHETIC-MOLD-000", "export")),
                   lambda: self.env.denied_companies.add("C-1"),
                   lambda: self.env.hidden.update({("Jifeng Mold", "read"): {"company"}}),
                   lambda: setattr(self.env.frappe.session, "user", "Guest")]
        for action in actions:
            with self.subTest(action=action):
                action()
                with self.assertRaises(PermissionError):
                    self.export()
                self.no_binary()
                self.env.denials.clear()
                self.env.denied_companies.clear()
                self.env.hidden.clear()
                self.env.frappe.session.user = "stock@example.invalid"

    def test_empty_export_still_requires_native_export_permission(self):
        self.env.denials.add(("Jifeng Mold", "export"))
        with self.assertRaises(PermissionError):
            self.export(search="not-found")
        self.no_binary()

    def test_unreadable_record_omitted_but_gain_or_loss_during_serialization_refuses(self):
        self.env.denials.add(("SYNTHETIC-MOLD-000", "read"))
        self.export()
        self.assertEqual(self.sheet().max_row, 1)
        with self.serialization_change(lambda: self.env.denials.clear()):
            with self.assertRaises(ValueError):
                self.export()
        self.no_binary()
        with self.serialization_change(lambda: self.env.denials.add(("SYNTHETIC-MOLD-000", "read"))):
            with self.assertRaises(ValueError):
                self.export()
        self.no_binary()

    def test_fresh_committed_source_change_despite_stale_original_snapshot(self):
        with self.serialization_change(lambda: self.env.documents["SYNTHETIC-MOLD-000"].update(remark="changed without modified"), stale=True):
            with self.assertRaises(ValueError):
                self.export()
        self.no_binary()

    def test_fresh_candidate_addition_refuses_partial_result(self):
        with self.serialization_change(lambda: self.add("mold", 1), stale=True):
            with self.assertRaises(ValueError):
                self.export()
        self.no_binary()

    def test_permission_metadata_change_refuses_even_if_cached_native_checks_allow(self):
        with self.serialization_change(lambda: self.auth["Property Setter"][0].update(value="1"), stale=True):
            with self.assertRaises(PermissionError):
                self.export()
        self.no_binary()

    def test_native_export_revocation_after_serialization_refuses(self):
        with self.serialization_change(lambda: self.env.denials.add(("SYNTHETIC-MOLD-000", "export"))):
            with self.assertRaises(PermissionError):
                self.export()
        self.no_binary()

    def test_illegal_oversize_nontext_invalid_check_nan_refuse_entire_file(self):
        row = self.env.documents["SYNTHETIC-MOLD-000"]
        for field, value in (("remark", "x" * 32768), ("remark", "x\x00y"), ("remark", "x\ud800y"),
                             ("remark", "x\ufffey"), ("remark", 123), ("scrapped", 2), ("amount", float("nan"))):
            with self.subTest(field=field):
                original = row.get(field)
                row[field] = value
                with self.assertRaises(ValueError):
                    self.export(include_inactive=1)
                self.no_binary()
                row[field] = original

    def test_fresh_cleanup_failures_refuse_response_and_restore_original_db(self):
        for failure in ("begin", "rollback", "close"):
            self.failures.add(failure)
            with self.assertRaises(RuntimeError):
                self.export()
            self.no_binary()
            self.failures.clear()

    def test_unverified_connection_or_same_connection_refuses(self):
        self.original.host = "unknown-host"
        with self.assertRaises(ValueError):
            self.export()
        self.no_binary()
        self.assertEqual(self.factory.call_count, 0)
        self.original.host = self.settings["host"]
        self.factory.side_effect = lambda **kw: self.original
        with self.assertRaises(ValueError):
            self.export()
        self.no_binary()

    def test_invalid_kind_filter_status_company_and_nontext_search_refuse(self):
        for args in (("Item", "C-1", "", 0), ("mold", " C-1", "", 0), ("mold", "C-1", "x" * 141, 0),
                     ("mold", "C-1", [], 0), ("mold", "C-1", "", 2), ("modeling", "C-1", "", 1)):
            with self.assertRaises(ValueError):
                self.export(*args)
            self.no_binary()


if __name__ == "__main__":
    unittest.main()
