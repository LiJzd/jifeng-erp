"""Real aggregation/output code under native permission/connection models.

These tests inspect actual XLSX/HTML bytes; they are not positive native or GUI
acceptance. Reuse the existing payroll fixture, no new business-data generator.
"""

from io import BytesIO
import importlib.util
import json
from pathlib import Path
import sys
import unittest
from unittest.mock import patch

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class TestSalaryAnalysisOutput(unittest.TestCase):
    def setUp(self):
        fixtures = load("analysis_output_fixtures", ROOT / "tests/test_salary_history_export.py")
        self.fixture = fixtures.TestSalaryHistoryExport()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.env, self.frappe = self.fixture.env, self.fixture.env.frappe
        self.env.metas["Jifeng Wage Entry"].fields.append(
            self.fixture.F.Record(fieldname="operation", fieldtype="Data", options="", permlevel=0, mask=0))
        self.env.records["Jifeng Wage Entry", "WAGE"].operation = "装配"
        environment = self.env

        def check_permission(document, right):
            if not environment.permission(document.doctype, right, document):
                raise PermissionError(right)

        self.env.Document.check_permission = check_permission
        self.env.Document.as_dict = lambda document: dict(document)
        self.A = load("analysis_output_aggregation", ROOT / "jifeng_erp/salary_analysis.py")
        modules = patch.dict(sys.modules, {"jifeng_erp.salary_analysis": self.A,
            "jifeng_erp.salary_history_export": self.fixture.E})
        modules.start()
        self.addCleanup(modules.stop)
        self.E = load("salary_analysis_output_under_test", ROOT / "jifeng_erp/salary_analysis_output.py")
        self.context = {**self.env.context, "employee": "", "currency": ""}

    def query(self):
        return self.A.query_salary_analysis(json.dumps(self.context))

    def output(self, kind="export", version=None):
        version = version or self.query()["result_version"]
        getattr(self.E, kind + "_salary_analysis")(json.dumps(self.context), version)
        return self.frappe.response["filecontent"]

    def workbook(self, content):
        workbook = load_workbook(BytesIO(content))
        self.addCleanup(workbook.close)
        return workbook

    def test_xlsx_matches_real_settlements_three_charts_and_all_saved_sources(self):
        query = self.query()
        workbook = self.workbook(self.output())
        self.assertEqual(workbook.sheetnames, ["工资结算", "部门收入占比", "分组收入占比",
            "工序计薪金额占比", "工资条来源", "工序计薪来源", "口径说明"])
        summary = dict(workbook["工资结算"].values)
        self.assertEqual(summary["总收入"], "1120")
        self.assertEqual(summary["净应发（未证明付款）"], "1070")
        self.assertEqual(summary["币种"], "CNY")
        self.assertEqual(list(workbook["工序计薪金额占比"].values)[1:], [
            ("未归类", "1000", "89.29"), ("装配", "120", "10.71")])
        self.assertEqual(list(workbook["分组收入占比"].values)[1:], [("未归类", "1120", "100.00")])
        source = list(workbook["工序计薪来源"].values)[1]
        self.assertEqual(source[:7], ("SLIP", "EMP", "Period employee name", "装配", "120", "WAGE", "AS"))
        self.assertEqual(query["slips"][0]["department"], "Period department")
        text = str([list(sheet.values) for sheet in workbook])
        self.assertNotIn("PRIVATE", text)
        self.assertNotIn("Current changed name", text)
        self.assertEqual(self.fixture.events[-2:], [("rollback", None), ("close", None)])

    def test_more_than_one_page_export_keeps_every_saved_salary_slip(self):
        # New rows reuse actual fixture details but omit a duplicate wage source.
        for index in range(20):
            slip = self.env.slip(name=f"SECOND-{index:02}")
            slip.earnings[1].additional_salary = ""
        workbook = self.workbook(self.output())
        self.assertEqual(workbook["工资条来源"].max_row, 22)
        self.assertEqual(dict(workbook["工资结算"].values)["已提交工资条数量"], "21")

    def test_formulas_and_large_decimal_text_are_literal_without_numeric_rounding(self):
        slip = self.env.records["Salary Slip", "SLIP"]
        slip.employee_name = '=HYPERLINK("https://example.invalid")'
        # Classification becomes unresolved, but saved financial totals stay
        # visible and must retain every digit rather than round through float.
        slip.gross_pay = "9007199254740993.017"
        slip.net_pay = "9007199254740943.017"
        workbook = self.workbook(self.output())
        cell = workbook["工资条来源"].cell(2, 3)
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(cell.value.startswith("="))
        self.assertEqual(workbook["工资条来源"].cell(2, 8).value, "9007199254740993.017")

    def test_print_escapes_employee_and_operation_text_and_contains_no_script(self):
        self.env.records["Salary Slip", "SLIP"].employee_name = '<img src=x onerror="bad()">'
        self.env.records["Jifeng Wage Entry", "WAGE"].operation = "<script>bad()</script>"
        html = self.output("print")
        self.assertIn("&lt;img", html)
        self.assertIn("&lt;script&gt;", html)
        self.assertNotIn("<script>", html)
        self.assertIn("thead{display:table-header-group}", html)
        self.assertIn("净应发（未证明付款）", html)

    def test_salary_slip_and_financial_source_export_rights_are_required(self):
        self.fixture.roles.clear()
        for granted in ({}, {"Salary Slip"}, {"Salary Slip", "Additional Salary"}):
            self.fixture.export_grants = set(granted)
            with self.subTest(granted=granted), self.assertRaises(PermissionError):
                self.output()
            self.assertNotIn("filecontent", self.frappe.response)
        self.fixture.export_grants = {"Salary Slip", "Additional Salary", "Jifeng Wage Entry"}
        self.assertTrue(self.output())

    def test_print_requires_exact_parent_and_each_source_print_right(self):
        for doctype, name in (("Salary Slip", "SLIP"), ("Additional Salary", "AS"), ("Jifeng Wage Entry", "WAGE")):
            with self.subTest(doctype=doctype):
                self.env.denials.add((doctype, name, "print"))
                with self.assertRaises(PermissionError):
                    self.output("print")
                self.env.denials.clear()
                self.assertNotIn("filecontent", self.frappe.response)

    def test_field_or_company_read_restriction_cannot_be_bypassed_by_output(self):
        self.env.hidden_fields["Salary Slip", None] = {"net_pay"}
        with self.assertRaises(PermissionError):
            self.output()
        self.env.hidden_fields.clear()
        self.env.denials.add(("Company", "COMPANY", "select"))
        with self.assertRaises(PermissionError):
            self.output("print")
        self.assertNotIn("filecontent", self.frappe.response)

    def test_source_changes_after_query_require_requery_instead_of_new_file(self):
        version = self.query()["result_version"]
        self.env.records["Salary Slip", "SLIP"].net_pay = 1069
        with self.assertRaisesRegex(ValueError, "重新查询"):
            self.output(version=version)
        self.assertNotIn("filecontent", self.frappe.response)

    def test_source_or_export_revocation_before_final_snapshot_publishes_no_bytes(self):
        self.fixture.on_begin = lambda: self.env.records["Additional Salary", "AS"].update(docstatus=2)
        with self.assertRaises(ValueError):
            self.output()
        self.assertNotIn("filecontent", self.frappe.response)
        self.env.records["Additional Salary", "AS"].docstatus = 1
        self.fixture.on_begin = lambda: self.fixture.roles.clear()
        with self.assertRaises(PermissionError):
            self.output()
        self.assertNotIn("filecontent", self.frappe.response)

    def test_cancelled_slips_are_excluded_and_empty_scope_keeps_headers(self):
        self.env.records["Salary Slip", "SLIP"].docstatus = 2
        workbook = self.workbook(self.output())
        self.assertEqual(dict(workbook["工资结算"].values)["净应发（未证明付款）"], "0")
        self.assertEqual(workbook["工资条来源"].max_row, 1)
        self.assertEqual(workbook["工序计薪来源"].max_row, 1)
        self.assertIn("当前筛选无来源记录", self.output("print"))

    def test_mixed_currencies_and_guest_never_produce_an_output(self):
        slip = self.env.slip(name="USD-SLIP")
        self.env.add("Currency", "USD")
        slip.currency = "USD"
        slip.earnings[1].additional_salary = ""
        with self.assertRaisesRegex(ValueError, "币种"):
            self.output(version="a" * 64)
        self.frappe.session.user = "Guest"
        with self.assertRaises(PermissionError):
            self.output("print", version="a" * 64)
        self.assertNotIn("filecontent", self.frappe.response)

    def test_two_output_endpoints_are_get_only(self):
        self.assertIn(("export_salary_analysis", {"methods": ["GET"]}), self.env.whitelisted)
        self.assertIn(("print_salary_analysis", {"methods": ["GET"]}), self.env.whitelisted)


if __name__ == "__main__":
    unittest.main()
