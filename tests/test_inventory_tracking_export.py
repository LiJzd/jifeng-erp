"""Real tracking helpers and literal XLSX with independent-connection models."""

from contextlib import contextmanager
from copy import deepcopy
from io import BytesIO
import importlib.util
from pathlib import Path
import sys
import types
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

ROOT = Path(__file__).resolve().parents[1]


def load(name, relative):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


F = load("tracking_export_fixture", "tests/test_inventory_tracking.py")


class Database:
    def __init__(self, model, **settings):
        self.__dict__.update(settings)
        self.model = model
        self.begin = Mock(side_effect=lambda **kw: model.event("begin", kw))
        self.rollback = Mock(side_effect=lambda: model.event("rollback"))
        self.close = Mock(side_effect=lambda: model.event("close"))

    def get_single_value(self, doctype, field, *, cache):
        self.model.assertEqual((doctype, field, cache), ("System Settings", "apply_strict_user_permissions", False))
        return self.model.old["strict"] if self.model.old and self is self.model.original else self.model.strict


class TestTrackingExport(unittest.TestCase):
    def setUp(self):
        self.fixture = F.TestInventoryTracking("runTest")
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.env, self.T, self.frappe = self.fixture.env, self.fixture.module, self.fixture.frappe
        self.frappe.permissions = types.SimpleNamespace(can_export=lambda dt: (dt, "export") not in self.env.role_denials)
        tool, modeling = types.ModuleType("jifeng_erp.tool_archive"), types.ModuleType("jifeng_erp.plate_modeling")
        tool.DOCTYPES, modeling.DOCTYPE = {}, "Jifeng Plate Modeling"
        modules = patch.dict(sys.modules, {"jifeng_erp.tool_archive": tool, "jifeng_erp.plate_modeling": modeling,
                                          "jifeng_erp.inventory_tracking": self.T})
        modules.start()
        self.addCleanup(modules.stop)
        self.tools = load("tracking_export_snapshot", "jifeng_erp/tool_archive_export.py")
        module = patch.dict(sys.modules, {"jifeng_erp.tool_archive_export": self.tools})
        module.start()
        self.addCleanup(module.stop)
        self.events, self.failures, self.old = [], set(), None
        self.strict, self.cached_strict = 0, 0
        self.frappe.get_system_settings = lambda key: self.cached_strict
        settings = {"socket": "/synthetic/stock.sock", "host": "synthetic-primary", "port": 13316,
                    "user": "synthetic-export", "password": "synthetic-only", "cur_db_name": "synthetic-db"}
        self.original = Database(self, **settings)
        self.frappe.local = F.Record(db=self.original, site="synthetic.localhost", conf=F.Record(
            db_type="mariadb", db_socket=settings["socket"], db_host=settings["host"], db_port=settings["port"],
            db_user=settings["user"], db_password=settings["password"], db_name=settings["cur_db_name"]))
        self.factory = Mock(side_effect=lambda **kw: Database(self, **kw))
        database = types.ModuleType("frappe.database")
        database.get_db = self.factory
        dbpatch = patch.dict(sys.modules, {"frappe.database": database})
        dbpatch.start()
        self.addCleanup(dbpatch.stop)
        self.auth = {"User": [{"name": self.frappe.session.user, "enabled": 1, "user_type": "System User"}],
                     "Has Role": [{"parenttype": "User", "parent": self.frappe.session.user, "role": "Stock Manager"}]}
        self.old_get_all, self.old_get_list, self.old_get_doc = self.fixture.get_all, self.fixture.get_list, self.fixture.get_doc
        self.frappe.get_all = Mock(side_effect=self.get_all)
        self.frappe.get_list = Mock(side_effect=self.get_list)
        self.frappe.get_doc = Mock(side_effect=self.get_doc)
        self.methods = []
        self.frappe.whitelist = lambda **kw: self.decorate(kw)
        model = self
        class Response(F.Record):
            def update(response, *args, **kw):
                model.assertIs(model.frappe.local.db, model.original)
                model.assertEqual(model.events[-2:], [("rollback", None), ("close", None)])
                super().update(*args, **kw)
        self.frappe.response = Response()
        self.E = load("tracking_export_subject", "jifeng_erp/inventory_tracking_export.py")

    def decorate(self, options):
        self.methods.append(options)
        return lambda fn: fn

    def event(self, name, options=None):
        self.events.append((name, options))
        if name in self.failures:
            raise RuntimeError("synthetic cleanup failure")

    @contextmanager
    def view(self):
        current = self.env.documents
        if self.old and self.frappe.local.db is self.original:
            self.env.documents = self.old["documents"]
        try:
            yield
        finally:
            self.env.documents = current

    def get_all(self, doctype, filters=None, fields=None, **kwargs):
        with self.view():
            if doctype in (self.T.catalog.PROFILE, "Item", "File", "Company", "Warehouse", "Stock Ledger Entry", "Stock Entry"):
                return self.old_get_all(doctype, filters, fields, **kwargs)
            auth = self.old["auth"] if self.old and self.frappe.local.db is self.original else self.auth
            def matches(row):
                return all(row.get(k) in v[1] if isinstance(v, list) and v[0] == "in" else row.get(k) == v
                           for k, v in (filters or {}).items())
            return [F.Record(deepcopy(row) if fields == ["*"] else {k: row.get(k) for k in fields})
                    for row in auth.get(doctype, []) if matches(row)]

    def get_list(self, *args, **kwargs):
        with self.view():
            return self.old_get_list(*args, **kwargs)

    def get_doc(self, *args, **kwargs):
        with self.view():
            return self.old_get_doc(*args, **kwargs)

    def export(self, kind="物料", **kwargs):
        self.frappe.response.clear()
        self.E.export_tracking(kind, "JF", **kwargs)

    def book(self):
        book = load_workbook(BytesIO(self.frappe.response["filecontent"]), data_only=False)
        self.addCleanup(book.close)
        return book

    def no_binary(self):
        self.assertNotIn("filecontent", self.frappe.response)
        self.assertIs(self.frappe.local.db, self.original)

    def changed(self, action):
        self.old = {"documents": deepcopy(self.env.documents), "auth": deepcopy(self.auth), "strict": self.strict}
        save = Workbook.save
        def change(book, stream):
            value = save(book, stream)
            action()
            return value
        return patch.object(Workbook, "save", change)

    def test_exact_16_17_column_order_precision_unknowns_and_no_private_sources(self):
        for kind in ("物料", "材料"):
            self.env.profile["archive_kind"] = kind
            self.export(kind)
            book = self.book()
            sheet = book.active
            self.assertEqual([c.value for c in sheet[1]], [self.E.LABELS[f] for f in self.T.COLUMN_FIELDS[kind]])
            self.assertEqual(sheet.max_column, 16 if kind == "物料" else 17)
            self.assertEqual([sheet.cell(2, i).value for i in (3, 4, 5, 6, 7)], ["3.00", "2.00", "6.00", "2.300000", None])
            cells = [c for s in book for row in s for c in row]
            self.assertFalse(any(c.data_type == "f" for c in cells))
            text = str([c.value for c in cells])
            for secret in ("SLE-A", "ITEM-A", "PROFILE-A", "/private/files/"):
                self.assertNotIn(secret, text)
        self.assertEqual(self.methods, [{"methods": ["GET", "POST"]}])
        self.frappe.db.commit.assert_not_called()
        self.frappe.db.set_value.assert_not_called()

    def test_complete_1001_rows_filters_and_sort_without_page_stitching(self):
        for i in range(1000):
            self.fixture.add_pair(i)
        with patch.object(self.T, "query_tracking", side_effect=AssertionError("UI pages cannot export")):
            self.export(sort={"field": "stock_qty", "direction": "desc"})
        sheet = self.book().active
        self.assertEqual(sheet.max_row, 1002)
        self.assertEqual(sheet.cell(2, 3).value, "3.00")
        self.assertEqual(sheet.cell(1002, 3).value, "0")
        self.export(column_filters=[{"field": "stock_qty", "operator": "in", "value": ["0"]}])
        self.assertEqual(self.book().active.max_row, 1001)

    def test_formula_strings_are_literals_and_unicode_cell_limit_is_enforced(self):
        self.env.profile.update(display_name="=HYPERLINK(\"bad\")", legacy_number="0007", supplier="@formula")
        self.export()
        sheet = self.book().active
        self.assertEqual((sheet.cell(2, 2).value, sheet.cell(2, 12).data_type), ("0007", "s"))
        self.assertEqual(sheet.cell(2, 12).value, '=HYPERLINK("bad")')
        self.env.profile.update(notes="😀" * 16384)
        with self.assertRaises(ValueError):
            self.export()
        self.no_binary()

    def test_role_export_denial_even_without_rows_and_guest_refusal(self):
        for doctype in (self.T.catalog.PROFILE, "Stock Ledger Entry"):
            self.env.role_denials.add((doctype, "export"))
            with self.assertRaises(PermissionError):
                self.export(search="no-matching-archive")
            self.no_binary()
            self.env.role_denials.clear()
        self.frappe.session.user = "Guest"
        with self.assertRaises(PermissionError):
            self.export()
        self.no_binary()

    def test_document_denials_cover_sources_excluded_by_column_filter(self):
        for doctype, name in ((self.T.catalog.PROFILE, "PROFILE-A"), ("Stock Ledger Entry", "SLE-A")):
            self.env.doc_denials.add((doctype, name, "export"))
            with self.assertRaises(PermissionError):
                self.export(column_filters=[{"field": "stock_qty", "operator": "in", "value": ["0"]}])
            self.no_binary()
            self.env.doc_denials.clear()

    def test_independent_snapshot_rejects_commit_hidden_by_original_connection(self):
        with self.changed(lambda: self.env.profile.update(notes="changed after serialization")):
            with self.assertRaises(ValueError):
                self.export()
        self.no_binary()
        self.assertEqual(self.events[-3:], [("begin", {"read_only": True}), ("rollback", None), ("close", None)])
        self.assertIsNot(self.factory.return_value, self.original)

    def test_permission_change_during_serialization_refuses_entire_file(self):
        with self.changed(lambda: self.auth["Has Role"].clear()):
            with self.assertRaises(PermissionError):
                self.export()
        self.no_binary()

    def test_initial_stale_strict_permission_cache_refuses_before_serialization(self):
        for raw, cached in ((1, 0), (0, 1), (0, "0")):
            self.strict, self.cached_strict = raw, cached
            with patch.object(self.E, "_checked") as collect:
                with self.assertRaises(PermissionError):
                    self.export()
                collect.assert_not_called()
            self.no_binary()
        self.strict, self.cached_strict = None, None
        self.export()
        self.assertEqual(self.book().active.max_row, 2)

    def test_committed_strict_permission_change_hidden_by_original_view_refuses(self):
        with self.changed(lambda: setattr(self, "strict", 1)):
            with self.assertRaises(PermissionError):
                self.export()
        self.no_binary()
        self.assertEqual(self.cached_strict, 0)
        self.assertEqual(self.events[-3:], [("begin", {"read_only": True}), ("rollback", None), ("close", None)])

    def test_read_mask_changes_cannot_recover_optional_values_in_export(self):
        self.env.hidden[self.T.catalog.PROFILE] = {"supplier"}
        self.export()
        self.assertIsNone(self.book().active.cell(2, 16).value)
        with self.changed(lambda: self.env.hidden[self.T.catalog.PROFILE].add("notes")):
            with self.assertRaises(ValueError):
                self.export()
        self.no_binary()

    def test_cleanup_failure_or_invalid_conditions_never_publish_bytes(self):
        for action in ("rollback", "close"):
            self.failures.add(action)
            with self.assertRaises(RuntimeError):
                self.export()
            self.no_binary()
            self.failures.clear()
        with patch.object(self.T, "_checked_tracking") as collect:
            with self.assertRaises(ValueError):
                self.export(sort={"field": "item_code", "direction": "asc"})
            collect.assert_not_called()
        self.no_binary()

    def test_complete_row_overflow_refuses_and_empty_scope_is_valid(self):
        self.fixture.add_pair(1)
        with patch.object(self.E, "MAX_EXPORT_ROWS", 1):
            with self.assertRaises(ValueError):
                self.export()
        self.no_binary()
        self.export(search="no-matching-archive")
        self.assertEqual(self.book().active.max_row, 1)
        self.assertEqual(self.book()["查询条件"].cell(8, 2).value, "0")

    def test_large_legal_conditions_are_preserved_in_bounded_cells_with_null_types(self):
        values = ["😀" * 990 + str(i) for i in range(20)]
        self.export(column_filters=[{"field": "notes", "operator": "in", "value": values}])
        sheet = self.book()["列筛选"]
        self.assertEqual([sheet.cell(i, 4).value for i in range(2, 22)], values)
        self.export(column_filters=[{"field": "notes", "operator": "in", "value": [None, ""]},
                                    {"field": "supplier", "operator": "not_in", "value": []}])
        sheet = self.book()["列筛选"]
        self.assertEqual([sheet.cell(i, 3).value for i in range(2, 5)], ["空值", "文本", "空选择列表"])


if __name__ == "__main__":
    unittest.main()
