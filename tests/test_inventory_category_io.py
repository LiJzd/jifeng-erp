"""Real XLSX encode/decode and scope/permission models; no browser or database."""

import base64
from copy import deepcopy
import importlib.util
from io import BytesIO
import sys
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

import test_inventory_archive as archive_fixture
import test_inventory_category_batch as batch_fixture


class TestInventoryCategoryIO(unittest.TestCase):
    def setUp(self):
        self.batch = batch_fixture.TestInventoryCategoryBatch(methodName="runTest")
        self.batch.setUp()
        self.env, self.category = self.batch.env, self.batch.module
        self.env.frappe.response = {}
        with patch.dict(sys.modules, {"frappe": self.env.frappe, "jifeng_erp.inventory_archive": self.env.module,
                                     "jifeng_erp.inventory_category": self.category}):
            self.io = self.load("category_io_under_test", "inventory_category_io.py")
            self.archive_import = self.load("archive_xlsx_under_test", "inventory_archive_import.py")

    def load(self, alias, filename):
        spec = importlib.util.spec_from_file_location(alias, archive_fixture.ROOT / "jifeng_erp" / filename)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module

    def bytes(self, workbook):
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def parse(self, workbook, **changes):
        args = {"kind": "物料", "company": "JF", "filename": "类别.xlsx", "content_base64": base64.b64encode(self.bytes(workbook)).decode()}
        args.update(changes)
        with patch.dict(sys.modules, {"jifeng_erp.inventory_archive_import": self.archive_import}):
            return self.io.parse_import(**args)

    def add_row(self, workbook, record_type, values):
        sheet = workbook["大类" if record_type == "major" else "小类"]
        row = sheet.max_row + 1
        for column, (field, _) in enumerate(self.io._columns_for(record_type), 1):
            self.io._text(sheet, row, column, values.get(field))

    def test_real_export_round_trip_is_readonly_and_keeps_leading_zero_and_zero_conversion(self):
        self.env.documents[(self.category.DICTIONARY, "M1")]["legacy_code"] = "0007"
        before = deepcopy(self.env.documents)
        self.io.export_categories("物料", "JF")
        workbook = load_workbook(BytesIO(self.env.frappe.response["filecontent"]))
        self.assertEqual(workbook.sheetnames, ["大类", "小类", "导入说明"])
        self.assertEqual(workbook["大类"].cell(2, 4).value, "0007")
        self.assertEqual(workbook["小类"].cell(3, 10).value, "0")
        parsed = self.parse(workbook)
        self.assertEqual(parsed["rows"], [])
        self.assertEqual(parsed["input_count"], 4)
        self.assertEqual(self.env.documents, before)
        self.assertEqual(self.env.saves, [])

    def test_formula_shaped_text_is_literal_and_actual_formulas_are_rejected(self):
        self.env.documents[(self.category.DICTIONARY, "M1")]["value"] = '=原文+"标签"'
        self.io.export_categories("物料", "JF")
        workbook = load_workbook(BytesIO(self.env.frappe.response["filecontent"]))
        self.assertEqual(workbook["大类"]["E2"].data_type, "s")
        self.assertEqual(self.parse(workbook)["rows"], [])
        workbook["大类"]["E2"] = "=1+1"
        with self.assertRaises(ValueError):
            self.parse(workbook)

    def test_new_parent_and_child_create_only_drafts_with_explicit_reference(self):
        workbook = self.io._workbook("物料", "JF")
        self.add_row(workbook, "major", {"reference": "新大类01", "legacy_code": "0009", "value": "新增大类"})
        self.add_row(workbook, "minor", {"reference": "新小类01", "parent_reference": "新大类01", "value": "新增小类", "conversion_value": "1.2300"})
        before = deepcopy(self.env.documents)
        parsed = self.parse(workbook)
        parent, child = parsed["rows"]
        self.assertEqual(child["parent_row_id"], parent["row_id"])
        self.assertEqual(parent["values"]["legacy_code"], "0009")
        self.assertEqual(child["values"]["conversion_value"], "1.2300")
        self.assertEqual(parent["values"]["inactive"], 0)
        self.assertEqual(self.env.documents, before)
        self.assertEqual(self.env.saves, [])

    def test_existing_edit_retains_identity_and_exact_source_version(self):
        self.io.export_categories("物料", "JF")
        workbook = load_workbook(BytesIO(self.env.frappe.response["filecontent"]))
        workbook["大类"].cell(2, 7, "新区域")
        parsed = self.parse(workbook)
        self.assertEqual(len(parsed["rows"]), 1)
        self.assertEqual(parsed["rows"][0]["name"], "M1")
        self.assertEqual(parsed["rows"][0]["values"], {"region": "新区域"})
        self.assertEqual(parsed["rows"][0]["version"], "2026-10-01T00:00:00")

    def test_numeric_code_is_rejected_instead_of_losing_leading_zeroes(self):
        workbook = self.io._workbook("物料", "JF")
        self.add_row(workbook, "major", {"reference": "M", "value": "测试", "legacy_code": "0001"})
        workbook["大类"].cell(2, 4, 1)
        with self.assertRaisesRegex(ValueError, "文本"):
            self.parse(workbook)

    def test_wrong_company_kind_and_stale_versions_refuse_import(self):
        self.io.export_categories("物料", "JF")
        workbook = load_workbook(BytesIO(self.env.frappe.response["filecontent"]))
        for change in ({"company": "OTHER"}, {"kind": "材料"}):
            with self.subTest(change=change), self.assertRaises(ValueError):
                self.parse(workbook, **change)
        self.env.documents[(self.category.DICTIONARY, "M1")]["modified"] = "2026-10-01T00:00:01"
        with self.assertRaises(RuntimeError):
            self.parse(workbook)

    def test_duplicate_or_missing_references_never_match_by_name(self):
        workbook = self.io._workbook("物料", "JF")
        self.add_row(workbook, "major", {"reference": "same", "value": "第一"})
        self.add_row(workbook, "major", {"reference": "same", "value": "第二"})
        with self.assertRaisesRegex(ValueError, "重复"):
            self.parse(workbook)
        workbook = self.io._workbook("物料", "JF")
        self.add_row(workbook, "minor", {"reference": "C", "parent_reference": "不可猜测", "value": "小类"})
        with self.assertRaisesRegex(ValueError, "引用"):
            self.parse(workbook)

    def test_export_and_import_respect_role_record_and_field_permissions(self):
        self.env.denials.add((self.category.DICTIONARY, "export"))
        with self.assertRaises(PermissionError):
            self.io.export_categories("物料", "JF")
        self.env.denials.clear()
        workbook = self.io._workbook("物料", "JF")
        self.add_row(workbook, "major", {"reference": "M", "value": "新增", "region": "不可写"})
        self.env.hidden[(self.category.DICTIONARY, "write")] = {"region"}
        with self.assertRaises(PermissionError):
            self.parse(workbook)
        self.assertEqual(self.env.saves, [])

    def test_unknown_columns_and_more_than_100_rows_refuse_without_mutation(self):
        workbook = self.io._workbook("物料", "JF")
        workbook["大类"].cell(1, 10, "公司覆盖")
        with self.assertRaises(ValueError):
            self.parse(workbook)
        workbook = self.io._workbook("物料", "JF")
        for index in range(101):
            self.add_row(workbook, "major", {"reference": "M" + str(index), "value": "新增" + str(index)})
        with self.assertRaises(ValueError):
            self.parse(workbook)
        self.assertEqual(self.env.saves, [])


if __name__ == "__main__":
    unittest.main()
