"""Spreadsheet intake preserves native prices, identifiers and draft-only writes."""

import base64
from datetime import datetime
import importlib.util
from io import BytesIO
from pathlib import Path
import sys
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]


def workbook(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    stream = BytesIO()
    book.save(stream)
    book.close()
    return base64.b64encode(stream.getvalue()).decode()


class Item(dict):
    __getattr__ = dict.get
    __setattr__ = dict.__setitem__


class TestPurchaseImport(unittest.TestCase):
    def setUp(self):
        self.frappe = Mock()
        self.frappe.ValidationError = type("ValidationError", (Exception,), {})
        self.frappe.PermissionError = type("PermissionError", (Exception,), {})
        self.frappe.DoesNotExistError = type("DoesNotExistError", (Exception,), {})
        self.frappe.flags = SimpleNamespace(mute_messages=False)
        self.frappe.response = {}
        self.frappe.session = SimpleNamespace(user="purchase@example.test")
        self.frappe.whitelist.side_effect = lambda **kwargs: lambda function: function
        self.frappe.throw.side_effect = lambda message, exception=self.frappe.ValidationError: (_ for _ in ()).throw(exception(message))
        with patch.dict(sys.modules, {"frappe": self.frappe, "frappe.model": Mock()}):
            spec = importlib.util.spec_from_file_location("purchase_import_under_test", ROOT / "jifeng_erp/purchase_import.py")
            self.module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(self.module)
        self.context = {"company": "QA", "supplier": "QA Supplier", "transaction_date": "2026-10-13"}

    def preview(self, text, native=None, resolve=None):
        def get_doc(values):
            item = Item(values["items"][0])
            def fill(for_validate):
                self.assertTrue(for_validate)
                item.update(item_name="测试商品", stock_uom="Nos", uom=item.get("uom") or "Nos", conversion_factor=1)
                if native:
                    native(item)
            return SimpleNamespace(items=[item], set_missing_item_details=fill)
        self.frappe.get_doc.side_effect = get_doc
        with patch.object(self.module, "_scope", return_value=(self.context, self.module.ROW_FIELDS)), \
                patch.object(self.module, "_resolve", side_effect=resolve or (lambda code, company: "ITEM-" + code)):
            return self.module.preview_import(self.context, text=text)

    def test_scanned_identifiers_leading_zeroes_and_explicit_zero_price(self):
        result = self.preview("商品编码/条码\t数量\t单价\n001234\t2.5\t0\n001234\t1\t8.25")
        self.assertTrue(result["can_apply"])
        self.assertEqual(result["rows"][0]["item"]["item_code"], "ITEM-001234")
        self.assertEqual(result["rows"][0]["item"]["rate"], 0)
        self.assertEqual(result["rows"][1]["item"]["amount"], 8.25)
        self.assertEqual(result["row_count"], 2)
        self.frappe.db.commit.assert_not_called()
        self.frappe.db.sql.assert_not_called()

    def test_one_bad_row_blocks_whole_batch_and_reports_original_row(self):
        result = self.preview("商品编码,数量,单价\nA,2,8\nB,0,8\nC,3,-1\nD,1")
        self.assertFalse(result["can_apply"])
        self.assertEqual(result["error_count"], 3)
        self.assertEqual([row["row_number"] for row in result["rows"] if row["error"]], [3, 4, 5])
        self.assertFalse(self.frappe.flags.mute_messages)

    def test_native_pricing_change_is_reported_before_draft_mutation(self):
        result = self.preview("item_code,qty,rate\nA,2,8", native=lambda item: item.update(rate=7))
        self.assertFalse(result["can_apply"])
        self.assertIn("定价规则", result["rows"][0]["error"])

    def test_permission_failure_hides_private_identifiers_and_restores_messages(self):
        def deny(code, company):
            raise self.frappe.PermissionError("private supplier detail")
        result = self.preview("item_code,qty,rate\nA,1,0", resolve=deny)
        self.assertFalse(result["can_apply"])
        self.assertNotIn("private", result["rows"][0]["error"])
        self.assertFalse(self.frappe.flags.mute_messages)

    def test_dates_units_and_warehouses_can_be_omitted(self):
        result = self.preview("item_code,qty,rate\nA,1,8")
        item = result["rows"][0]["item"]
        self.assertEqual(item["schedule_date"], "2026-10-13")
        self.assertEqual(item["uom"], "Nos")

    def test_early_delivery_date_and_nonfinite_numbers_are_row_errors(self):
        result = self.preview("item_code,qty,rate,schedule_date\nA,1,8,2026-10-12\nB,NaN,8,2026-10-13\nC,1,Infinity,2026-10-13")
        self.assertEqual(result["error_count"], 3)

    def test_native_hidden_nonempty_field_does_not_leak(self):
        with patch.object(self.module, "ROW_FIELDS", self.module.ROW_FIELDS | {"private_value"}):
            with patch.object(self.module, "_scope", return_value=(self.context, self.module.ROW_FIELDS - {"private_value"})), \
                    patch.object(self.module, "_resolve", return_value="ITEM"):
                item = Item(item_code="ITEM", qty=1, rate=8, private_value="secret")
                self.frappe.get_doc.return_value = SimpleNamespace(items=[item], set_missing_item_details=Mock())
                result = self.module.preview_import(self.context, text="item_code,qty,rate\nA,1,8")
        self.assertFalse(result["can_apply"])
        self.assertIsNone(result["rows"][0]["item"])

    def test_csv_supports_utf8_bom_chinese_windows_encoding_and_quoted_codes(self):
        for raw in ("商品编码,数量,单价\n\"A,B\",2,8".encode("utf-8-sig"), "商品编码,数量,单价\n中文编号,2,8".encode("gb18030")):
            result = self.module._data(filename="采购.csv", content_base64=base64.b64encode(raw).decode())
            self.assertEqual(len(result), 1)
            self.assertEqual(result[0][1]["qty"], "2")

    def test_xlsx_preserves_text_codes_and_native_excel_dates(self):
        encoded = workbook([["商品编码/条码", "数量", "单价", "交货日期"], ["00123", 2, 0, datetime(2026, 10, 14)]])
        row = self.module._data(filename="采购.xlsx", content_base64=encoded)[0][1]
        self.assertEqual(row["code"], "00123")
        self.assertEqual(self.module._date(row["schedule_date"]), "2026-10-14")

    def test_numeric_codes_are_not_silently_reconstructed(self):
        encoded = workbook([["商品编码", "数量", "单价"], [123, 1, 8]])
        row = self.module._data(filename="采购.xlsx", content_base64=encoded)[0][1]
        with self.assertRaisesRegex(self.frappe.ValidationError, "文本"):
            self.module._resolve(row["code"], "QA")

    def test_formula_cells_and_duplicate_or_unknown_headers_are_rejected(self):
        with self.assertRaisesRegex(self.frappe.ValidationError, "公式"):
            self.module._data(filename="采购.xlsx", content_base64=workbook([["商品编码", "数量", "单价"], ["A", "=1+1", 8]]))
        for text in ("商品编码,条码,数量,单价\nA,B,1,8", "商品编码,数量,备注\nA,1,8"):
            with self.assertRaisesRegex(self.frappe.ValidationError, "列名"):
                self.module._data(text=text)

    def test_row_limit_and_simultaneous_sources_are_rejected(self):
        with self.assertRaisesRegex(self.frappe.ValidationError, "500"):
            self.module._data(text="item_code,qty,rate\n" + "A,1,8\n" * 501)
        with self.assertRaisesRegex(self.frappe.ValidationError, "一种"):
            self.module._data(text="A", filename="采购.xlsx")

    def test_blank_csv_rows_keep_original_error_row_numbers(self):
        result = self.preview("item_code,qty,rate\nA,1,8\n\nB,0,8")
        self.assertEqual(result["rows"][1]["row_number"], 4)

    def test_template_roundtrip_blank_rows_do_not_count_as_orders(self):
        self.module.download_template()
        raw = self.frappe.response["filecontent"]
        book = load_workbook(BytesIO(raw))
        self.assertEqual(book.active["A2"].number_format, "@")
        for column, value in enumerate(["00123", 2, 0], 1):
            book.active.cell(2, column, value)
        stream = BytesIO()
        book.save(stream)
        book.close()
        rows = self.module._data(filename="采购.xlsx", content_base64=base64.b64encode(stream.getvalue()).decode())
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][1]["code"], "00123")

    def test_guest_cannot_preview_or_download_templates(self):
        self.frappe.session.user = "Guest"
        with self.assertRaises(self.frappe.PermissionError):
            self.module._scope(self.context)
        with self.assertRaises(self.frappe.PermissionError):
            self.module.download_template()


if __name__ == "__main__":
    unittest.main()
